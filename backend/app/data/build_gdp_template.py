"""Build the GD Puasa export template ``gdp_template.xlsx``.

Run manually when the master ``Gula Darah Terkendali (1).xlsx`` (its headers,
the Terkendali formulas in Puskesmas!S3:W3, the conditional formatting, or the
Rekap recap formulas) changes.

Output: ``backend/app/data/gdp_template.xlsx`` — a trimmed copy of the master
that keeps all three sheets (Panduan, Puskesmas, Rekap), the two Puskesmas
header rows, the column merges, the conditional formatting, and the legacy
array formulas (S3:W3 on Puskesmas, A3:AA3 on Rekap). The sample patient data
(Puskesmas rows 3..end) is stripped so the asset is tiny; the export endpoint
fills rows from row 3 down and patches the formula / CF ranges to the real row
count.

The master uses *legacy* array formulas (``<f t="array" ref=...>``) with no
``metadata.xml`` / ``calcChain`` — openpyxl round-trips them faithfully, so we
can load+save without breaking the spill.

Run from repo root::

    python -m app.data.build_gdp_template   # (cwd: backend/)
    # or: python backend/app/data/build_gdp_template.py
"""

from __future__ import annotations

import os
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[3]
SRC_PATH = Path(
    os.environ.get("GDP_MASTER_XLSX", str(REPO_ROOT / "Gula Darah Terkendali (1).xlsx"))
)
OUT_PATH = Path(__file__).resolve().parent / "gdp_template.xlsx"

# First data row in the Puskesmas sheet (rows 1-2 are headers).
DATA_START_ROW = 3
# Columns A..R (1..18) hold filled data; S..W (19..23) hold the formulas.
LAST_DATA_COL = 18


def main() -> None:
    if not SRC_PATH.exists():
        raise SystemExit(f"master workbook not found: {SRC_PATH}")

    wb = openpyxl.load_workbook(SRC_PATH, data_only=False)
    ws = wb["Puskesmas"]

    # Strip sample patient rows but keep row 3 (it carries the S3:W3 array
    # formulas). Clear only the data cells A3:R3 — leave S3:W3 untouched.
    max_row = ws.max_row
    if max_row > DATA_START_ROW:
        ws.delete_rows(DATA_START_ROW + 1, max_row - DATA_START_ROW)
    for col in range(1, LAST_DATA_COL + 1):
        ws.cell(row=DATA_START_ROW, column=col).value = None

    # delete_rows leaves per-row dimension entries behind, which re-emit 33k
    # empty <row> elements. Drop them so the asset stays small.
    for r in list(ws.row_dimensions.keys()):
        if r > DATA_START_ROW:
            del ws.row_dimensions[r]

    wb.save(OUT_PATH)
    size_kb = round(OUT_PATH.stat().st_size / 1024, 1)
    print(f"wrote {OUT_PATH} ({size_kb} KB, Puskesmas max_row={ws.max_row})")


if __name__ == "__main__":
    main()
