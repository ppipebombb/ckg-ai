"""Build the Hipertensi export template ``hipertensi_template.xlsx``.

Run manually when the master ``Kertas Bantu Hipertensi dan DM Terkendali.xlsx``
(its headers, the Terkendali MAP/LAMBDA formulas in ``Hipertensi - Puskesmas``
AE..BG, or the ``Hipertensi - Rekap`` recap formulas) changes.

Output: ``backend/app/data/hipertensi_template.xlsx`` — a trimmed copy of the
master that keeps ONLY the two ``Hipertensi - *`` sheets, the three
``Hipertensi - Puskesmas`` header rows, the column merges, and the row-4 formula
cells (the spilling ``MAP`` array formulas AF..AU and the per-row ``Status``
formulas AV..BG). The sample patient data (rows 4..end, columns A..AD) is
stripped so the asset is tiny; the export endpoint fills rows from row 4 down and
patches the array-formula bounds + translates the per-row formulas to the real
row count.

The cross-sheet refs in ``Hipertensi - Rekap`` reference the *sheet name*
(``'Hipertensi - Puskesmas'!...``) so both sheets must keep their exact names.

Run from repo root::

    python -m app.data.build_hipertensi_template   # (cwd: backend/)
    # or: python backend/app/data/build_hipertensi_template.py
"""

from __future__ import annotations

import os
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[3]
SRC_PATH = Path(
    os.environ.get(
        "HIPERTENSI_MASTER_XLSX",
        str(REPO_ROOT / "Kertas Bantu Hipertensi dan DM Terkendali.xlsx"),
    )
)
OUT_PATH = Path(__file__).resolve().parent / "hipertensi_template.xlsx"

_SHEET = "Hipertensi - Puskesmas"
_REKAP = "Hipertensi - Rekap"
# Rows 1-3 are headers; data starts at row 4.
DATA_START_ROW = 4
# Columns A..AD (1..30) hold filled data; AE..BG (31..59) hold formulas.
LAST_DATA_COL = 30


def main() -> None:
    if not SRC_PATH.exists():
        raise SystemExit(f"master workbook not found: {SRC_PATH}")

    wb = openpyxl.load_workbook(SRC_PATH, data_only=False)

    # Keep only the two Hipertensi sheets (drop DM / legacy sheets).
    for name in list(wb.sheetnames):
        if name not in (_SHEET, _REKAP):
            del wb[name]

    ws = wb[_SHEET]

    # Strip sample patient rows but keep row 4 (it carries the AF4:AU4 array
    # formulas and the AV4:BG4 per-row formulas). Clear only the data cells
    # A4:AD4 — leave AE4:BG4 untouched.
    max_row = ws.max_row
    if max_row > DATA_START_ROW:
        ws.delete_rows(DATA_START_ROW + 1, max_row - DATA_START_ROW)
    for col in range(1, LAST_DATA_COL + 1):
        ws.cell(row=DATA_START_ROW, column=col).value = None

    # delete_rows leaves per-row dimension entries behind, which re-emit empty
    # <row> elements. Drop them so the asset stays small.
    for r in list(ws.row_dimensions.keys()):
        if r > DATA_START_ROW:
            del ws.row_dimensions[r]

    wb.save(OUT_PATH)
    size_kb = round(OUT_PATH.stat().st_size / 1024, 1)
    print(f"wrote {OUT_PATH} ({size_kb} KB, {_SHEET} max_row={ws.max_row})")


if __name__ == "__main__":
    main()
