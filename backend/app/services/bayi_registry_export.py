"""Kertas Kerja Registri Bayi Kuning .xlsx export — Ikterus + Ikterus Berat.

Simpler than the HT/DM/Obesitas exports: no live formulas, no month grid — the
ikterus band, diagnosis and rujukan are text the backend already resolved. One
row per bayi. The Ikterus sheet carries two "Pemantauan" column groups (follow-up
visits); the Ikterus Berat sheet is baseline only. Small registries (tens of
rows), so plain openpyxl is fine — the raw-XML technique the siblings use exists
only to dodge openpyxl's O(n²) merge cost on their wide month grids.

Column layout follows the dirjen "Register Sheet Pasien PJB, Ikterus.xlsx"
sheets "Bayi Kuning_Ikterus" / "Bayi Kuning_Ikterus Berat" column-for-column.
"""

from __future__ import annotations

import io
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_IDENT = ["NIK", "Nama", "Jenis Kelamin", "Tanggal Lahir\n(DD/MM/YYYY)", "No. Tlp", "Alamat"]
_HASIL = ["Tanggal Berkunjung", "Ikterus", "Diagnosis", "Rujuk Eksternal"]
_PEMANTAUAN = ["Ikterus", "Diagnosis", "Rujuk Eksternal"]

_TITLE = {
    "ikterus": "KERTAS KERJA REGISTRI BAYI KUNING - IKTERUS",
    "ikterus_berat": "KERTAS KERJA REGISTRI BAYI KUNING - IKTERUS BERAT",
}
_SYARAT = {
    "ikterus": (
        "Syarat: bayi (umur 0) dengan klasifikasi MTBM 'Ikterus' atau 'Ikterus berat'. "
        "Klasifikasi diambil dari isian EPUS bila ada; jika tidak, diturunkan dari "
        "kuesioner 'Memeriksa Ikterus' (kuning; kuning <24 jam / >14 hari / sampai "
        "telapak tangan-kaki = Ikterus berat). Cut off mulai 1 Januari 2026."
    ),
    "ikterus_berat": (
        "Syarat: bayi (umur 0) dengan klasifikasi MTBM 'Ikterus berat' "
        "(kuning timbul <24 jam, atau >14 hari, atau kuning sampai telapak tangan/kaki). "
        "Cut off mulai 1 Januari 2026."
    ),
}


def _fmt_date(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return date.fromisoformat(str(iso)).strftime("%d/%m/%Y")
    except ValueError:
        return str(iso)


def build_bayi_registry_workbook(
    pk_name: str, year: int, bayi: list[dict], sheet: str
) -> bytes:
    """Render the Bayi Kuning registry (one ``sheet`` = "ikterus" |
    "ikterus_berat") for (puskesmas, year) to .xlsx bytes."""
    with_pemantauan = sheet == "ikterus"
    ncols = len(_IDENT) + len(_HASIL) + (2 * len(_PEMANTAUAN) if with_pemantauan else 0)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registri Bayi Kuning"
    ws.freeze_panes = "A6"

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    group_fill = PatternFill("solid", fgColor="BDD7EE")
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    last = get_column_letter(ncols)
    ws.cell(1, 1, _TITLE[sheet]).font = Font(bold=True, size=13)
    ws.merge_cells(f"A1:{last}1")
    ws.cell(2, 1, _SYARAT[sheet]).alignment = left
    ws.merge_cells(f"A2:{last}2")
    ws.cell(3, 1, f"Nama Puskesmas: {pk_name}").font = Font(bold=True)
    ws.cell(3, min(5, ncols), f"Tahun: {year}").font = Font(bold=True)

    # Group header (row 4) + column header (row 5).
    def grp(c1: int, c2: int, text: str) -> None:
        ws.cell(4, c1, text)
        for c in range(c1, c2 + 1):
            cell = ws.cell(4, c)
            cell.fill = group_fill
            cell.font = Font(bold=True)
            cell.alignment = center
            cell.border = border
        if c2 > c1:
            ws.merge_cells(f"{get_column_letter(c1)}4:{get_column_letter(c2)}4")

    n_ident = len(_IDENT)
    n_hasil = len(_HASIL)
    grp(1, n_ident, "Identitas Pasien")
    grp(n_ident + 1, n_ident + n_hasil, "Hasil Pemeriksaan Bayi Kuning")
    if with_pemantauan:
        base = n_ident + n_hasil
        grp(base + 1, base + 3, "Pemantauan 1 Ikterus")
        grp(base + 4, base + 6, "Pemantauan 2 Ikterus")

    headers = list(_IDENT) + list(_HASIL)
    if with_pemantauan:
        headers += list(_PEMANTAUAN) + list(_PEMANTAUAN)
    ikterus_col_label = "Ikterus berat" if sheet == "ikterus_berat" else "Ikterus"
    headers[n_ident + 1] = ikterus_col_label  # the baseline band header
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(5, i, h)
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = border

    r = 6
    for b in bayi:
        ik = b.get("ikterus") or {}
        base = ik.get("baseline") or {}
        pem = ik.get("pemantauan") or []
        vals = [
            b.get("nik"), b.get("nama"), b.get("jenis_kelamin"),
            _fmt_date(b.get("tanggal_lahir")), b.get("no_tlp"), b.get("alamat"),
            _fmt_date(b.get("tanggal_berkunjung")),
            base.get("klasifikasi"), base.get("diagnosis"), base.get("rujuk_eksternal"),
        ]
        if with_pemantauan:
            for i in range(2):
                p = pem[i] if i < len(pem) else {}
                vals += [p.get("klasifikasi"), p.get("diagnosis"), p.get("rujuk_eksternal")]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(r, i, "" if v is None else str(v))
            cell.border = border
            cell.alignment = left if i in (2, 6) else center
        r += 1

    widths = [20, 22, 12, 14, 14, 30, 16, 16, 26, 22]
    if with_pemantauan:
        widths += [16, 26, 22, 16, 26, 22]
    for i, w in enumerate(widths[:ncols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    _add_doc_sheet(wb, sheet)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _add_doc_sheet(wb: openpyxl.Workbook, sheet: str) -> None:
    ws = wb.create_sheet("Formula & Logika")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 120
    section_fill = PatternFill("solid", fgColor="BDD7EE")
    wrap = Alignment(vertical="top", wrap_text=True)
    rows = [
        ("title", "CARA MEMBACA KERTAS KERJA BAYI KUNING", ""),
        ("blank", "", ""),
        ("section", "1. SIAPA YANG MASUK DAFTAR INI?", ""),
        ("item", "Bayi baru lahir (umur 0) yang dinilai kuning",
         "Sheet 'Ikterus' memuat bayi dengan klasifikasi 'Ikterus' atau 'Ikterus berat'. "
         "Sheet 'Ikterus Berat' hanya memuat yang 'Ikterus berat'."),
        ("item", "Sumber klasifikasi",
         "Diambil dari isian klasifikasi EPUS (MTBM) bila tersedia; jika tidak, DIHITUNG "
         "backend dari kuesioner 'Memeriksa Ikterus'. Tanda sumber (EPUS / Hitung) ada di dashboard."),
        ("blank", "", ""),
        ("section", "2. ATURAN KLASIFIKASI (MTBM)", ""),
        ("item", "Tidak ada ikterus", "'Apakah bayi kuning' = Tidak."),
        ("item", "Ikterus", "Kuning, tanpa kriteria berat."),
        ("item", "Ikterus berat",
         "Kuning DAN salah satu: timbul < 24 jam, ATAU > 14 hari, ATAU kuning sampai "
         "telapak tangan/kaki."),
        ("blank", "", ""),
        ("section", "3. KOLOM", ""),
        ("item", "Tanggal Berkunjung", "Kunjungan pertama (skrining) tempat ikterus dinilai."),
        ("item", "Pemantauan 1 / 2",
         "Kunjungan ikterus berikutnya (hanya di sheet Ikterus). Ikterus Berat tidak "
         "memakai kolom pemantauan."),
        ("item", "Rujuk Eksternal",
         "Faskes tujuan rujukan pada kunjungan itu (dari tabel 'Data Rujukan External' EPUS); "
         "kosong bila tidak dirujuk."),
    ]
    r = 1
    for kind, a, b in rows:
        if kind == "title":
            ws.cell(r, 1, a).font = Font(bold=True, size=13)
        elif kind == "section":
            c = ws.cell(r, 1, a)
            c.font = Font(bold=True, size=11)
            c.fill = section_fill
            ws.cell(r, 2, "").fill = section_fill
        elif kind == "item":
            ca = ws.cell(r, 1, a)
            ca.font = Font(bold=True)
            ca.alignment = wrap
            ws.cell(r, 2, b).alignment = wrap
        r += 1
