"""Shared fast renderer for the "Full Review Diagnose" export workbooks.

The Hipertensi (``hipertensi_diagnose_export``) and GD Puasa
(``gdp_diagnose_export``) diagnose views share an identical layout: one row per
visit-day reading, grouped per patient, with ``Nama Puskesmas`` / ``Nama
Pasien`` / ``NIK`` merged vertically across each patient's block, and the two
value columns colored green/red. Only the two value-column headers, the reading
keys, the green/red rule, and the summary tally differ — callers pass those in.

Why raw XML instead of openpyxl cell-by-cell: openpyxl's ``ws.merge_cells`` is
O(n²) — it rescans every existing range on each call — so a puskesmas with tens
of thousands of patients took *hours*; the per-cell ``.fill`` / ``.border`` /
``.alignment`` setters plus ``wb.save`` pile a linear-but-heavy cost on top.
Instead we let openpyxl build a tiny *seed* workbook carrying the handful of
cell styles we need (so styles.xml / theme / structure are authored by the
tested library), read those style indices back, then emit every row plus the
``<mergeCells>`` block as raw XML. ~125k rows: hours → ~4s, byte-equivalent
output.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Indonesian month names (index 0 unused so month number maps directly).
_MONTHS_ID = [
    "",
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
]

_SHEET = "Diagnosa"
_SUMMARY_SHEET = "Ringkasan"
# Fixed 7-column layout: Puskesmas, Nama, NIK, Bulan, Tanggal, value1, value2.
_COL_PUSKESMAS, _COL_NAMA, _COL_NIK, _COL_BULAN, _COL_TANGGAL, _COL_V1, _COL_V2 = range(
    1, 8
)
_FIXED_HEADERS = ["Nama Puskesmas", "Nama Pasien", "NIK", "Bulan", "Tanggal"]
_COL_WIDTHS = {1: 24, 2: 22, 3: 20, 4: 12, 5: 14, 6: 10, 7: 10}
_DATA_START_ROW = 2

# Shared green/red fills (callers' fill rules must return one of these or None).
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
_HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Excel 1900 date system epoch (matches openpyxl): serial = (d - epoch).days.
_EXCEL_EPOCH = date(1899, 12, 30)


def build_diagnose_workbook(
    pk_name: str,
    year: int,
    rows: list,
    *,
    value_headers: tuple[str, str],
    val_keys: tuple[str, str],
    value_fills: Callable[[float | None, float | None], tuple],
    summary_extra: Callable[[int, int, int], list[tuple[str, object]]],
) -> bytes:
    """Return the .xlsx bytes for a Full Review Diagnose view.

    ``rows`` is a list of ``{"nama", "nik", "readings": [{"date", <k1>, <k2>}]}``
    where ``val_keys = (k1, k2)``. ``value_fills(v1, v2)`` returns the
    ``(fill1, fill2)`` pair (each ``GREEN`` / ``RED`` / ``None``).
    ``summary_extra(total_v1, total_v2, total_bacaan)`` returns the summary rows
    appended after ``Jumlah Pasien``.
    """
    seed, s, diag_part, summary_part = _build_seed()
    letter = {c: get_column_letter(c) for c in range(1, 8)}
    fill_style = {None: s["center"], GREEN: s["green"], RED: s["red"]}
    k1, k2 = val_keys

    # ── Diagnosa sheetData (header + one row per reading) ──────────────────
    body = ['<row r="1">']
    for col, label in enumerate([*_FIXED_HEADERS, *value_headers], start=1):
        body.append(
            f'<c r="{letter[col]}1" s="{s["header"]}" t="inlineStr">'
            f"<is><t>{_xml_escape(label)}</t></is></c>"
        )
    body.append("</row>")

    merges: list[str] = []
    total_pasien = total_v1 = total_v2 = total_bacaan = 0
    pk_esc = _xml_escape(pk_name)
    r = _DATA_START_ROW
    for row in rows:
        readings = row["readings"]
        if not readings:
            continue
        total_pasien += 1
        block_start = r
        nama = row["nama"] or "—"
        nik = row["nik"]
        for idx, rd in enumerate(readings):
            total_bacaan += 1
            v1 = rd[k1]
            v2 = rd[k2]
            if v1 is not None:
                total_v1 += 1
            if v2 is not None:
                total_v2 += 1
            f1, f2 = value_fills(v1, v2)
            cells = [f'<row r="{r}">']
            # Nama Puskesmas / Nama Pasien / NIK: value only on the block's
            # first row (the rest of the block is empty, merged below).
            if idx == 0:
                for col, val in (
                    (_COL_PUSKESMAS, pk_esc),
                    (_COL_NAMA, _xml_escape(nama)),
                    (_COL_NIK, _xml_escape(nik)),
                ):
                    sp = ' xml:space="preserve"' if val != val.strip() else ""
                    cells.append(
                        f'<c r="{letter[col]}{r}" s="{s["name"]}" t="inlineStr">'
                        f"<is><t{sp}>{val}</t></is></c>"
                    )
            else:
                for col in (_COL_PUSKESMAS, _COL_NAMA, _COL_NIK):
                    cells.append(f'<c r="{letter[col]}{r}" s="{s["name"]}"/>')
            cells.append(
                f'<c r="D{r}" s="{s["center"]}" t="inlineStr"><is><t>'
                f"{_MONTHS_ID[rd['date'].month]}</t></is></c>"
            )
            cells.append(
                f'<c r="E{r}" s="{s["date"]}" t="n">'
                f"<v>{(rd['date'] - _EXCEL_EPOCH).days}</v></c>"
            )
            cells.append(
                f'<c r="F{r}" s="{fill_style[f1]}"><v>{_num_str(v1)}</v></c>'
                if v1 is not None
                else f'<c r="F{r}" s="{fill_style[f1]}"/>'
            )
            cells.append(
                f'<c r="G{r}" s="{fill_style[f2]}"><v>{_num_str(v2)}</v></c>'
                if v2 is not None
                else f'<c r="G{r}" s="{fill_style[f2]}"/>'
            )
            cells.append("</row>")
            body.append("".join(cells))
            r += 1
        block_end = r - 1
        if block_end > block_start:
            for col in (_COL_PUSKESMAS, _COL_NAMA, _COL_NIK):
                merges.append(
                    f'<mergeCell ref="{letter[col]}{block_start}:'
                    f'{letter[col]}{block_end}"/>'
                )

    diag_data = "<sheetData>" + "".join(body) + "</sheetData>"
    merge_xml = (
        f'<mergeCells count="{len(merges)}">' + "".join(merges) + "</mergeCells>"
        if merges
        else ""
    )
    diag_max_row = r - 1 if r > _DATA_START_ROW else 1

    # ── Ringkasan sheetData ────────────────────────────────────────────────
    summary_rows = [
        ("Nama Puskesmas", pk_name),
        ("Tahun", year),
        ("Jumlah Pasien", total_pasien),
        ("", ""),
        *summary_extra(total_v1, total_v2, total_bacaan),
    ]
    sbody = []
    for i, (label, value) in enumerate(summary_rows, start=1):
        lc = _summary_text_cell(f"A{i}", s["sum_bold"], str(label))
        if isinstance(value, (int, float)):
            vc = f'<c r="B{i}" s="{s["sum_plain"]}" t="n"><v>{_num_str(value)}</v></c>'
        else:
            vc = _summary_text_cell(f"B{i}", s["sum_plain"], str(value))
        sbody.append(f'<row r="{i}">{lc}{vc}</row>')
    summary_data = "<sheetData>" + "".join(sbody) + "</sheetData>"

    # ── Splice both sheets into the seed and re-zip ────────────────────────
    zin = zipfile.ZipFile(io.BytesIO(seed))
    diag_xml = zin.read(diag_part).decode("utf-8")
    diag_xml = re.sub(
        r"<sheetData>.*?</sheetData>", diag_data, diag_xml, count=1, flags=re.S
    )
    if merge_xml:
        diag_xml = diag_xml.replace("</sheetData>", "</sheetData>" + merge_xml, 1)
    diag_xml = re.sub(
        r'(<dimension ref="A1:G)\d+(")', rf"\g<1>{diag_max_row}\g<2>", diag_xml
    )
    sum_xml = zin.read(summary_part).decode("utf-8")
    sum_xml = re.sub(
        r"<sheetData>.*?</sheetData>", summary_data, sum_xml, count=1, flags=re.S
    )
    sum_xml = re.sub(
        r'(<dimension ref="A1:B)\d+(")', rf"\g<1>{len(summary_rows)}\g<2>", sum_xml
    )

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in zin.infolist():
            if item.filename == diag_part:
                zf.writestr(item, diag_xml.encode("utf-8"))
            elif item.filename == summary_part:
                zf.writestr(item, sum_xml.encode("utf-8"))
            else:
                zf.writestr(item, zin.read(item.filename))
    return out.getvalue()


def _build_seed() -> tuple[bytes, dict[str, str], str, str]:
    """Build the openpyxl skeleton (freeze panes, column widths, one cell of
    each style variant) and return its bytes, the ``{variant: style-id}`` map,
    and the Diagnosa / Ringkasan worksheet part names."""
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = _SUMMARY_SHEET
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 28
    diag = wb.create_sheet(_SHEET)
    for col, width in _COL_WIDTHS.items():
        diag.column_dimensions[get_column_letter(col)].width = width
    diag.freeze_panes = "A2"

    center = Alignment(horizontal="center", vertical="center")
    top = Alignment(vertical="center")
    bold = Font(bold=True)
    # One cell per distinct style so styles.xml carries them all.
    h = diag.cell(1, 1, value="H")
    h.font, h.fill, h.alignment, h.border = bold, _HEADER_FILL, center, _BORDER
    n = diag.cell(2, 1, value="x")
    n.alignment, n.border = top, _BORDER
    cb = diag.cell(2, 4, value="x")
    cb.alignment, cb.border = center, _BORDER
    dt = diag.cell(2, 5, value=date(2024, 1, 1))
    dt.number_format, dt.alignment, dt.border = "yyyy-mm-dd", center, _BORDER
    g = diag.cell(2, 6, value=1)
    g.alignment, g.border, g.fill = center, _BORDER, GREEN
    rc = diag.cell(3, 6, value=1)
    rc.alignment, rc.border, rc.fill = center, _BORDER, RED
    sb = summary.cell(1, 1, value="L")
    sb.font = bold
    summary.cell(1, 2, value=1)

    buf = io.BytesIO()
    wb.save(buf)
    seed = buf.getvalue()

    z = zipfile.ZipFile(io.BytesIO(seed))
    diag_part = _sheet_part(z, _SHEET)
    summary_part = _sheet_part(z, _SUMMARY_SHEET)
    diag_xml = z.read(diag_part).decode("utf-8")
    sum_xml = z.read(summary_part).decode("utf-8")

    def style_id(xml: str, ref: str) -> str:
        m = re.search(rf'<c r="{ref}"[^>]*?\bs="(\d+)"', xml)
        return m.group(1) if m else "0"

    style = {
        "header": style_id(diag_xml, "A1"),
        "name": style_id(diag_xml, "A2"),
        "center": style_id(diag_xml, "D2"),
        "date": style_id(diag_xml, "E2"),
        "green": style_id(diag_xml, "F2"),
        "red": style_id(diag_xml, "F3"),
        "sum_bold": style_id(sum_xml, "A1"),
        "sum_plain": style_id(sum_xml, "B1"),
    }
    return seed, style, diag_part, summary_part


def _sheet_part(z: zipfile.ZipFile, name: str) -> str:
    """Resolve a sheet name to its ``xl/worksheets/sheetN.xml`` part."""
    wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid = None
    for tag in re.finditer(r"<sheet\b[^>]*?/?>", wb_xml):
        nm = re.search(r'name="([^"]+)"', tag.group(0))
        ri = re.search(r'r:id="([^"]+)"', tag.group(0))
        if nm and ri and nm.group(1) == name:
            rid = ri.group(1)
            break
    for tag in re.finditer(r"<Relationship\b[^>]*?/?>", rels):
        i = re.search(r'Id="([^"]+)"', tag.group(0))
        t = re.search(r'Target="([^"]+)"', tag.group(0))
        if i and t and i.group(1) == rid:
            target = t.group(1)
            return target.lstrip("/") if target.startswith("/") else "xl/" + target
    raise ValueError(f"worksheet part for {name!r} not found")


def _summary_text_cell(ref: str, style: str, text: str) -> str:
    """A string summary cell. Empty text serializes as a self-closed
    ``t="inlineStr"`` cell — exactly how openpyxl writes ``value=""``."""
    if text == "":
        return f'<c r="{ref}" s="{style}" t="inlineStr"/>'
    return (
        f'<c r="{ref}" s="{style}" t="inlineStr"><is><t>'
        f"{_xml_escape(text)}</t></is></c>"
    )


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _num_str(value: float) -> str:
    """Serialize a number the way openpyxl does: whole floats lose the ``.0``."""
    f = float(value)
    return str(int(f)) if f.is_integer() else repr(f)
