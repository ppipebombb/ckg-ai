"""xlsx export for the "Gap Tatalaksana" list.

Deliberately plain openpyxl, unlike ``hipertensi_export`` / ``dm_registry_export``:
those two are complicated because they must round-trip the dirjen V8juni2026
template and its array formulas. This list has no template and no formulas — it
is a flat follow-up worksheet a puskesmas prints or works through by phone — so
a straight ``Workbook()`` with a styled header row is the whole job.

One sheet, one row per patient, in the same order the page shows them (oldest
registration month first). Empty source fields are written as empty cells; the
reader sees a blank, the on-screen table shows an em dash.
"""

from __future__ import annotations

import io
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_SHEET = "Gap Tatalaksana"

# (header, key, width). ``rerata`` is rendered as one "sys/dia" column because
# that is how the registry legend states the interpretation threshold.
_COLUMNS: list[tuple[str, str, int]] = [
    ("No", "_no", 6),
    ("NIK", "nik", 20),
    ("Nama", "nama", 28),
    ("Jenis Kelamin", "jenis_kelamin", 14),
    ("Tanggal Lahir", "tanggal_lahir", 14),
    ("No Telp / HP", "no_tlp", 18),
    ("Alamat", "alamat", 42),
    ("Tanggal Berkunjung CKG", "tanggal_berkunjung", 20),
    ("Rerata TD (Sistol/Diastol)", "_rerata", 22),
    ("Interpretasi", "interpretasi", 16),
    ("Bulan Masuk Registri", "registration_ym", 20),
]

_MONTHS_ID = [
    "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]


def _fmt_date(raw: str | None) -> str:
    """ISO date string → ``dd-mm-yyyy``. Written as text, not a date cell: this
    sheet is read and phoned from, never computed on, and a text cell cannot be
    re-interpreted as US month/day order by a differently-localised Excel."""
    if not raw:
        return ""
    try:
        return date.fromisoformat(raw).strftime("%d-%m-%Y")
    except ValueError:
        return str(raw)


def _fmt_ym(raw: str | None) -> str:
    """``"2025-08"`` → ``"Agustus 2025"``."""
    if not raw:
        return ""
    try:
        y, m = raw.split("-")
        return f"{_MONTHS_ID[int(m)]} {y}"
    except (ValueError, IndexError):
        return str(raw)


def _num(v: float | None) -> str:
    return "" if v is None else f"{v:g}"


def _fmt_rerata(p: dict) -> str:
    sys, dia = p.get("rerata_sys"), p.get("rerata_dia")
    if sys is None and dia is None:
        return ""
    return f"{_num(sys)}/{_num(dia)}"


def build_hipertensi_gap_workbook(
    pk_name: str, year: int | None, patients: list[dict]
) -> bytes:
    """Render the Gap Tatalaksana rows to xlsx bytes. ``patients`` is the
    already-filtered list (year / search applied by the caller) so the file
    always matches what the user had on screen."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _SHEET

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Row 1: what this list is, so a printed copy is not ambiguous.
    scope = f"Tahun masuk registri: {year}" if year else "Semua tahun"
    ws.cell(1, 1, f"Gap Tatalaksana Hipertensi — {pk_name} ({scope})").font = Font(
        bold=True, size=13
    )
    ws.cell(
        2,
        1,
        "Pasien hipertensi (CKG) yang belum pernah tercatat mendapat obat "
        "antihipertensi. Sel kosong berarti data tidak tersedia di sumber.",
    )

    header_row = 4
    for idx, (title, _key, width) in enumerate(_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
        c = ws.cell(header_row, idx, title)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    for i, p in enumerate(patients, start=1):
        row = header_row + i
        values = {
            "_no": i,
            "_rerata": _fmt_rerata(p),
            "tanggal_lahir": _fmt_date(p.get("tanggal_lahir")),
            "tanggal_berkunjung": _fmt_date(p.get("tanggal_berkunjung")),
            "registration_ym": _fmt_ym(p.get("registration_ym")),
        }
        for idx, (_title, key, _width) in enumerate(_COLUMNS, start=1):
            v = values[key] if key in values else (p.get(key) or "")
            c = ws.cell(row, idx, v)
            c.alignment = center if key in ("_no", "_rerata") else left_top
            c.border = border

    ws.freeze_panes = f"A{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
