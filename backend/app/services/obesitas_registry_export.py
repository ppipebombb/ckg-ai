"""Kertas Kerja Registri Obesitas .xlsx export.

Sibling of ``lipid_registry_export`` / ``dm_registry_export`` /
``hipertensi_registry_export``; same rendering technique, different grid. One
block of N rows per patient (NIK). Identitas (A–I) and the baseline "Hasil
Pemeriksaan Antopometri" (J–M) are written once and vertically merged across the
block; the Follow Up grid (N–BU, 12 months × **5** columns) fills each month's
column-group independently — EVERY kunjungan in the month stacked in date order.
Block height N is the tallest column-group.

Layout matches the dirjen sheet column-for-column, including the sheet's decision
to file Tanggal Berkunjung and both Riwayat columns under "Identitas Pasien"
rather than under the measurement band. **There is no Jenis Obat column** — the
Obesitas sheet has none, unlike all three siblings.

A month column-group is written ONLY for months the scan emitted. Because the
scan stamps ``Pasien Missed Visit`` on the control WINDOW alone (baseline +3 … +6
— see ``obesitas_registry_scan._control_months``), a month outside that window
with no kunjungan simply has no cells, and stays blank. The live formula's
"nothing at all → Missed Visit" branch therefore only ever fires on a row the
scan deliberately emitted as a missed control. Do not "fix" the formula to stamp
every empty month: that would put the workbook at odds with the dashboard.

The derived columns are LIVE Excel formulas (with cached values, recalculated on
open): IMT (both baseline and per month) implements BB/(TB/100)²; Interpretasi
(M) implements the B1 syarat over that IMT; each month's Interpretasi hasil
implements the L16 target against the block's own baseline BB cell. A second
worksheet ("Formula & Logika") documents every rule.

Rendered the fast way (raw XML) for the reasons given in
``hipertensi_registry_export``: openpyxl's ``merge_cells`` is O(n²).
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.obesitas_registry_scan import (
    _BB_MAX_KG,
    _BB_MIN_KG,
    _FU_MISSED_VISIT,
    _FU_TERKENDALI,
    _FU_TIDAK_TERKENDALI,
    _IMT_MIN_OBESITAS_I,
    _IMT_MIN_OBESITAS_II,
    _INTERP_NORMAL,
    _INTERP_OBESITAS_I,
    _INTERP_OBESITAS_II,
    _TARGET_PENURUNAN_PCT,
    _TB_MAX_CM,
    _TB_MIN_CM,
)

_MONTHS_ID = [
    "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]

# Column layout (1-based): identitas A–I (NIK…Riwayat DM, per the sheet's own
# band), baseline antropometri J–L (BB/TB/IMT), interpretasi M, then 12
# month-blocks of 5 columns each starting at N (14). BU (73) is the last column.
_FOLLOWUP_FIRST_COL = 14
_MONTH_STRIDE = 5
_LAST_COL = _FOLLOWUP_FIRST_COL + 12 * _MONTH_STRIDE - 1  # 73 = BU
_INTERP_COL = 13  # M

_IDENT_HEADERS = [
    "NIK", "Nama", "Jenis Kelamin", "Tanggal Lahir\n(DD/MM/YYYY)", "No. Tlp",
    "Alamat", "Tanggal Berkunjung", "Riwayat diagnosis HT",
    "Riwayat diagnosis DM",
]
_BASELINE_HEADERS = ["BB\n(kg)", "TB\n(cm)", "IMT\n(kg/m2)"]
_MONTH_SUBHEADERS = [
    "Tanggal Pemeriksaan\n(DD/MM/YYYY)", "BB\n(kg)", "TB\n(cm)", "IMT\n(kg/m2)",
    "Interpretasi hasil",
]

_SHEET = "Registri Obesitas"
_DATE_FMT = "DD/MM/YYYY"
_DATA_START = 8
_GROUP_ROW, _MONTH_ROW, _COL_ROW = 5, 6, 7

# Excel 1900 date system epoch (matches openpyxl): serial = (d - epoch).days.
_EXCEL_EPOCH = date(1899, 12, 30)

# Threshold literals for the formulas, rendered without a trailing ".0".
_IMT_1 = int(_IMT_MIN_OBESITAS_I)
_IMT_2 = int(_IMT_MIN_OBESITAS_II)
_TARGET = int(_TARGET_PENURUNAN_PCT)


def _month_base_col(m: int) -> int:
    """First column index (1-based) of month ``m`` (1..12). Jan→14 (N)."""
    return _FOLLOWUP_FIRST_COL + (m - 1) * _MONTH_STRIDE


def build_obesitas_registry_workbook(
    pk_name: str, year: int, patients: list[dict]
) -> bytes:
    seed, s, part = _build_seed()
    letter = {c: get_column_letter(c) for c in range(1, _LAST_COL + 1)}

    body: list[str] = []
    merges: list[str] = []
    _emit_header(body, merges, s, letter, pk_name, year)

    r = _DATA_START
    for p in patients:
        r = _emit_block(body, merges, s, letter, p, r)
    max_row = r - 1 if r > _DATA_START else _COL_ROW

    sheet_data = "<sheetData>" + "".join(body) + "</sheetData>"
    merge_xml = (
        f'<mergeCells count="{len(merges)}">' + "".join(merges) + "</mergeCells>"
        if merges
        else ""
    )

    zin = zipfile.ZipFile(io.BytesIO(seed))
    xml = zin.read(part).decode("utf-8")
    # The replacement MUST be a callable, not the string itself. `sheet_data`
    # carries patient free text (nama / alamat / puskesmas name), and a STRING
    # replacement is parsed as a regex template: one backslash anywhere in that
    # text raises re.PatternError ("bad escape \D", "invalid group reference 2")
    # and 500s the whole download. A callable is returned verbatim. ``\`` in an
    # ePuskesmas address ("JL X RT.01\RW.02") is not exotic.
    xml = re.sub(
        r"<sheetData/?>(?:.*?</sheetData>)?",
        lambda _m: sheet_data,
        xml,
        count=1,
        flags=re.S,
    )
    if merge_xml:
        xml = xml.replace("</sheetData>", "</sheetData>" + merge_xml, 1)
    xml = re.sub(
        r'<dimension ref="[^"]*"/>',
        f'<dimension ref="A1:{get_column_letter(_LAST_COL)}{max_row}"/>',
        xml,
        count=1,
    )

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in zin.infolist():
            zf.writestr(
                item,
                xml.encode("utf-8")
                if item.filename == part
                else zin.read(item.filename),
            )
    return out.getvalue()


# ── header ─────────────────────────────────────────────────────────────────
def _emit_header(body, merges, s, letter, pk_name: str, year: int) -> None:
    body.append(
        f'<row r="1" ht="18" customHeight="1">'
        f'{_text(letter[1], 1, s["title"], "KERTAS KERJA REGISTRI OBESITAS")}</row>'
    )
    body.append(
        f'<row r="2">{_text(letter[1], 2, s["plain"], f"Syarat Registri Obesitas: IMT >= {_IMT_1} kg/m2 (Obesitas I: IMT {_IMT_1} s/d di bawah {_IMT_2}; Obesitas II: IMT {_IMT_2} ke atas). IMT dihitung dari BB dan TB: BB / (TB dalam meter)^2. Riwayat HT / DM ditampilkan sebagai informasi, bukan syarat masuk. Follow Up mulai bulan berikutnya setelah bulan kunjungan CKG; kontrol dijadwalkan 3-6 bulan setelahnya. Format dirjen ini berlaku mulai 1 Januari 2026. Lihat sheet \'Formula & Logika\'.")}</row>'
    )
    body.append(
        f'<row r="3">{_text(letter[1], 3, s["bold"], f"Nama Puskesmas: {pk_name}")}'
        f'{_text(letter[8], 3, s["bold"], f"Tahun: {year}")}</row>'
    )
    merges.append(f'<mergeCell ref="A1:{letter[_INTERP_COL]}1"/>')
    merges.append(f'<mergeCell ref="A2:{letter[_INTERP_COL]}2"/>')
    merges.append(f'<mergeCell ref="A3:{letter[6]}3"/>')

    grp, hdr, blkg = s["group"], s["header"], s["blank_group"]
    rows: dict[int, dict[int, str]] = {_GROUP_ROW: {}, _MONTH_ROW: {}, _COL_ROW: {}}

    def gcell(r, c, text, style):
        rows[r][c] = (
            _text(letter[c], r, style, text)
            if text
            else f'<c r="{letter[c]}{r}" s="{style}"/>'
        )

    # Row 5 group band.
    gcell(_GROUP_ROW, 1, "Identitas Pasien", grp)
    gcell(_GROUP_ROW, 10, "Hasil Pemeriksaan Antopometri", grp)
    gcell(
        _GROUP_ROW, _INTERP_COL, "Interpretasi hasil\n(pada tanggal berkunjung)", grp
    )
    gcell(_GROUP_ROW, _FOLLOWUP_FIRST_COL, "Follow Up", grp)
    merges.append(f'<mergeCell ref="A{_GROUP_ROW}:{letter[9]}{_MONTH_ROW}"/>')
    merges.append(
        f'<mergeCell ref="{letter[10]}{_GROUP_ROW}:{letter[12]}{_MONTH_ROW}"/>'
    )
    merges.append(
        f'<mergeCell ref="{letter[_INTERP_COL]}{_GROUP_ROW}:{letter[_INTERP_COL]}{_COL_ROW}"/>'
    )
    merges.append(
        f'<mergeCell ref="{letter[_FOLLOWUP_FIRST_COL]}{_GROUP_ROW}:{letter[_LAST_COL]}{_GROUP_ROW}"/>'
    )

    # Row 6 month names.
    for m in range(1, 13):
        bc = _month_base_col(m)
        gcell(_MONTH_ROW, bc, _MONTHS_ID[m], grp)
        merges.append(
            f'<mergeCell ref="{letter[bc]}{_MONTH_ROW}:'
            f'{letter[bc + _MONTH_STRIDE - 1]}{_MONTH_ROW}"/>'
        )

    # Row 7 column headers.
    for i, label in enumerate(_IDENT_HEADERS):
        gcell(_COL_ROW, 1 + i, label, hdr)
    for i, label in enumerate(_BASELINE_HEADERS):
        gcell(_COL_ROW, 10 + i, label, hdr)
    for m in range(1, 13):
        bc = _month_base_col(m)
        for i, label in enumerate(_MONTH_SUBHEADERS):
            gcell(_COL_ROW, bc + i, label, hdr)

    for r in (_GROUP_ROW, _MONTH_ROW, _COL_ROW):
        cells = [f'<row r="{r}" ht="30" customHeight="1">']
        style_for = grp if r != _COL_ROW else blkg
        for c in range(1, _LAST_COL + 1):
            cells.append(rows[r].get(c) or f'<c r="{letter[c]}{r}" s="{style_for}"/>')
        cells.append("</row>")
        body.append("".join(cells))


# ── formulas ───────────────────────────────────────────────────────────────
def _imt_formula(bb: str, tb: str) -> str:
    """IMT = BB / (TB/100)^2, rounded to 1 decimal — the precision the sheet's
    own bands are written at, and the same rounding
    ``obesitas_registry_scan.imt`` applies, so the cached value and the
    recalculated one always agree.

    COUNT-guarded on both cells: ``N("")`` is 0, and a division by an empty TB
    would render #DIV/0! across the whole column.
    """
    return (
        f"IF(AND(COUNT({bb})>0,COUNT({tb})>0,N({tb})>0),"
        f"ROUND(N({bb})/(N({tb})/100)^2,1),\"\")"
    )


def _baseline_interpretasi_formula(r: int) -> str:
    """Live Excel implementation of
    ``obesitas_registry_scan._interpretasi_baseline``.

    Reads the IMT the neighbouring cell (L) already computed rather than
    recomputing it, so the label can never disagree with the number printed next
    to it — including at the boundary, where the 1-decimal rounding decides.

    Riwayat (H, I) is deliberately absent — like the Dislipidemia register and
    unlike the HT/DM ones, a recorded diagnosis does not pin the label here.
    """
    v = f"L{r}"
    return (
        f'IF(COUNT({v})=0,"",'
        f'IF(N({v})>={_IMT_2},"{_INTERP_OBESITAS_II}",'
        f'IF(N({v})>={_IMT_1},"{_INTERP_OBESITAS_I}","{_INTERP_NORMAL}")))'
    )


def _followup_interpretasi_formula(letter, bc: int, r: int, r0: int) -> str:
    """Live Excel implementation of
    ``obesitas_registry_scan._interpretasi_followup``.

    Month columns from ``bc``: +0 Tanggal, +1 BB, +2 TB, +3 IMT. ``r0`` is the
    first row of the patient's block, i.e. where the merged baseline BB (J) sits
    — the weight the ">5% penurunan" is measured against.

    Branch order: nothing recorded at all → Missed Visit; no weight to compare
    (a height-only visit, or a patient with no baseline weight) → blank; else the
    strict >5% test. The "everything empty → Missed Visit" branch is correct here
    only because the scan emits a row for a month at all ONLY when there was a
    kunjungan or a missed CONTROL — see this module's docstring.
    """
    tgl = f"{letter[bc]}{r}"
    bb, tb = f"{letter[bc + 1]}{r}", f"{letter[bc + 2]}{r}"
    base = f"J{r0}"
    return (
        f'IF(COUNT({tgl},{bb},{tb})=0,"{_FU_MISSED_VISIT}",'
        f'IF(OR(COUNT({bb})=0,COUNT({base})=0,N({base})<=0),"",'
        f'IF(ROUND((N({base})-N({bb}))/N({base})*100,1)>{_TARGET},'
        f'"{_FU_TERKENDALI}","{_FU_TIDAK_TERKENDALI}")))'
    )


# ── per-patient block ────────────────────────────────────────────────────
def _emit_block(body, merges, s, letter, p: dict, r0: int) -> int:
    months = {int(k): (v or []) for k, v in (p.get("followup") or {}).items()}
    # One row per kunjungan (no drug stacking — this sheet has no Jenis Obat).
    n = max([1, *(len(months.get(m, [])) for m in range(1, 13))])
    r_end = r0 + n - 1
    grid: dict[int, dict[int, str]] = {r0 + i: {} for i in range(n)}

    # Identitas A–I (value on r0).
    grid[r0][1] = _text(letter[1], r0, s["text"], p.get("nik"))
    grid[r0][2] = _text(letter[2], r0, s["left"], p.get("nama"))
    grid[r0][3] = _text(letter[3], r0, s["text"], p.get("jenis_kelamin"))
    grid[r0][4] = _date_cell(letter[4], r0, s["date"], p.get("tanggal_lahir"))
    grid[r0][5] = _text(letter[5], r0, s["text"], p.get("no_tlp"))
    grid[r0][6] = _text(letter[6], r0, s["left"], p.get("alamat"))
    grid[r0][7] = _date_cell(letter[7], r0, s["date"], p.get("tanggal_berkunjung"))
    grid[r0][8] = _text(letter[8], r0, s["text"], p.get("riwayat_ht"))
    grid[r0][9] = _text(letter[9], r0, s["text"], p.get("riwayat_dm"))

    # Baseline antropometri J–L + interpretasi M (value on r0).
    grid[r0][10] = _num_cell(letter[10], r0, s["text"], p.get("bb"))
    grid[r0][11] = _num_cell(letter[11], r0, s["text"], p.get("tb"))
    grid[r0][12] = _formula_cell(
        letter[12], r0, s["text"],
        _imt_formula(f"J{r0}", f"K{r0}"),
        p.get("imt"),
    )
    grid[r0][_INTERP_COL] = _formula_cell(
        letter[_INTERP_COL], r0, s["text"],
        _baseline_interpretasi_formula(r0),
        p.get("interpretasi"),
    )

    # Follow Up months (independent stacks).
    for m in range(1, 13):
        bc = _month_base_col(m)
        rr = r0
        for v in months.get(m, []):
            grid[rr][bc] = _date_cell(letter[bc], rr, s["date"], v.get("tanggal"))
            grid[rr][bc + 1] = _num_cell(letter[bc + 1], rr, s["text"], v.get("bb"))
            grid[rr][bc + 2] = _num_cell(letter[bc + 2], rr, s["text"], v.get("tb"))
            grid[rr][bc + 3] = _formula_cell(
                letter[bc + 3], rr, s["text"],
                _imt_formula(f"{letter[bc + 1]}{rr}", f"{letter[bc + 2]}{rr}"),
                v.get("imt"),
            )
            grid[rr][bc + 4] = _formula_cell(
                letter[bc + 4], rr, s["text"],
                _followup_interpretasi_formula(letter, bc, rr, r0),
                v.get("interpretasi"),
            )
            rr += 1

    # Vertical merges for identitas + baseline (A..M) across the block.
    if n > 1:
        for col in range(1, _INTERP_COL + 1):
            merges.append(f'<mergeCell ref="{letter[col]}{r0}:{letter[col]}{r_end}"/>')

    blank = s["blank"]
    for rr in range(r0, r_end + 1):
        cells = [f'<row r="{rr}">']
        row_cells = grid[rr]
        for c in range(1, _LAST_COL + 1):
            cells.append(row_cells.get(c) or f'<c r="{letter[c]}{rr}" s="{blank}"/>')
        cells.append("</row>")
        body.append("".join(cells))
    return r_end + 1


# ── cell emitters ──────────────────────────────────────────────────────────
def _text(col_letter: str, row: int, style: str, value) -> str:
    if value is None or value == "":
        return f'<c r="{col_letter}{row}" s="{style}"/>'
    text = str(value)
    esc = _xml_escape(text)
    sp = ' xml:space="preserve"' if text != text.strip() else ""
    return (
        f'<c r="{col_letter}{row}" s="{style}" t="inlineStr">'
        f"<is><t{sp}>{esc}</t></is></c>"
    )


def _num_cell(col_letter: str, row: int, style: str, value) -> str:
    if value is None or value == "":
        return f'<c r="{col_letter}{row}" s="{style}"/>'
    return f'<c r="{col_letter}{row}" s="{style}"><v>{_num_str(value)}</v></c>'


def _formula_cell(col_letter: str, row: int, style: str, formula: str, value) -> str:
    """A live formula with the scan's computed value cached (the seed workbook
    sets fullCalcOnLoad, so Excel re-evaluates on open)."""
    f = _xml_escape(formula)
    if value is None or value == "":
        return f'<c r="{col_letter}{row}" s="{style}" t="str"><f>{f}</f><v></v></c>'
    if isinstance(value, (int, float)):
        return f'<c r="{col_letter}{row}" s="{style}"><f>{f}</f><v>{_num_str(value)}</v></c>'
    return (
        f'<c r="{col_letter}{row}" s="{style}" t="str"><f>{f}</f>'
        f"<v>{_xml_escape(str(value))}</v></c>"
    )


def _date_cell(col_letter: str, row: int, style: str, value) -> str:
    d = value if isinstance(value, date) else _parse_iso(value)
    if d is None:
        return _text(col_letter, row, style, value)
    return f'<c r="{col_letter}{row}" s="{style}"><v>{(d - _EXCEL_EPOCH).days}</v></c>'


def _parse_iso(value) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return None


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _num_str(value) -> str:
    f = float(value)
    return str(int(f)) if f.is_integer() else repr(f)


# ── "Formula & Logika" documentation sheet ─────────────────────────────────
_FORMULA_SHEET_ROWS: list[tuple[str, str, str]] = [
    # (kind, col A, col B) — kind: "title" | "section" | "item" | "blank"
    # Deliberately plain, non-technical Indonesian — read by puskesmas staff.
    ("title", "CARA MEMBACA KERTAS KERJA REGISTRI OBESITAS", ""),
    ("blank", "", ""),
    ("item", "Intinya: satu daftar, dua pertanyaan",
     "Bagian 'Hasil Pemeriksaan Antopometri' (hijau) menjawab: seberapa berat pasien SAAT skrining CKG? "
     "Bagian 'Follow Up' menjawab: pada kontrol berikutnya, apakah berat badannya sudah turun cukup banyak? "
     "Karena pertanyaannya berbeda, labelnya juga berbeda — dan itu memang benar, bukan error."),
    ("blank", "", ""),
    ("section", "1. SIAPA YANG MASUK DAFTAR INI?", ""),
    ("item", "Masuk daftar jika IMT-nya 25 ke atas",
     f"Obesitas I: IMT {_IMT_1} sampai di bawah {_IMT_2}. Obesitas II: IMT {_IMT_2} ke atas. "
     f"Pasien dengan IMT di bawah {_IMT_1} tidak ditampilkan, begitu juga pasien yang berat atau "
     "tinggi badannya tidak tercatat (IMT-nya tidak bisa dihitung)."),
    ("item", "Riwayat HT dan DM hanya informasi, bukan syarat",
     "Kolom 'Riwayat diagnosis HT' dan 'Riwayat diagnosis DM' ditampilkan supaya terlihat pasien mana yang "
     "punya penyakit penyerta — TAPI tidak menentukan siapa yang masuk daftar. Yang menentukan hanya IMT. "
     "Pasien dengan IMT 32 tetap masuk daftar walaupun belum pernah didiagnosis hipertensi maupun diabetes."),
    ("item", "Beda dengan kertas kerja Hipertensi dan Diabetes",
     "Di dua kertas kerja itu, pasien yang sudah punya diagnosis SELALU diberi label penyakitnya walaupun "
     "hasil ukur hari itu bagus. Di sini tidak: label mengikuti hasil pengukuran, karena daftar ini memang "
     "disusun dari pengukuran berat dan tinggi badan."),
    ("item", "Berlaku mulai 1 Januari 2026",
     "Format dirjen untuk registri Obesitas ini berlaku mulai 1 Januari 2026. Tahun sebelumnya boleh dipilih, "
     "tetapi pencatatan berat dan tinggi badannya biasanya jauh lebih sedikit."),
    ("blank", "", ""),
    ("section", "2. BAGAIMANA IMT DIHITUNG?", ""),
    ("item", "Rumusnya",
     "IMT = Berat Badan (kg) dibagi kuadrat Tinggi Badan (meter). Contoh: BB 78 kg, TB 165 cm = 1,65 m; "
     "IMT = 78 / (1,65 x 1,65) = 28,7. Dibulatkan 1 angka di belakang koma."),
    ("item", "Kolom IMT dihitung sendiri oleh sistem",
     "Baik ASIK maupun ePuskesmas tidak mengirimkan angka IMT untuk dewasa — ASIK hanya menyimpan berat "
     "dan tinggi badan. Jadi IMT di sini dihitung ulang dari BB dan TB. Kolom IMT di file Excel ini juga "
     "berupa rumus hidup: kalau BB atau TB dikoreksi, IMT dan labelnya ikut berubah otomatis."),
    ("item", "Kalau salah satu tidak terisi",
     "IMT tidak bisa dihitung, sehingga kolom Interpretasi dibiarkan kosong dan pasien itu tidak masuk daftar. "
     "Ini bukan berarti pasiennya tidak obesitas — berarti pengukurannya belum lengkap."),
    ("blank", "", ""),
    ("section", "3. ARTI LABEL INTERPRETASI (saat kunjungan CKG)", ""),
    ("item", _INTERP_OBESITAS_I, f"IMT {_IMT_1} sampai di bawah {_IMT_2} kg/m2."),
    ("item", _INTERP_OBESITAS_II, f"IMT {_IMT_2} kg/m2 ke atas."),
    ("item", _INTERP_NORMAL,
     f"Berat dan tinggi badan sudah diukur, tetapi IMT-nya masih di bawah {_IMT_1}. "
     "Pasien seperti ini tidak muncul di daftar. Perlu dicatat: label ini hanya berarti 'di bawah ambang "
     "obesitas' — format dirjen tidak membagi lagi wilayah di bawah 25, jadi kurus dan berat badan lebih "
     "tidak dibedakan di sini."),
    ("item", "Kosong",
     "Berat atau tinggi badan tidak tercatat pada kunjungan itu, jadi IMT tidak bisa dihitung."),
    ("blank", "", ""),
    ("section", "4. CARA BACA FOLLOW UP", ""),
    ("item", "Mulai kapan",
     "Mulai bulan SETELAH bulan skrining CKG. Bulan skrining sendiri tidak diulang di Follow Up karena "
     "sudah ada di bagian hijau."),
    ("item", "Kontrolnya 3-6 BULAN setelah skrining",
     "Format dirjen menilai target 'penurunan BB pada 3-6 bulan berikutnya'. Jadi bulan kontrolnya adalah "
     "bulan ke-3, ke-4, ke-5, dan ke-6 setelah bulan skrining CKG. Ini berbeda dengan kertas kerja "
     "Hipertensi dan Diabetes yang kontrolnya bulanan."),
    ("item", "Target tercapai",
     f"'{_FU_TERKENDALI}' — berat badan turun LEBIH DARI {_TARGET}% dibandingkan berat saat skrining CKG. "
     "Perhatikan: penurunan tepat 5,0% belum dihitung tercapai; harus lebih dari itu."),
    ("item", "Target tidak tercapai",
     f"'{_FU_TIDAK_TERKENDALI}' — penurunan berat badan {_TARGET}% atau kurang, termasuk kalau beratnya "
     "tetap atau justru naik."),
    ("item", "Pembandingnya selalu berat saat skrining CKG",
     "Bukan berat kunjungan sebelumnya. Jadi persentase penurunan selalu dihitung dari titik awal yang sama, "
     "dan bisa dibandingkan antar bulan."),
    ("item", "Missed Visit hanya di bulan kontrol",
     f"'{_FU_MISSED_VISIT}' hanya muncul di bulan ke-3 sampai ke-6 setelah skrining yang terlewat tanpa "
     "pemeriksaan. Bulan ke-1, ke-2, dan bulan ke-7 dan seterusnya sengaja dibiarkan kosong — memang tidak "
     "ada kontrol yang dijadwalkan di bulan itu, jadi tidak adil kalau dihitung terlewat. Bulan yang belum "
     "terjadi juga dibiarkan kosong."),
    ("item", "Semua kunjungan ditampilkan",
     "Kalau pasien ditimbang beberapa kali dalam sebulan, SEMUA kunjungannya ditampilkan berurutan di kolom "
     "bulan itu — masing-masing dengan tanggal, hasil, dan statusnya sendiri. Penimbangan di luar bulan "
     "kontrol tetap dicatat dan tetap dinilai."),
    ("item", "Kunjungan yang hanya mencatat tinggi badan",
     "Tetap ditampilkan (kolom TB terisi), tetapi kolom Interpretasi dibiarkan kosong — tanpa berat badan "
     "tidak ada penurunan yang bisa dihitung."),
    ("blank", "", ""),
    ("section", "5. KENAPA KOLOM FOLLOW UP BANYAK YANG KOSONG?", ""),
    ("item", "Karena berat badan jarang ditimbang ulang di kunjungan biasa",
     "Berat dan tinggi badan hampir selalu diukur saat skrining CKG, tetapi pada kunjungan berobat biasa "
     "sering tidak dicatat lagi di ePuskesmas. Kolom yang kosong berarti pengukurannya memang tidak "
     "tercatat — bukan kesalahan sistem. Justru inilah yang perlu diperbaiki: makin rutin berat badan "
     "ditimbang dan dicatat, makin berguna kertas kerja ini."),
    ("blank", "", ""),
    ("section", "6. DARI MANA ANGKANYA DIAMBIL?", ""),
    ("item", "ASIK lebih diutamakan, ePuskesmas mengisi yang kosong",
     "Identitas, berat badan, tinggi badan, dan kedua kolom Riwayat diambil dari data skrining CKG (ASIK) "
     "lebih dulu; nilai dari ePuskesmas dipakai untuk mengisi yang kosong — per kolom, bukan semua-atau-tidak."),
    ("item", "Follow Up selalu dari ePuskesmas",
     "Kunjungan setelah skrining CKG dicatat di ePuskesmas."),
    ("item", "Angka yang tidak masuk akal diabaikan",
     f"Berat badan di luar {int(_BB_MIN_KG)}-{int(_BB_MAX_KG)} kg atau tinggi badan di luar "
     f"{int(_TB_MIN_CM)}-{int(_TB_MAX_CM)} cm dianggap salah ketik dan diperlakukan sebagai "
     "'tidak terukur' — bukan dipaksa masuk ke perhitungan."),
    ("blank", "", ""),
    ("section", "7. DASAR ATURAN (untuk yang ingin memeriksa)", ""),
    ("item", f"Ambang Obesitas I ({_IMT_1}) dan Obesitas II ({_IMT_2})",
     "Legenda format dirjen 'Register Sheet Pasien HT, DM, Dislipidemia dan Obesitas', sheet 'Obesitas', "
     "bagian 'Syarat Registri Obesitas'."),
    ("item", "Target penurunan berat badan dan jadwal kontrol 3-6 bulan",
     "Legenda format dirjen pada sheet yang sama, bagian kategori evaluasi pasien Obesitas."),
    ("item", "Pengukuran berat dan tinggi badan di CKG",
     "Juknis CKG — KMK 84/2026, paket 'Gizi (BB - TB - Lingkar Perut)': "
     "https://kesprimkom.kemkes.go.id/assets/uploads/contents/others/2026kepmenkes084.pdf"),
    ("item", "Format & istilah registri",
     "Dokumen dirjen 'Register Sheet Pasien HT, DM, Dislipidemia dan Obesitas', sheet 'Obesitas'."),
]


def _add_formula_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Formula & Logika")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 130
    title_font = Font(bold=True, size=13)
    section_font = Font(bold=True, size=11)
    wrap_top = Alignment(vertical="top", wrap_text=True)
    section_fill = PatternFill("solid", fgColor="BDD7EE")
    r = 1
    for kind, a, b in _FORMULA_SHEET_ROWS:
        if kind == "title":
            c = ws.cell(r, 1, a)
            c.font = title_font
        elif kind == "section":
            c = ws.cell(r, 1, a)
            c.font = section_font
            c.fill = section_fill
            ws.cell(r, 2, "").fill = section_fill
        elif kind == "item":
            ca = ws.cell(r, 1, a)
            ca.font = Font(bold=True)
            ca.alignment = wrap_top
            cb = ws.cell(r, 2, b)
            cb.alignment = wrap_top
        r += 1


# ── seed workbook (styles + column widths + freeze) ────────────────────────
def _build_seed() -> tuple[bytes, dict[str, str], str]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _SHEET
    ws.freeze_panes = "A8"
    wb.calculation.fullCalcOnLoad = True

    widths = {
        1: 20, 2: 22, 3: 12, 4: 14, 5: 14, 6: 30,
        7: 16, 8: 18, 9: 18, 10: 10, 11: 10, 12: 11, 13: 26,
    }
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    for m in range(1, 13):
        bc = _month_base_col(m)
        for off, w in ((0, 14), (1, 10), (2, 10), (3, 11), (4, 32)):
            ws.column_dimensions[get_column_letter(bc + off)].width = w

    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    bold = Font(bold=True)
    title_font = Font(bold=True, size=13)
    group_fill = PatternFill("solid", fgColor="BDD7EE")
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    variants = {
        "title": (1, 1, title_font, None, None, None),
        "plain": (2, 1, None, None, None, None),
        "bold": (3, 1, bold, None, None, None),
        "group": (5, 1, bold, group_fill, center, border),
        "header": (7, 1, bold, header_fill, center, border),
        "blank_group": (5, 2, None, group_fill, None, border),
        "text": (8, 1, None, None, center, border),
        "left": (8, 2, None, None, left, border),
        "blank": (8, 3, None, None, None, border),
    }
    refs: dict[str, tuple[int, int]] = {}
    for name, (r, c, font, fill, align, bd) in variants.items():
        cell = ws.cell(r, c, "x")
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if align:
            cell.alignment = align
        if bd:
            cell.border = bd
        refs[name] = (r, c)
    dcell = ws.cell(9, 1, value=date(2024, 1, 1))
    dcell.number_format = _DATE_FMT
    dcell.alignment = center
    dcell.border = border
    refs["date"] = (9, 1)

    _add_formula_sheet(wb)

    buf = io.BytesIO()
    wb.save(buf)
    seed = buf.getvalue()

    z = zipfile.ZipFile(io.BytesIO(seed))
    part = _sheet_part(z, _SHEET)
    sheet_xml = z.read(part).decode("utf-8")

    def style_id(r: int, c: int) -> str:
        ref = f"{get_column_letter(c)}{r}"
        m = re.search(rf'<c r="{ref}"[^>]*?\bs="(\d+)"', sheet_xml)
        return m.group(1) if m else "0"

    styles = {name: style_id(r, c) for name, (r, c) in refs.items()}
    return seed, styles, part


def _sheet_part(z: zipfile.ZipFile, name: str) -> str:
    wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid = None
    for tag in re.finditer(r"<sheet\b[^>]*?/?>", wb_xml):
        nm = re.search(r'name="([^"]+)"', tag.group(0))
        ri = re.search(r'r:id="([^"]+)"', tag.group(0))
        if nm and ri and nm.group(1) == name:
            rid = ri.group(1)
            break
    for tag in re.finditer(r"<Relationship\b[^>]*?/?>", rels):
        i = re.search(r'Id="([^"]+)"', tag.group(0))
        t = re.search(r'Target="([^"]+)"', tag.group(0))
        if i and t and i.group(1) == rid:
            target = t.group(1)
            return target.lstrip("/") if target.startswith("/") else "xl/" + target
    raise ValueError(f"worksheet part for {name!r} not found")
