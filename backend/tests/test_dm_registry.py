"""Tests for the Diabetes Melitus registry scan: inclusion per the dirjen syarat
(DO note N15), the N15 "tidak valid" combinations, the Prediabetes band,
ASIK-first per-field sourcing with EPUS fallback, next-month follow-up with the
U17 target bands, Missed Visit fill, cross-date dedup, and the DM-only drug
filter.

Mirrors ``test_hipertensi_registry.py``: real Fernet encryption over synthetic
EPUS/ASIK blobs plus a fake Session, so the pipeline runs without Postgres. A
**past** report year (2024) is used so the Missed-Visit fill (capped at the
current month only for the current year) deterministically runs through December
regardless of when the test runs.
"""

import io
import json
import uuid
from datetime import date
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.api.deps import Principal, get_dashboard_principal, get_db
from app.core.security import encrypt_json
from app.main import app
from app.models.patient import MatchStatus
from app.services.dm_registry_export import (
    _LAST_COL,
    _MONTH_STRIDE,
    _month_base_col,
    build_dm_registry_workbook,
)
from app.services.dm_registry_scan import (
    _FU_MISSED_VISIT,
    _FU_TERKENDALI,
    _FU_TIDAK_TERKENDALI,
    _INTERP_DM,
    _INTERP_PREDIABETES,
    _INTERP_TIDAK_VALID,
    _interpretasi_baseline,
    _interpretasi_followup,
    count_interpretasi_bands,
    scan_dm_registry,
)

_YEAR = 2024  # past year → Missed Visit fill runs to December (deterministic)

_RIWAYAT_Q = "Apakah Anda pernah dinyatakan diabetes atau kencing manis oleh Dokter?"
_GDS2_Q = (
    "Gula Darah Sewaktu Kedua (GDS 2). Lakukan jika hasil GDS 1 Prediabetes "
    "(≥140-199mg/dl) atau Hiperglikemia (≥200mg/dl) dan BELUM PERNAH "
    "didiagnosis Diabetes"
)


# ── blob builders ──────────────────────────────────────────────────────────
def _epus(
    *, nik, nama="PASIEN X", gds=None, gdp=None, gd2pp=None, hba1c=None,
    icd=None, obat=None, rw_dm=None, lab=None,
):
    """Synthetic EPUS visit blob. ``lab`` is a list of (Pemeriksaan, Hasil)."""
    pem = {}
    if gds is not None:
        pem["Pemeriksaan Gula"] = str(gds)
    if gdp is not None:
        pem["Pemeriksaan Gula Darah Puasa"] = str(gdp)
    if gd2pp is not None:
        pem["Pemeriksaan Gula Darah 2 Jam PP"] = str(gd2pp)
    if hba1c is not None:
        pem["HbA1c"] = str(hba1c)

    tabs = {"Resep": (
        {"tables": {"Resep": [{"Nama Obat": o} for o in obat]}}
        if obat else {"fields": {}}
    )}
    ptm_fields = {}
    if pem:
        ptm_fields["Pemeriksaan"] = pem
    if rw_dm is not None:
        ptm_fields["Riwayat PTM pada Diri Sendiri"] = {"Penyakit Diabetes": rw_dm}
    if ptm_fields:
        tabs["PTM"] = {"fields": ptm_fields}
    if lab:
        tabs["Laboratorium"] = {
            "tables": {
                "Ubah Data Laboratorium": [
                    {"Pemeriksaan": p, "Hasil": str(h)} for p, h in lab
                ]
            }
        }
    return {
        "data_pasien": {
            "NIK": nik,
            "Nama Pasien": nama,
            "Jenis Kelamin": "P",
            "Tempat/Tgl Lahir": "BEKASI/ 27-04-1980",
            "No Telp / HP": "08123456789",
            "Alamat": "jl pakis RT RW Kel PEKAYON",
        },
        "penyakit_khusus": ([{"ICDX": icd, "Penyakit": ""}] if icd else []),
        "tabs": tabs,
    }


def _asik(
    *, gds1=None, gds2=None, gdp=None, gd2pp=None, hba1c=None, riwayat=None,
    obat=None, diagnosis=None, identitas=None, split=False,
):
    """Raw ASIK blob. With ``split=True`` the readings are spread across the
    several gula-darah layanan the live data actually uses, to prove the scan
    walks all of them rather than one form."""
    main = {}
    lanjutan = {}
    hb_form = {}
    if riwayat is not None:
        main[_RIWAYAT_Q] = riwayat
    if gds1 is not None:
        main["Gula Darah Sewaktu (GDS) (mg/dl)"] = str(gds1)
    if gds2 is not None:
        main[_GDS2_Q] = str(gds2)
    if gdp is not None:
        target = lanjutan if split else main
        key = ("Gula Darah Puasa Lanjutan (GDP) (mg/dl)" if split
               else "Gula Darah Puasa (GDP) (mg/dl)")
        target[key] = str(gdp)
    if gd2pp is not None:
        (lanjutan if split else main)["Gula Darah 2 Jam PP (mg/dl)"] = str(gd2pp)
    if hba1c is not None:
        hb_form["Pemeriksaan Hb1AC"] = str(hba1c)

    nakes = [{"layanan": "Pemeriksaan Gula Darah Dewasa Lansia", "form_data": main}]
    if lanjutan:
        nakes.append({
            "layanan": "Pemeriksaan Gula Darah Lanjutan (GDP & GD 2 PP)",
            "form_data": lanjutan,
        })
    if hb_form:
        nakes.append({
            "layanan": "Pemeriksaan Gula Darah Lanjutan (HbA1C)",
            "form_data": hb_form,
        })

    out = {"pelayanan_nakes": nakes, "pemeriksaan_mandiri": []}
    if identitas is not None:
        out["detail_data"] = identitas
    if obat is not None or diagnosis is not None:
        form = {}
        if diagnosis is not None:
            form["Diagnosis"] = diagnosis
        for i, o in enumerate(obat or []):
            form["Pilih obat" if i == 0 else f"Pilih obat ({i + 1})"] = o
        out["tatalaksana"] = {
            "rows": [{
                "kelompok_skrinning": "Pemeriksaan Gula Darah Dewasa Lansia",
                "tatalaksana_name": "Diabetes Melitus",
                "status": "selesai",
                "form_data": form,
            }]
        }
    return out


def _row(nik, fd, status, epus, *, asik=None, group=None, nama="PASIEN X"):
    return SimpleNamespace(
        id=uuid.uuid4(),
        nik=nik,
        nama=nama,
        filter_date=fd,
        match_status=status,
        match_group_id=group,
        scraped_epus_data=encrypt_json(epus),
        scraped_asik_data=encrypt_json(asik) if asik is not None else None,
    )


class _FakeResult:
    def __init__(self, rows):
        self._rows = rows

    def __iter__(self):
        return iter(self._rows)


class _FakeDB:
    """Returns the main blob rows on the first execute(), and the baseline ASIK
    rows on the second (the scan issues exactly those two queries)."""

    def __init__(self, rows):
        self._main = sorted(rows, key=lambda r: (r.nik, r.filter_date))
        self._asik = [
            SimpleNamespace(id=r.id, scraped_asik_data=r.scraped_asik_data)
            for r in rows
        ]
        self.calls = 0

    def execute(self, _stmt):
        self.calls += 1
        return _FakeResult(self._main if self.calls == 1 else self._asik)


def _scan(rows):
    payload = scan_dm_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    return {p["nik"]: p for p in payload["patients"]}


M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY


# ── pure classification: baseline bands ────────────────────────────────────
def _g(**kw):
    base = {"gds1": None, "gds2": None, "gdp": None, "gd2pp": None, "hba1c": None}
    base.update(kw)
    return base


def test_baseline_riwayat_ya_is_dm_regardless_of_readings():
    """A diagnosed patient stays DM even with a perfectly normal reading."""
    assert _interpretasi_baseline(_g(gdp=90), riwayat_ya=True) == _INTERP_DM
    assert _interpretasi_baseline(_g(), riwayat_ya=True) == _INTERP_DM


def test_baseline_dm_thresholds_per_n15():
    """N15: GDP >=126, GD2PP >=200, GDS2 >=200 each independently mean DM."""
    assert _interpretasi_baseline(_g(gdp=126), riwayat_ya=False) == _INTERP_DM
    # GD2PP needs a paired GDP to be valid; pair it with a normal one.
    assert _interpretasi_baseline(_g(gdp=90, gd2pp=200), riwayat_ya=False) == _INTERP_DM
    assert _interpretasi_baseline(
        _g(gds1=150, gds2=200), riwayat_ya=False
    ) == _INTERP_DM


def test_baseline_prediabetes_band():
    assert _interpretasi_baseline(_g(gdp=100), riwayat_ya=False) == _INTERP_PREDIABETES
    assert _interpretasi_baseline(_g(gdp=125), riwayat_ya=False) == _INTERP_PREDIABETES
    assert _interpretasi_baseline(_g(gds1=140), riwayat_ya=False) == _INTERP_PREDIABETES
    assert _interpretasi_baseline(_g(gds1=199), riwayat_ya=False) == _INTERP_PREDIABETES
    assert _interpretasi_baseline(
        _g(gdp=90, gd2pp=140), riwayat_ya=False
    ) == _INTERP_PREDIABETES


def test_baseline_normal_and_blank():
    assert _interpretasi_baseline(_g(gdp=99), riwayat_ya=False) == "Normal"
    assert _interpretasi_baseline(_g(gds1=139), riwayat_ya=False) == "Normal"
    assert _interpretasi_baseline(_g(), riwayat_ya=False) == ""


def test_dm_outranks_prediabetes():
    """A DM-level GDP wins even when another reading only reaches Prediabetes."""
    assert _interpretasi_baseline(
        _g(gds1=150, gdp=130), riwayat_ya=False
    ) == _INTERP_DM


def test_n15_invalid_gd2pp_without_gdp_is_masked():
    """GD2PP without GDP is 'tidak valid' — it must not by itself produce DM."""
    assert _interpretasi_baseline(_g(gd2pp=250), riwayat_ya=False) == _INTERP_TIDAK_VALID
    # ...but a valid reading beside it still interprets.
    assert _interpretasi_baseline(
        _g(gds1=150, gd2pp=250), riwayat_ya=False
    ) == _INTERP_PREDIABETES


def test_n15_invalid_gds2_without_gds1_is_masked():
    """GDS2 without GDS1 is 'tidak valid' — the confirmation protocol requires
    a first reading, so an orphan GDS2 cannot establish DM on its own."""
    assert _interpretasi_baseline(_g(gds2=250), riwayat_ya=False) == _INTERP_TIDAK_VALID
    assert _interpretasi_baseline(_g(gdp=110, gds2=250), riwayat_ya=False) == (
        _INTERP_PREDIABETES
    )


# ── pure classification: follow-up bands ───────────────────────────────────
def test_followup_target_tercapai():
    assert _interpretasi_followup(_g(hba1c=6.9)) == _FU_TERKENDALI
    assert _interpretasi_followup(_g(gdp=80)) == _FU_TERKENDALI
    assert _interpretasi_followup(_g(gdp=130)) == _FU_TERKENDALI
    assert _interpretasi_followup(_g(gd2pp=179)) == _FU_TERKENDALI


def test_followup_target_tidak_tercapai():
    assert _interpretasi_followup(_g(hba1c=7)) == _FU_TIDAK_TERKENDALI
    assert _interpretasi_followup(_g(gdp=131)) == _FU_TIDAK_TERKENDALI
    assert _interpretasi_followup(_g(gd2pp=180)) == _FU_TIDAK_TERKENDALI


def test_followup_hypoglycaemia_is_not_terkendali():
    """GDP <80 is undefined in the legend; hypoglycaemia is not 'terkendali'."""
    assert _interpretasi_followup(_g(gdp=79)) == _FU_TIDAK_TERKENDALI


def test_followup_any_failing_signal_wins():
    assert _interpretasi_followup(_g(hba1c=6.0, gd2pp=200)) == _FU_TIDAK_TERKENDALI


def test_followup_gds_alone_is_not_interpretable():
    """The U17 legend defines no GDS target band — a GDS-only visit is blank,
    NOT 'tidak tercapai'."""
    assert _interpretasi_followup(_g(gds1=300)) == ""
    assert _interpretasi_followup(_g()) == ""


# ── end-to-end scan ────────────────────────────────────────────────────────
def test_full_patient_with_followup_and_missed_visits():
    rows = [
        _row("A", date(_YEAR, 2, 10), M,
             _epus(nik="A", icd="E11", obat=["METFORMIN 500 MG"]),
             asik=_asik(gds1=180, gdp=140, riwayat="Ya")),
        # Jan visit BEFORE the baseline date → outside the registry window
        _row("A", date(_YEAR, 1, 5), E, _epus(nik="A", gdp=100)),
        # Mar follow-up visit (EPUS only)
        _row("A", date(_YEAR, 3, 20), E,
             _epus(nik="A", gdp=160, obat=["METFORMIN 500 MG"])),
    ]
    p = _scan(rows)["A"]
    assert p["tanggal_berkunjung"] == f"{_YEAR}-02-10"
    assert p["riwayat_dm"] == "Ya"
    assert p["gds1"] == 180 and p["gdp"] == 140
    assert p["interpretasi"] == _INTERP_DM
    assert p["obat"] == ["METFORMIN 500 MG"]
    # Follow-up starts the month AFTER the CKG month (Mar) → December.
    assert set(p["followup"].keys()) == {str(m) for m in range(3, 13)}
    assert "1" not in p["followup"] and "2" not in p["followup"]
    fu3 = p["followup"]["3"][0]
    assert fu3["tanggal"] == f"{_YEAR}-03-20"
    assert fu3["gdp"] == 160
    assert fu3["interpretasi"] == _FU_TIDAK_TERKENDALI  # GDP 160 > 130
    assert fu3["obat"] == ["METFORMIN 500 MG"]
    # A month with no visit is a Missed Visit row with no data.
    miss = p["followup"]["4"][0]
    assert miss["interpretasi"] == _FU_MISSED_VISIT
    assert miss["tanggal"] is None and miss["gdp"] is None and miss["obat"] == []


def test_no_matched_visit_is_excluded():
    """Every registry row must be anchored on a CKG (MATCHED) visit."""
    rows = [_row("N", date(_YEAR, 3, 1), E, _epus(nik="N", gdp=200, icd="E11"))]
    assert _scan(rows) == {}


def test_normal_reading_without_diagnosis_is_excluded():
    rows = [
        _row("Q", date(_YEAR, 3, 1), M, _epus(nik="Q"),
             asik=_asik(gds1=100, gdp=90, riwayat="Tidak")),
    ]
    assert _scan(rows) == {}


def test_prediabetes_row_is_included():
    rows = [
        _row("P", date(_YEAR, 3, 1), M, _epus(nik="P"),
             asik=_asik(gdp=110, riwayat="Tidak")),
    ]
    p = _scan(rows)["P"]
    assert p["riwayat_dm"] == "Tidak"
    assert p["interpretasi"] == _INTERP_PREDIABETES


def test_epus_diagnosis_alone_qualifies():
    """Riwayat may come from EPUS when ASIK never asked — sticky Ya."""
    rows = [
        _row("D", date(_YEAR, 4, 2), M, _epus(nik="D", icd="E11"), asik=_asik()),
    ]
    p = _scan(rows)["D"]
    assert p["riwayat_dm"] == "Ya"
    assert p["sources"]["riwayat_dm"] == "EPUS"
    assert p["interpretasi"] == _INTERP_DM


def test_asik_tatalaksana_diagnosis_qualifies():
    rows = [
        _row("T", date(_YEAR, 4, 2), M, _epus(nik="T"),
             asik=_asik(diagnosis="E11 - Type 2 diabetes mellitus")),
    ]
    p = _scan(rows)["T"]
    assert p["riwayat_dm"] == "Ya"
    assert p["sources"]["riwayat_dm"] == "ASIK"


def test_asik_readings_spread_across_layanan_are_all_found():
    """The CKG glucose values live in several forms — the scan must walk them
    all, not just 'Pemeriksaan Gula Darah Dewasa Lansia'."""
    rows = [
        _row("S", date(_YEAR, 5, 5), M, _epus(nik="S"),
             asik=_asik(gds1=150, gdp=90, gd2pp=210, hba1c=8.1,
                        riwayat="Tidak", split=True)),
    ]
    p = _scan(rows)["S"]
    assert p["gds1"] == 150 and p["gdp"] == 90
    assert p["gd2pp"] == 210 and p["hba1c"] == 8.1
    assert p["interpretasi"] == _INTERP_DM  # GD2PP 210, validly paired with GDP


def test_asik_wins_per_field_epus_fills_gaps():
    """Unlike hipertensi's TD pair, glucose is sourced per field: ASIK supplies
    GDS1, ePuskesmas still contributes the GDP it alone recorded."""
    rows = [
        _row("F", date(_YEAR, 6, 1), M,
             _epus(nik="F", gdp=135, gds=999),
             asik=_asik(gds1=150, riwayat="Tidak")),
    ]
    p = _scan(rows)["F"]
    assert p["gds1"] == 150      # ASIK wins the field it has
    assert p["gdp"] == 135       # EPUS keeps the field ASIK lacks
    assert p["sources"]["gula_darah"] == "ASIK"
    assert p["interpretasi"] == _INTERP_DM  # GDP 135 >= 126


def test_epus_lab_table_supplies_gdp_and_hba1c():
    rows = [
        _row("L", date(_YEAR, 6, 1), M,
             _epus(nik="L", lab=[("PEMERIKSAAN KIMIA KLINIK / Glukosa Puasa", 145),
                                 ("KIMIA KLINIK / HbA1c", 9.2)]),
             asik=_asik(riwayat="Tidak")),
    ]
    p = _scan(rows)["L"]
    assert p["gdp"] == 145
    assert p["hba1c"] == 9.2
    assert p["interpretasi"] == _INTERP_DM


def test_epus_lab_table_takes_the_last_matching_row():
    """A re-test in the same visit's lab table supersedes the earlier result —
    the same per-array-latest rule gdp_report._extract_lab_gdp uses, so the two
    DM reports never quote different GDPs for the same patient."""
    rows = [
        _row("L2", date(_YEAR, 6, 1), M,
             _epus(nik="L2", lab=[("KIMIA KLINIK / Glukosa Puasa", 95),
                                  ("KIMIA KLINIK / Glukosa Puasa", 145),
                                  ("KIMIA KLINIK / HbA1c", 5.4),
                                  ("KIMIA KLINIK / HbA1c", 9.2)]),
             asik=_asik(riwayat="Tidak")),
    ]
    p = _scan(rows)["L2"]
    assert p["gdp"] == 145
    assert p["hba1c"] == 9.2


def test_ptm_pemeriksaan_wins_over_lab_table_for_both_gdp_and_hba1c():
    """The lab table only fills PTM's gaps — one rule for GDP and HbA1C alike
    (previously HbA1C silently overrode the PTM value while GDP did not)."""
    rows = [
        _row("L3", date(_YEAR, 6, 1), M,
             _epus(nik="L3", gdp=130, hba1c=8.0,
                   lab=[("KIMIA KLINIK / Glukosa Puasa", 145),
                        ("KIMIA KLINIK / HbA1c", 9.2)]),
             asik=_asik(riwayat="Tidak")),
    ]
    p = _scan(rows)["L3"]
    assert p["gdp"] == 130
    assert p["hba1c"] == 8.0


def test_ptm_generic_gula_is_gds_only_when_gdp_absent():
    """EPUS 'Pemeriksaan Gula' is a generic field — it must not duplicate GDP."""
    with_gdp = [
        _row("G1", date(_YEAR, 6, 1), M, _epus(nik="G1", gds=150, gdp=140),
             asik=_asik(riwayat="Tidak")),
    ]
    assert _scan(with_gdp)["G1"]["gds1"] is None  # suppressed, GDP present

    without_gdp = [
        _row("G2", date(_YEAR, 6, 1), M, _epus(nik="G2", gds=150),
             asik=_asik(riwayat="Tidak")),
    ]
    assert _scan(without_gdp)["G2"]["gds1"] == 150


def test_non_dm_drugs_are_filtered_out():
    rows = [
        _row("R", date(_YEAR, 2, 1), M,
             _epus(nik="R", icd="E11",
                   obat=["AMLODIPINE 10 MG", "METFORMIN 500 MG", "PARACETAMOL"]),
             asik=_asik()),
    ]
    assert _scan(rows)["R"]["obat"] == ["METFORMIN 500 MG"]


def test_asik_obat_wins_over_epus_resep():
    rows = [
        _row("O", date(_YEAR, 2, 1), M,
             _epus(nik="O", icd="E11", obat=["METFORMIN 500 MG"]),
             asik=_asik(obat=["Glimepiride 2 mg Tablet", "Insulin Glargine"])),
    ]
    p = _scan(rows)["O"]
    assert p["obat"] == ["Glimepiride 2 mg Tablet", "Insulin Glargine"]
    assert p["sources"]["obat"] == "ASIK"


def test_baseline_month_visits_excluded_from_followup():
    """The CKG month is the baseline block; follow-up is 'kunjungan bulan
    berikutnya' — later visits in the SAME month stay out."""
    rows = [
        _row("Z", date(_YEAR, 1, 20), M, _epus(nik="Z", icd="E11"), asik=_asik()),
        _row("Z", date(_YEAR, 1, 10), E, _epus(nik="Z", gdp=150)),  # before baseline
        _row("Z", date(_YEAR, 1, 25), E, _epus(nik="Z", gdp=120)),  # same month, after
    ]
    p = _scan(rows)["Z"]
    assert "1" not in p["followup"]
    assert p["followup"]["2"][0]["interpretasi"] == _FU_MISSED_VISIT


def test_all_visits_of_month_listed_in_date_order():
    rows = [
        _row("V", date(_YEAR, 2, 1), M, _epus(nik="V", icd="E11"), asik=_asik()),
        _row("V", date(_YEAR, 3, 5), E, _epus(nik="V", gdp=180)),
        _row("V", date(_YEAR, 3, 20), E, _epus(nik="V", gdp=110)),
    ]
    fu3 = _scan(rows)["V"]["followup"]["3"]
    assert len(fu3) == 2
    assert fu3[0]["tanggal"] == f"{_YEAR}-03-05"
    assert fu3[0]["interpretasi"] == _FU_TIDAK_TERKENDALI
    assert fu3[1]["tanggal"] == f"{_YEAR}-03-20"
    assert fu3[1]["interpretasi"] == _FU_TERKENDALI


def test_followup_visit_without_glucose_is_not_listed():
    """A visit with no glucose reading at all is a no-show for THIS registry —
    it must fall through to the Missed Visit fill, not appear as a blank row."""
    rows = [
        _row("W", date(_YEAR, 2, 1), M, _epus(nik="W", icd="E11"), asik=_asik()),
        _row("W", date(_YEAR, 3, 5), E, _epus(nik="W")),  # visit, but no glucose
    ]
    fu3 = _scan(rows)["W"]["followup"]["3"]
    assert len(fu3) == 1
    assert fu3[0]["interpretasi"] == _FU_MISSED_VISIT
    assert fu3[0]["tanggal"] is None


def test_cross_date_scrape_twins_collapse():
    """Two rows sharing a match_group_id are the same real visit scraped on two
    filter_dates — only the earliest survives."""
    gid = uuid.uuid4()
    rows = [
        _row("C", date(_YEAR, 1, 5), M, _epus(nik="C", icd="E11"), asik=_asik()),
        _row("C", date(_YEAR, 3, 5), E, _epus(nik="C", gdp=150), group=gid),
        _row("C", date(_YEAR, 3, 9), E, _epus(nik="C", gdp=150), group=gid),
    ]
    fu3 = _scan(rows)["C"]["followup"]["3"]
    assert len(fu3) == 1
    assert fu3[0]["tanggal"] == f"{_YEAR}-03-05"


def test_identitas_asik_wins_epus_fills_gaps():
    rows = [
        _row("I", date(_YEAR, 3, 1), M, _epus(nik="I", nama="EPUS NAMA", icd="E11"),
             asik=_asik(identitas={
                 "data_individu": {
                     "Nama": "ASIK NAMA",
                     "Jenis Kelamin": "Perempuan",
                     "Tanggal Lahir": "25 Januari 2000",
                 },
                 "data_domisili": {},
             })),
    ]
    p = _scan(rows)["I"]
    assert p["nama"] == "ASIK NAMA"           # ASIK wins
    assert p["tanggal_lahir"] == "2000-01-25"  # ASIK Indonesian-month format
    assert p["no_tlp"] == "08123456789"        # ASIK lacks it → EPUS fills
    assert p["alamat"].startswith("jl pakis")


def test_rows_sorted_by_nama_then_nik():
    rows = [
        _row("2", date(_YEAR, 3, 1), M,
             _epus(nik="2", nama="ZULFA", icd="E11"), asik=_asik()),
        _row("1", date(_YEAR, 3, 1), M,
             _epus(nik="1", nama="ANDI", icd="E11"), asik=_asik()),
    ]
    payload = scan_dm_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    assert [p["nama"] for p in payload["patients"]] == ["ANDI", "ZULFA"]


def _sample_payload():
    """A realistic scan payload: one DM patient (riwayat + follow-up) and one
    Prediabetes patient. Produced by the real scan so the route tests exercise
    the actual serialization contract, not a hand-written stand-in."""
    rows = [
        _row("111", date(_YEAR, 2, 10), M,
             _epus(nik="111", nama="BUDI", icd="E11", obat=["METFORMIN 500 MG"]),
             asik=_asik(gds1=180, gdp=140, riwayat="Ya")),
        _row("111", date(_YEAR, 3, 20), E, _epus(nik="111", gdp=160)),
        _row("222", date(_YEAR, 2, 11), M, _epus(nik="222", nama="ANI"),
             asik=_asik(gdp=110, riwayat="Tidak")),
    ]
    return scan_dm_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)


@pytest.fixture
def dm_client(monkeypatch, stub_redis, warm_calls):
    """TestClient for the DM registry routes: puskesmas exists, admin principal,
    stub redis, and request_warm recorded instead of enqueuing Celery."""
    pid = uuid.uuid4()
    fake_db = MagicMock()
    fake_db.scalar.return_value = "PUSKESMAS TEBET"

    app.dependency_overrides[get_db] = lambda: fake_db
    app.dependency_overrides[get_dashboard_principal] = lambda: Principal(
        "admin", uuid.uuid4(), None
    )

    import app.api.routes.dm_report as dm

    monkeypatch.setattr(dm, "redis_client", lambda: stub_redis)
    monkeypatch.setattr(
        dm,
        "request_warm",
        lambda rc, cache_key, report_type, kwargs: warm_calls.append(
            (cache_key, report_type)
        ),
    )
    with TestClient(app) as client:
        yield client, pid, stub_redis
    app.dependency_overrides.clear()


def _seed(rc, pid, payload):
    from app.api.routes.dm_report import _registry_cache_key

    rc.set(_registry_cache_key(pid, _YEAR), json.dumps(payload))


def test_route_cold_cache_returns_computing_and_enqueues_warm(dm_client, warm_calls):
    client, pid, _rc = dm_client
    r = client.get("/dm-reports/registry", params={"puskesmas_id": str(pid),
                                                  "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["computing"] is True
    assert body["items"] == [] and body["total"] == 0
    assert body["cache_hit"] is False
    # The warm must be dispatched under the report_type cron.warm_one knows.
    assert [t for _k, t in warm_calls] == ["dm_registry"]


def test_route_serves_cached_payload(dm_client):
    client, pid, rc = dm_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/dm-reports/registry", params={"puskesmas_id": str(pid),
                                                  "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["computing"] is False and body["cache_hit"] is True
    assert body["total"] == 2
    assert body["total_dm"] == 1 and body["total_prediabetes"] == 1
    assert [i["nama"] for i in body["items"]] == ["ANI", "BUDI"]  # sorted by nama
    budi = next(i for i in body["items"] if i["nama"] == "BUDI")
    assert budi["puskesmas_name"] == "PUSKESMAS TEBET"
    assert budi["tahun_pelaporan"] == _YEAR
    assert budi["riwayat_dm"] == "Ya"
    assert budi["gds1"] == 180 and budi["gdp"] == 140
    assert budi["interpretasi"] == _INTERP_DM
    # followup dict keys survive the str->int coercion, values keep their shape
    assert budi["followup"]["3"][0]["gdp"] == 160
    assert budi["followup"]["3"][0]["interpretasi"] == _FU_TIDAK_TERKENDALI


def test_route_search_filters_and_rebands(dm_client):
    client, pid, rc = dm_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/dm-reports/registry", params={"puskesmas_id": str(pid),
                                                   "year": _YEAR, "q": "ani"})
    body = r.json()
    assert body["total"] == 1
    assert [i["nama"] for i in body["items"]] == ["ANI"]
    # Bands respect the filter, like `total`.
    assert body["total_dm"] == 0 and body["total_prediabetes"] == 1


def test_route_search_matches_nik(dm_client):
    client, pid, rc = dm_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/dm-reports/registry", params={"puskesmas_id": str(pid),
                                                   "year": _YEAR, "q": "111"})
    assert [i["nik"] for i in r.json()["items"]] == ["111"]


def test_route_summary(dm_client):
    client, pid, rc = dm_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/dm-reports/registry/summary",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 2
    assert body["total_dm"] == 1 and body["total_prediabetes"] == 1
    assert body["riwayat_dm_ya"] == 1
    assert body["dalam_pengobatan"] == 1  # only BUDI has a drug
    assert body["computing"] is False


def test_route_summary_cold_cache_enqueues_warm(dm_client, warm_calls):
    client, pid, _rc = dm_client
    r = client.get("/dm-reports/registry/summary",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.json()["computing"] is True
    assert [t for _k, t in warm_calls] == ["dm_registry"]


def test_route_clear_cache_deletes_and_requeues(dm_client, warm_calls):
    from app.api.routes.dm_report import _registry_cache_key

    client, pid, rc = dm_client
    _seed(rc, pid, _sample_payload())
    key = _registry_cache_key(pid, _YEAR)
    assert rc.get(key) is not None

    r = client.request("DELETE", "/dm-reports/registry/cache",
                       params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 204
    assert rc.get(key) is None
    assert [t for _k, t in warm_calls] == ["dm_registry"]


def test_route_404_when_puskesmas_missing(dm_client):
    client, pid, _rc = dm_client
    app.dependency_overrides[get_db] = lambda: SimpleNamespace(
        scalar=lambda *_a, **_k: None
    )
    r = client.get("/dm-reports/registry", params={"puskesmas_id": str(pid),
                                                   "year": _YEAR})
    assert r.status_code == 404


# ── xlsx export ────────────────────────────────────────────────────────────
def _wb(patients=None):
    payload = _sample_payload()
    content = build_dm_registry_workbook(
        "PUSKESMAS TEBET", _YEAR, patients if patients is not None else payload["patients"]
    )
    return openpyxl.load_workbook(io.BytesIO(content))


def test_export_grid_geometry_matches_the_dirjen_layout():
    """12 months x 7 columns starting at O (15) → last column CT (98). The
    baseline block has NO HbA1C column — that only exists in Follow Up."""
    assert _MONTH_STRIDE == 7
    assert _month_base_col(1) == 15
    assert _month_base_col(12) == 15 + 11 * 7
    assert _LAST_COL == 98


def test_export_sheets_and_headers():
    wb = _wb()
    assert wb.sheetnames == ["Registri Diabetes Melitus", "Formula & Logika"]
    ws = wb["Registri Diabetes Melitus"]
    # Row 7 = column headers. Baseline runs G..M, Jenis Obat at N.
    assert ws.cell(7, 1).value == "NIK"
    assert ws.cell(7, 7).value == "Tanggal Berkunjung"
    assert ws.cell(7, 8).value == "Riwayat diagnosis DM"
    assert ws.cell(7, 9).value.startswith("GDS 1")
    assert ws.cell(7, 10).value.startswith("GDS 2")
    assert ws.cell(7, 11).value.startswith("GDP")
    assert ws.cell(7, 12).value.startswith("GD 2PP")
    assert ws.cell(7, 13).value == "Interpretasi hasil"
    assert ws.cell(6, 15).value == "Januari"  # month band
    # Each month group: Tanggal, GDS, GDP, GD2PP, HbA1C, Interpretasi, Jenis Obat
    bc = _month_base_col(3)
    assert ws.cell(7, bc).value.startswith("Tanggal Pemeriksaan")
    assert [ws.cell(7, bc + i).value.split("\n")[0] for i in range(1, 5)] == [
        "GDS", "GDP", "GD2PP", "HbA1C",
    ]
    assert ws.cell(7, bc + 5).value == "Interpretasi hasil"
    assert ws.cell(7, bc + 6).value == "Jenis Obat"


def test_export_writes_patient_values():
    wb = _wb()
    ws = wb["Registri Diabetes Melitus"]
    # Patients are sorted by nama → ANI (Prediabetes) first, then BUDI.
    assert ws.cell(8, 1).value == "222"
    assert ws.cell(8, 2).value == "ANI"
    budi_row = next(
        r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == "BUDI"
    )
    assert ws.cell(budi_row, 8).value == "Ya"       # Riwayat DM
    assert ws.cell(budi_row, 9).value == 180        # GDS1
    assert ws.cell(budi_row, 11).value == 140       # GDP
    assert ws.cell(budi_row, 14).value == "METFORMIN 500 MG"
    # March follow-up: GDP 160 in the month-3 group.
    bc = _month_base_col(3)
    assert ws.cell(budi_row, bc + 2).value == 160


def test_export_interpretasi_columns_are_live_formulas():
    """The derived columns must ship as auditable formulas, not frozen text."""
    wb = _wb()
    ws = wb["Registri Diabetes Melitus"]
    baseline = ws.cell(8, 13).value
    assert isinstance(baseline, str) and baseline.startswith("=")
    assert _INTERP_DM in baseline and _INTERP_PREDIABETES in baseline
    assert _INTERP_TIDAK_VALID in baseline
    fu = ws.cell(8, _month_base_col(3) + 5).value
    assert isinstance(fu, str) and fu.startswith("=")
    assert _FU_MISSED_VISIT in fu and _FU_TERKENDALI in fu


def test_export_followup_formula_guards_empty_gdp():
    """Regression: N("") is 0, so an unguarded `<80` test would mark every empty
    month 'tidak terkendali'. The COUNT guard must be present."""
    wb = _wb()
    ws = wb["Registri Diabetes Melitus"]
    fu = ws.cell(8, _month_base_col(3) + 5).value
    assert "COUNT(" in fu
    # the <80 branch is always paired with a COUNT guard on the same cell
    assert "<80" in fu
    gdp_ref = f"{openpyxl.utils.get_column_letter(_month_base_col(3) + 2)}8"
    assert f"AND(COUNT({gdp_ref})>0,OR(N({gdp_ref})>130,N({gdp_ref})<80))" in fu


def test_export_empty_registry_still_valid_workbook():
    wb = _wb(patients=[])
    ws = wb["Registri Diabetes Melitus"]
    assert ws.cell(7, 1).value == "NIK"  # headers survive
    assert ws.cell(8, 1).value is None   # no data rows


def test_export_formula_sheet_documents_the_judgement_calls():
    """The three rules that are NOT in the dirjen legend must be written down
    for the puskesmas staff reading the file."""
    ws = _wb()["Formula & Logika"]
    text = "\n".join(
        str(ws.cell(r, c).value or "")
        for r in range(1, ws.max_row + 1)
        for c in (1, 2)
    )
    assert "Prediabetes" in text
    assert "hipoglikemia" in text.casefold()      # GDP <80 rule
    assert "GDS" in text and "KOSONG" in text     # GDS has no target band
    assert "GD2PP tanpa GDP" in text              # N15 validity
    assert "Metformin" in text                    # drug families


def test_export_route_returns_xlsx(dm_client):
    client, pid, rc = dm_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/dm-reports/registry/export",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "Registri Diabetes Melitus - PUSKESMAS TEBET" in (
        r.headers["content-disposition"]
    )
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Registri Diabetes Melitus", "Formula & Logika"]


def test_count_interpretasi_bands_sums_to_total():
    """Every qualifying row is either DM or Prediabetes — the invariant the
    summary card depends on."""
    rows = [
        _row("X1", date(_YEAR, 3, 1), M, _epus(nik="X1", icd="E11"), asik=_asik()),
        _row("X2", date(_YEAR, 3, 1), M, _epus(nik="X2"),
             asik=_asik(gdp=110, riwayat="Tidak")),
        _row("X3", date(_YEAR, 3, 1), M, _epus(nik="X3"),
             asik=_asik(gdp=200, riwayat="Tidak")),
    ]
    payload = scan_dm_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    patients = payload["patients"]
    dm, pre = count_interpretasi_bands(patients)
    assert (dm, pre) == (2, 1)
    assert dm + pre == len(patients)
