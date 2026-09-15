"""Gap Tatalaksana: the endpoint's in-memory filters and the xlsx export.

The scan-side invariant (the list IS the gap in Chart 2) lives in
``test_hipertensi_charts``; this file covers what sits between that payload and
the user.
"""

import io

import openpyxl

from app.api.routes.hipertensi_report import _gap_filtered
from app.services.hipertensi_gap_export import build_hipertensi_gap_workbook


def _p(nik, nama, ym, **kw):
    row = {
        "nik": nik,
        "nama": nama,
        "jenis_kelamin": "Perempuan",
        "tanggal_lahir": "1980-01-01",
        "no_tlp": "081234",
        "alamat": "JL MAWAR 1",
        "tanggal_berkunjung": "2025-03-04",
        "rerata_sys": 160.0,
        "rerata_dia": 100.0,
        "interpretasi": "Hipertensi",
        "registration_ym": ym,
    }
    row.update(kw)
    return row


_PAYLOAD = {
    "patients": [
        _p("1", "BUDI", "2025-03"),
        _p("2", "SITI", "2025-11"),
        _p("3", "AGUS", "2026-02"),
    ]
}


def test_filter_by_registration_year():
    """``year`` is the year the patient ENTERED the registry — an untreated
    2025 registrant stays under 2025 forever, it is not a reporting year."""
    assert [p["nik"] for p in _gap_filtered(_PAYLOAD, 2025, None)] == ["1", "2"]
    assert [p["nik"] for p in _gap_filtered(_PAYLOAD, 2026, None)] == ["3"]
    assert len(_gap_filtered(_PAYLOAD, None, None)) == 3


def test_search_matches_name_or_nik_case_insensitively():
    assert [p["nik"] for p in _gap_filtered(_PAYLOAD, None, "sit")] == ["2"]
    assert [p["nik"] for p in _gap_filtered(_PAYLOAD, None, "  BUD ".strip())] == ["1"]
    assert [p["nik"] for p in _gap_filtered(_PAYLOAD, None, "3")] == ["3"]
    assert _gap_filtered(_PAYLOAD, None, "zzz") == []


def test_filters_compose():
    assert [p["nik"] for p in _gap_filtered(_PAYLOAD, 2025, "sit")] == ["2"]


def _sheet(content: bytes):
    return openpyxl.load_workbook(io.BytesIO(content)).active


def test_export_emits_one_row_per_patient_with_readable_dates():
    ws = _sheet(build_hipertensi_gap_workbook("PKM Uji", 2025, _PAYLOAD["patients"][:2]))
    header = [c.value for c in ws[4]]
    assert header[:3] == ["No", "NIK", "Nama"]
    assert "No Telp / HP" in header and "Alamat" in header
    assert ws.max_row == 6  # header row 4 + 2 patients

    first = {header[i]: c.value for i, c in enumerate(ws[5])}
    assert first["NIK"] == "1"
    assert first["Tanggal Lahir"] == "01-01-1980"
    assert first["Tanggal Berkunjung CKG"] == "04-03-2025"
    assert first["Rerata TD (Sistol/Diastol)"] == "160/100"
    assert first["Bulan Masuk Registri"] == "Maret 2025"
    assert "PKM Uji" in ws.cell(1, 1).value


def test_export_leaves_missing_fields_blank():
    row = _p("9", "TANPA KONTAK", "2025-04", no_tlp="", alamat="",
             tanggal_lahir=None, rerata_sys=None, rerata_dia=None)
    ws = _sheet(build_hipertensi_gap_workbook("PKM Uji", None, [row]))
    header = [c.value for c in ws[4]]
    got = {header[i]: c.value for i, c in enumerate(ws[5])}
    assert got["Nama"] == "TANPA KONTAK"
    # An empty string is written as a truly blank cell (openpyxl reads it back
    # as None) — the reader sees nothing, never the literal "None".
    for blank in ("No Telp / HP", "Alamat", "Tanggal Lahir", "Rerata TD (Sistol/Diastol)"):
        assert got[blank] in (None, "")
    assert "Semua tahun" in ws.cell(1, 1).value


def test_export_handles_an_empty_list():
    ws = _sheet(build_hipertensi_gap_workbook("PKM Uji", 2026, []))
    assert ws.max_row == 4  # just the title/subtitle/header block
