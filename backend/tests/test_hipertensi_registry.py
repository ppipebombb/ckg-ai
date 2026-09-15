"""End-to-end-ish tests for the Hipertensi registry: scan logic (inclusion per
the dirjen syarat, ASIK-first sourcing for identitas/TD/riwayat/obat with EPUS
same-date fallback, next-month follow-up with ALL visits per month listed,
Missed Visit fill, cross-date dedup, HT-only drug filter) → Pydantic row → xlsx
(live formulas + Formula & Logika sheet).

Uses real Fernet encryption (``encrypt_json``) over synthetic EPUS/ASIK blobs
and a fake Session so the whole pipeline runs without Postgres. A **past**
report year (2024) is used so the Missed-Visit fill (capped at the current month
only for the current year) deterministically runs through December regardless of
when the test runs.
"""

import uuid
from datetime import date
from types import SimpleNamespace

import openpyxl

from app.core.security import encrypt_json
from app.models.patient import MatchStatus
from app.schemas.hipertensi_registry import HipertensiRegistryRow
from app.services.hipertensi_registry_export import (
    _month_base_col,
    build_hipertensi_registry_workbook,
)
from app.services.hipertensi_registry_scan import (
    _FU_MISSED_VISIT,
    _FU_TERKENDALI,
    _FU_TIDAK_TERKENDALI,
    _interpretasi_baseline,
    count_interpretasi_bands,
    scan_hipertensi_registry,
)

_YEAR = 2024  # past year → Missed Visit fill runs to December (deterministic)

_RIWAYAT_Q = "Apakah Anda pernah dinyatakan tekanan darah tinggi?"


# ── blob builders ──────────────────────────────────────────────────────────
def _epus(sys, dia, *, nik, nama="PASIEN X", icd=None, obat=None, rw_ht=None):
    resep = (
        {"tables": {"Resep": [{"Nama Obat": o} for o in obat]}}
        if obat
        else {"fields": {}}
    )
    tabs = {
        "Anamnesa": {
            "fields": {
                "Periksa Fisik": {
                    "Sistole": str(sys) if sys is not None else None,
                    "Diastole": str(dia) if dia is not None else None,
                }
            }
        },
        "Resep": resep,
    }
    if rw_ht is not None:
        tabs["PTM"] = {
            "fields": {"Riwayat PTM pada Diri Sendiri": {"Penyakit Hipertensi": rw_ht}}
        }
    return {
        "data_pasien": {
            "NIK": nik,
            "Nama Pasien": nama,
            "Jenis Kelamin": "P",
            "Tempat/Tgl Lahir": "BEKASI/ 27-04-1980",
            "No Telp / HP": "08123456789",
            "Alamat": "jl pakis RT RW Kel PEKAYON",
        },
        "penyakit_khusus": ([{"ICDX": icd, "Penyakit": ""}] if icd else []),
        "tabs": tabs,
    }


def _asik(sys2=None, dia2=None, riwayat=None, *, sys1=None, dia1=None,
          obat=None, diagnosis=None, identitas=None):
    """Raw ASIK blob. ``sys1``/``dia1`` are the ASIK first TD reading (the
    source of truth); ``sys2``/``dia2`` the second; ``obat``/``diagnosis`` go in
    a recorded Tekanan Darah tatalaksana form; ``identitas`` is the
    ``detail_data`` block."""
    fd = {}
    if sys1 is not None:
        fd["Tekanan Darah Sistolik"] = str(sys1)
    if dia1 is not None:
        fd["Tekanan darah diastolik"] = str(dia1)
    if sys2 is not None:
        fd["Tekanan Darah Sistolik Ke-2"] = str(sys2)
    if dia2 is not None:
        fd["Tekanan darah diastolik ke-2"] = str(dia2)
    if riwayat is not None:
        fd[_RIWAYAT_Q] = riwayat
    out = {
        "pelayanan_nakes": [
            {"layanan": "Tekanan Darah Dewasa Lansia", "form_data": fd}
        ],
        "pemeriksaan_mandiri": [],
    }
    if identitas is not None:
        out["detail_data"] = identitas
    if obat is not None or diagnosis is not None:
        form = {}
        if diagnosis is not None:
            form["Diagnosis"] = diagnosis
        for i, o in enumerate(obat or []):
            form["Pilih obat" if i == 0 else f"Pilih obat ({i + 1})"] = o
        out["tatalaksana"] = {
            "rows": [
                {
                    "kelompok_skrinning": "Tekanan Darah Dewasa Lansia",
                    "tatalaksana_name": "Hipertensi",
                    "status": "selesai",
                    "form_data": form,
                }
            ]
        }
    return out


def _row(nik, fd, status, epus, *, asik=None, group=None, nama="PASIEN X"):
    return SimpleNamespace(
        id=uuid.uuid4(),
        nik=nik,
        nama=nama,
        filter_date=fd,
        match_status=status,
        match_group_id=group,
        scraped_epus_data=encrypt_json(epus),
        scraped_asik_data=encrypt_json(asik) if asik is not None else None,
    )


class _FakeResult:
    def __init__(self, rows):
        self._rows = rows

    def __iter__(self):
        return iter(self._rows)


class _FakeDB:
    """Returns the main blob rows on the first execute(), and the baseline
    ASIK/merged rows on the second (the scan issues exactly those two queries)."""

    def __init__(self, rows):
        self._main = sorted(rows, key=lambda r: (r.nik, r.filter_date))
        self._asik = [
            SimpleNamespace(id=r.id, scraped_asik_data=r.scraped_asik_data)
            for r in rows
        ]
        self.calls = 0

    def execute(self, _stmt):
        self.calls += 1
        return _FakeResult(self._main if self.calls == 1 else self._asik)


def _scan(rows):
    payload = scan_hipertensi_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    return {p["nik"]: p for p in payload["patients"]}


# ── tests ──────────────────────────────────────────────────────────────────
def test_full_patient_with_followup_and_missed_visits():
    M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY
    rows = [
        _row("A", date(_YEAR, 2, 10), M,
             _epus(175, 100, nik="A", icd="I10", obat=["AMLODIPINE", "CAPTOPRIL"]),
             asik=_asik(160, 95)),
        # Jan visit BEFORE the baseline date → not part of the registry window
        _row("A", date(_YEAR, 1, 5), E, _epus(130, 85, nik="A")),
        # Mar follow-up visit (EPUS only)
        _row("A", date(_YEAR, 3, 20), E, _epus(150, 92, nik="A", obat=["AMLODIPINE"])),
    ]
    p = _scan(rows)["A"]
    assert p["tanggal_berkunjung"] == f"{_YEAR}-02-10"
    assert p["riwayat_ht"] == "Ya"  # I10
    assert p["rerata_sys"] == 167.5 and p["rerata_dia"] == 97.5
    assert p["interpretasi"] == "Hipertensi"
    assert p["obat"] == ["AMLODIPINE", "CAPTOPRIL"]  # both are HT drugs
    # Follow-up starts the month AFTER the CKG month (Mar) → December.
    assert set(p["followup"].keys()) == {str(m) for m in range(3, 13)}
    assert "1" not in p["followup"] and "2" not in p["followup"]
    fu3 = p["followup"]["3"][0]
    assert fu3["tanggal"] == f"{_YEAR}-03-20"
    assert fu3["interpretasi"] == _FU_TIDAK_TERKENDALI  # 150/92
    assert fu3["obat"] == ["AMLODIPINE"]
    # A month with no visit is a Missed Visit row with no data.
    miss = p["followup"]["4"][0]
    assert miss["interpretasi"] == _FU_MISSED_VISIT
    assert miss["tanggal"] is None and miss["sys"] is None and miss["obat"] == []


def test_baseline_month_visits_excluded_from_followup():
    """The CKG month is represented by the baseline (green) block; follow-up
    evaluation is "pada kunjungan bulan berikutnya" — so even later visits in
    the same month stay out of the follow-up grid."""
    M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY
    rows = [
        _row("Z", date(_YEAR, 1, 20), M, _epus(160, 95, nik="Z", icd="I10"), asik=_asik()),
        _row("Z", date(_YEAR, 1, 10), E, _epus(150, 95, nik="Z")),  # before baseline
        _row("Z", date(_YEAR, 1, 25), E, _epus(135, 88, nik="Z")),  # same month, after
    ]
    p = _scan(rows)["Z"]
    assert "1" not in p["followup"]
    assert p["followup"]["2"][0]["interpretasi"] == _FU_MISSED_VISIT


def test_terkendali_label_next_month():
    M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY
    rows = [
        _row("T", date(_YEAR, 5, 1), M, _epus(150, 95, nik="T", icd="I10"), asik=_asik()),
        _row("T", date(_YEAR, 6, 12), E, _epus(130, 85, nik="T")),
    ]
    p = _scan(rows)["T"]
    assert p["followup"]["6"][0]["interpretasi"] == _FU_TERKENDALI  # 130/85 < 140/90


def test_all_visits_of_month_listed():
    """Multiple kunjungan in one month → ALL are listed in date order, each
    with its own independent status (client feedback: show every visit, not
    just the last one)."""
    M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY
    rows = [
        _row("L", date(_YEAR, 2, 1), M, _epus(160, 100, nik="L", icd="I10"), asik=_asik()),
        _row("L", date(_YEAR, 3, 5), E, _epus(150, 95, nik="L", obat=["CAPTOPRIL 25 MG"])),
        _row("L", date(_YEAR, 3, 20), E, _epus(132, 84, nik="L", obat=["AMLODIPIN 5 MG TAB"])),
    ]
    fu3 = _scan(rows)["L"]["followup"]["3"]
    assert len(fu3) == 2
    assert fu3[0]["tanggal"] == f"{_YEAR}-03-05"
    assert fu3[0]["interpretasi"] == _FU_TIDAK_TERKENDALI  # 150/95
    assert fu3[0]["obat"] == ["CAPTOPRIL 25 MG"]
    assert fu3[1]["tanggal"] == f"{_YEAR}-03-20"
    assert fu3[1]["interpretasi"] == _FU_TERKENDALI  # 132/84
    assert fu3[1]["obat"] == ["AMLODIPIN 5 MG TAB"]


# The next three pin the rerata ≥140/90 → "Hipertensi" rule across every combo
# of (prior diagnosis?, 2nd reading?). Neither input may resurrect a separate
# "pending confirmation" label — a high rerata is Hipertensi, full stop.
def test_no_diagnosis_high_rerata_included_riwayat_tidak():
    """Syarat 2: no diagnosis record anywhere, CKG reading ≥140/90 from a SINGLE
    measurement (no TD2), Riwayat = Tidak → INCLUDED and labelled 'Hipertensi'."""
    rows = [
        _row("B", date(_YEAR, 4, 23), MatchStatus.MATCHED,
             _epus(142, 89, nik="B"), asik=_asik()),
    ]
    out = _scan(rows)
    assert "B" in out  # included via Syarat 2
    p = out["B"]
    assert p["riwayat_ht"] == "Tidak"
    assert p["rerata_sys"] == 142 and p["rerata_dia"] == 89  # TD1 alone
    assert p["interpretasi"] == "Hipertensi"


def test_high_with_td2_no_history_is_hipertensi():
    """A high reading with a 2nd measurement (TD2 present), no prior
    diagnosis → 'Hipertensi'."""
    rows = [
        _row("HT2", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(145, 95, nik="HT2"), asik=_asik(150, 96)),
    ]
    p = _scan(rows)["HT2"]
    assert p["riwayat_ht"] == "Tidak"
    assert p["td_sys2"] == 150 and p["td_dia2"] == 96  # TD2 present
    assert p["interpretasi"] == "Hipertensi"


def test_high_single_reading_with_history_is_hipertensi():
    """A single high reading WITH a prior diagnosis (Riwayat = Ya) →
    'Hipertensi'."""
    rows = [
        _row("HH", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(160, 100, nik="HH", icd="I10"), asik=_asik()),
    ]
    p = _scan(rows)["HH"]
    assert p["riwayat_ht"] == "Ya"
    assert p["td_sys2"] is None  # single reading
    assert p["interpretasi"] == "Hipertensi"


def test_no_diagnosis_normal_excluded_pre_hipertensi_included():
    """Without a diagnosis: a Normal rerata stays out of the registry, but a
    Pre-Hipertensi rerata (130–139/85–89) is now INCLUDED for early monitoring
    (syarat widened 2026-07)."""
    rows = [
        _row("D", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(110, 70, nik="D"), asik=_asik()),  # Normal → excluded
        _row("P", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(135, 88, nik="P"), asik=_asik()),  # Pre-Hipertensi → included
    ]
    out = _scan(rows)
    assert "D" not in out  # Normal + no diagnosis stays out
    assert "P" in out  # included via widened Syarat 2
    assert out["P"]["riwayat_ht"] == "Tidak"
    assert out["P"]["interpretasi"] == "Pre-Hipertensi"


def test_high_reading_outside_baseline_does_not_include():
    """Inclusion anchors on the CKG visit: a high TD at a later non-CKG visit
    alone (no diagnosis, normal CKG rerata) does not put the NIK in the
    registry — that case belongs to the Full Review Diagnose view."""
    M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY
    rows = [
        _row("H", date(_YEAR, 3, 1), M, _epus(118, 76, nik="H"), asik=_asik()),
        _row("H", date(_YEAR, 8, 1), E, _epus(180, 110, nik="H")),
    ]
    assert "H" not in _scan(rows)


def test_epus_only_nik_excluded():
    rows = [_row("C", date(_YEAR, 5, 1), MatchStatus.EPUS_ONLY,
                 _epus(180, 110, nik="C", icd="I10"))]
    assert "C" not in _scan(rows)  # no MATCHED/CKG visit → excluded


def test_riwayat_from_epus_self_report():
    rows = [
        _row("R1", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(118, 76, nik="R1", rw_ht="Ya"), asik=_asik()),
    ]
    p = _scan(rows)["R1"]
    assert p["riwayat_ht"] == "Ya"
    assert p["interpretasi"] == "Hipertensi"  # diagnosed → Hipertensi even when controlled


def test_riwayat_from_asik_answer():
    rows = [
        _row("R2", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(118, 76, nik="R2"), asik=_asik(riwayat="Ya")),
    ]
    assert _scan(rows)["R2"]["riwayat_ht"] == "Ya"


def test_td1_from_asik_preferred_over_epus():
    """ASIK is the source of truth for the baseline reading: the ASIK first TD
    wins over the EPUS visit reading on the same date."""
    rows = [
        _row("AT", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(120, 78, nik="AT"), asik=_asik(sys1=160, dia1=95)),
    ]
    p = _scan(rows)["AT"]
    assert p["td_sys1"] == 160 and p["td_dia1"] == 95  # ASIK wins over EPUS 120/78
    # Single high reading, no prior diagnosis → Hipertensi.
    assert p["rerata_sys"] == 160 and p["interpretasi"] == "Hipertensi"


def test_td1_epus_fallback_when_asik_absent():
    """No ASIK first reading → the same-date EPUS visit reading fills TD1."""
    rows = [
        _row("EF", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(150, 95, nik="EF"), asik=_asik()),  # ASIK has no TD reading
    ]
    p = _scan(rows)["EF"]
    assert p["td_sys1"] == 150 and p["td_dia1"] == 95  # EPUS fallback
    # Single high reading, no prior diagnosis → Hipertensi.
    assert p["interpretasi"] == "Hipertensi"


def test_td2_from_raw_asik():
    rows = [
        _row("M1", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(150, 95, nik="M1", icd="I10"), asik=_asik(160, 99)),
    ]
    p = _scan(rows)["M1"]
    assert p["td_sys2"] == 160 and p["td_dia2"] == 99  # TD2 from raw ASIK
    assert p["rerata_sys"] == 155  # (150 EPUS TD1 + 160 ASIK TD2) / 2


def test_identitas_from_asik_preferred_over_epus():
    """Identitas comes from ASIK detail_data; EPUS fills only missing fields.
    The ASIK 'DD NamaBulan YYYY' birthdate is parsed to ISO."""
    ident = {
        "data_individu": {
            "Nama": "BUDI ASIK", "Jenis Kelamin": "Laki-Laki",
            "Tanggal Lahir": "25 Januari 2000", "No. HP/WA orang tua": "0811",
        },
        "data_domisili": {"Alamat Domisili": "JL ASIK NO 1"},
    }
    rows = [
        _row("ID", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(150, 95, nik="ID", nama="BUDI EPUS"),
             asik=_asik(identitas=ident)),
    ]
    p = _scan(rows)["ID"]
    assert p["nama"] == "BUDI ASIK"
    assert p["jenis_kelamin"] == "Laki-Laki"
    assert p["tanggal_lahir"] == "2000-01-25"
    assert p["alamat"] == "JL ASIK NO 1"


def test_riwayat_from_asik_tatalaksana_diagnosis():
    """A recorded ASIK tatalaksana Diagnosis with an HT ICD → Riwayat HT = Ya
    (even with a normal CKG reading and no EPUS diagnosis)."""
    rows = [
        _row("DX", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(118, 76, nik="DX"),
             asik=_asik(diagnosis="I10 - Essential (primary) hypertension")),
    ]
    p = _scan(rows)["DX"]
    assert p["riwayat_ht"] == "Ya"
    assert p["interpretasi"] == "Hipertensi"  # diagnosed → Hipertensi even when controlled


def test_obat_from_asik_tatalaksana_preferred_over_epus():
    rows = [
        _row("OB", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(150, 95, nik="OB", obat=["CAPTOPRIL 25 MG"]),
             asik=_asik(obat=["Amlodipine Besilate 5 mg Tablet"])),
    ]
    assert _scan(rows)["OB"]["obat"] == ["Amlodipine Besilate 5 mg Tablet"]


def test_obat_epus_fallback_when_no_asik_tatalaksana():
    rows = [
        _row("OE", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(150, 95, nik="OE", obat=["AMLODIPIN 5 MG TAB"]), asik=_asik()),
    ]
    assert _scan(rows)["OE"]["obat"] == ["AMLODIPIN 5 MG TAB"]


def test_sources_asik_wins():
    """Provenance: ASIK TD1, TD2, and obat win; riwayat comes from the EPUS
    diagnosis (no ASIK riwayat answer / tatalaksana diagnosis)."""
    rows = [
        _row("SP", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(120, 78, nik="SP", icd="I10", obat=["AMLODIPIN 5 MG TAB"]),
             asik=_asik(160, 99, sys1=158, dia1=96,
                        obat=["Amlodipine Besilate 5 mg Tablet"])),
    ]
    src = _scan(rows)["SP"]["sources"]
    assert src["td1"] == "ASIK"        # ASIK first reading won over EPUS
    assert src["td2"] == "ASIK"        # TD2 only ever from ASIK
    assert src["obat"] == "ASIK"       # ASIK tatalaksana prescription won
    assert src["riwayat_ht"] == "EPUS"  # EPUS I10; no ASIK riwayat signal


def test_sources_epus_fallback_and_nulls():
    """Provenance: TD1/obat fall back to ePuskesmas when ASIK has none; TD2 is
    null (no 2nd reading); riwayat from the ASIK answer."""
    rows = [
        _row("SE", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(150, 95, nik="SE", obat=["AMLODIPIN 5 MG TAB"]),
             asik=_asik(riwayat="Ya")),  # ASIK answer Ya, no TD / no obat
    ]
    src = _scan(rows)["SE"]["sources"]
    assert src["td1"] == "EPUS"         # EPUS same-date fallback
    assert src["td2"] is None           # no 2nd reading anywhere
    assert src["obat"] == "EPUS"        # EPUS Resep fallback
    assert src["riwayat_ht"] == "ASIK"  # ASIK self-report answer won


def test_td2_pushes_rerata_over_threshold():
    """TD1 alone is below 140/90 but the ASIK second reading lifts the rerata
    over the bar → included via Syarat 2."""
    rows = [
        _row("M2", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(138, 88, nik="M2"), asik=_asik(145, 95)),
    ]
    p = _scan(rows)["M2"]
    assert p["rerata_sys"] == 141.5 and p["rerata_dia"] == 91.5
    assert p["interpretasi"] == "Hipertensi"
    assert p["riwayat_ht"] == "Tidak"


def test_obat_filtered_to_antihypertensives():
    rows = [
        _row("O1", date(_YEAR, 5, 1), MatchStatus.MATCHED,
             _epus(150, 95, nik="O1", icd="I10",
                   obat=["PARACETAMOL TAB 500 MG", "AMLODIPIN 5 MG TAB",
                         "VITAMIN B COMPLEX", "Kaptopril 12,5 mg-JKN",
                         "HIDROKLORTIAZID 25 MG (HCT) TAB"]),
             asik=_asik()),
    ]
    assert _scan(rows)["O1"]["obat"] == [
        "AMLODIPIN 5 MG TAB", "Kaptopril 12,5 mg-JKN",
        "HIDROKLORTIAZID 25 MG (HCT) TAB",
    ]


def test_pre_hipertensi_band_per_ckg_juknis():
    """Rerata bands per KMK 84/2026, which apply only when Riwayat HT = Tidak:
    130–139/85–89 = Pre-Hipertensi; 120–129/80–84 is Normal (the legacy JNC7
    120/80 cut-off no longer applies). Tested on the pure classifier; the
    scan-level inclusion of a Pre-Hipertensi row (and exclusion of Normal) is
    pinned by test_no_diagnosis_normal_excluded_pre_hipertensi_included. A
    Riwayat = Ya patient is always Hipertensi (see test_riwayat_* tests)."""
    assert _interpretasi_baseline(132, 86, riwayat_ya=False) == "Pre-Hipertensi"
    assert _interpretasi_baseline(125, 82, riwayat_ya=False) == "Normal"
    assert _interpretasi_baseline(120, 80, riwayat_ya=True) == "Hipertensi"


def test_count_interpretasi_bands():
    """(hipertensi, pre_hipertensi) split over registry patients; the two sum to
    the registry total (every listed patient is one band or the other)."""
    patients = [
        {"interpretasi": "Hipertensi"},
        {"interpretasi": "Hipertensi"},
        {"interpretasi": "Pre-Hipertensi"},
    ]
    assert count_interpretasi_bands(patients) == (2, 1)
    assert count_interpretasi_bands([]) == (0, 0)


def test_cross_date_twin_deduped_in_followup():
    gid = uuid.uuid4()
    M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY
    rows = [
        _row("E", date(_YEAR, 2, 10), M, _epus(160, 95, nik="E", icd="I10"), asik=_asik()),
        # One real March visit scraped on two filter_dates → counts once.
        _row("E", date(_YEAR, 3, 10), E, _epus(160, 95, nik="E", icd="I10"), group=gid),
        _row("E", date(_YEAR, 3, 14), E, _epus(160, 95, nik="E", icd="I10"), group=gid),
    ]
    fu3 = _scan(rows)["E"]["followup"]["3"]
    assert len(fu3) == 1  # twins are one visit, listed once
    assert fu3[0]["tanggal"] == f"{_YEAR}-03-10"  # earliest date is the real one


def test_same_date_rescrape_deduped_keeping_freshest():
    """Two rows with the SAME filter_date (a re-scrape duplicate of one real
    visit, no match_group_id) collapse to one kunjungan; the later-created row
    (last in the (filter_date, created_at) order) wins."""
    M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY
    rows = [
        _row("S", date(_YEAR, 2, 1), M, _epus(160, 100, nik="S", icd="I10"), asik=_asik()),
        _row("S", date(_YEAR, 3, 7), E, _epus(150, 95, nik="S")),  # stale scrape
        _row("S", date(_YEAR, 3, 7), E, _epus(150, 95, nik="S", obat=["AMLODIPIN 5 MG TAB"])),
    ]
    fu3 = _scan(rows)["S"]["followup"]["3"]
    assert len(fu3) == 1  # one kunjungan, not a doubled row
    assert fu3[0]["tanggal"] == f"{_YEAR}-03-07"
    assert fu3[0]["obat"] == ["AMLODIPIN 5 MG TAB"]  # freshest scrape wins


def test_december_baseline_has_no_followup_months():
    rows = [
        _row("X", date(_YEAR, 12, 5), MatchStatus.MATCHED,
             _epus(150, 95, nik="X", icd="I10"), asik=_asik()),
    ]
    assert _scan(rows)["X"]["followup"] == {}


def test_payload_to_pydantic_and_xlsx():
    rows = [
        _row("A", date(_YEAR, 2, 10), MatchStatus.MATCHED,
             _epus(175, 100, nik="A", icd="I10", obat=["AMLODIPINE", "CAPTOPRIL"]),
             asik=_asik(160, 95)),
        # Two March visits → both stack in the March column-group.
        _row("A", date(_YEAR, 3, 20), MatchStatus.EPUS_ONLY,
             _epus(150, 92, nik="A", obat=["AMLODIPINE"])),
        _row("A", date(_YEAR, 3, 25), MatchStatus.EPUS_ONLY,
             _epus(135, 85, nik="A")),
        _row("B", date(_YEAR, 4, 23), MatchStatus.MATCHED,
             _epus(142, 89, nik="B"), asik=_asik()),
    ]
    patients = list(_scan(rows).values())
    # Pydantic coercion (str month keys → int, ISO → date, null tanggal ok).
    rowmodels = [
        HipertensiRegistryRow(puskesmas_name="PK", tahun_pelaporan=_YEAR, **p)
        for p in patients
    ]
    assert all(isinstance(rm.tanggal_berkunjung, date) for rm in rowmodels)

    content = build_hipertensi_registry_workbook("Puskesmas Test", _YEAR, patients)
    import io

    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    assert ws.cell(1, 1).value == "KERTAS KERJA REGISTRI HIPERTENSI"
    assert ws.cell(8, 1).value == "A"  # NIK on the anchor row
    assert ws.cell(8, 9).value == 175  # TD Sistolik 1
    # Derived cells are live formulas (M=rerata sys, O=interpretasi).
    assert str(ws.cell(8, 13).value).startswith("=IF(COUNT(I8,K8)")
    interp_formula = str(ws.cell(8, 15).value)  # formula text
    # Riwayat HT = Ya (col H) short-circuits to Hipertensi; otherwise the three
    # rerata bands (M,N). The formula must NOT consult TD2 (K,L) — no confirm gate.
    assert '"Hipertensi"' in interp_formula
    assert '"Pre-Hipertensi"' in interp_formula
    assert '"Normal"' in interp_formula
    assert "Meningkat" not in interp_formula
    assert 'H8="Ya"' in interp_formula  # diagnosis short-circuit
    assert "K8" not in interp_formula and "L8" not in interp_formula  # TD2 not consulted
    # March follow-up: BOTH visits stack in the month column-group, each with
    # its own date, TD, and status formula.
    bc = _month_base_col(3)
    assert ws.cell(8, bc).value.date() == date(_YEAR, 3, 20)
    assert ws.cell(9, bc).value.date() == date(_YEAR, 3, 25)
    assert ws.cell(8, bc + 1).value == 150 and ws.cell(9, bc + 1).value == 135
    assert str(ws.cell(8, bc + 3).value).startswith("=IF(COUNT(")
    assert str(ws.cell(9, bc + 3).value).startswith("=IF(COUNT(")
    # The documentation sheet rides along.
    assert "Formula & Logika" in wb.sheetnames


def test_month_column_mapping():
    from openpyxl.utils import get_column_letter

    assert get_column_letter(_month_base_col(1)) == "Q"
    assert get_column_letter(_month_base_col(12)) == "BT"
    assert get_column_letter(_month_base_col(12) + 4) == "BX"
