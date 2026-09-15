"""Tests for the Dislipidemia registry scan.

Mirrors ``test_dm_registry.py`` (real Fernet encryption over synthetic EPUS/ASIK
blobs plus a fake Session, so the pipeline runs without Postgres), and adds
dedicated coverage for the three ways this registry deliberately DIVERGES from
its DM/hipertensi siblings:

  1. syarat is measurement-only — a recorded HT/DM diagnosis does NOT admit a
     patient with a clean panel,
  2. riwayat does NOT pin the baseline label,
  3. control is QUARTERLY — ``Pasien Missed Visit`` appears only on control
     months, and the months in between stay blank.

Plus the two label-collision traps that would silently corrupt values: the CKG
riwayat QUESTION contains the word "kolesterol", and lab rows are named
"Kolesterol HDL"/"Kolesterol LDL".

A **past** report year (2024) is used so the quarterly control fill (capped at
the current month only for the current year) runs deterministically regardless
of when the test runs.
"""

import io
import json
import uuid
from datetime import date
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.api.deps import Principal, get_dashboard_principal, get_db
from app.core.security import encrypt_json
from app.main import app
from app.models.patient import MatchStatus
from app.services.lipid_registry_export import (
    _LAST_COL,
    _MONTH_STRIDE,
    _OBAT_COL,
    _month_base_col,
    build_lipid_registry_workbook,
)
from app.services.lipid_registry_scan import (
    _FU_MISSED_VISIT,
    _FU_TERKENDALI,
    _FU_TIDAK_TERKENDALI,
    _INTERP_DISLIPIDEMIA,
    _INTERP_NORMAL,
    _abnormal_analytes,
    _control_months,
    _interpretasi_baseline,
    _interpretasi_followup,
    count_analyte_bands,
    scan_lipid_registry,
)

_YEAR = 2024  # past year → quarterly control fill runs to December

_RIWAYAT_LIPID_Q = (
    "Apakah Anda pernah didiagnosa atau mendapatkan hasil pemeriksaan "
    "kolesterol (lemak darah) tinggi?"
)
_RIWAYAT_HT_Q = (
    "Apakah Anda pernah dinyatakan tekanan darah tinggi/hipertensi oleh Dokter?"
)
_RIWAYAT_DM_Q = (
    "Apakah Anda pernah dinyatakan diabetes atau kencing manis oleh Dokter?"
)
_LIPID_LAYANAN = (
    "Skrining Laboratorium =>40 thn Perempuan Fungsi Ginjal, Hati, Profil Lipid"
)


# ── blob builders ──────────────────────────────────────────────────────────
def _epus(
    *, nik, nama="PASIEN X", kol=None, ldl=None, hdl=None, trig=None,
    icd=None, obat=None, lab=None,
):
    """Synthetic EPUS visit blob. ``lab`` is a list of (Pemeriksaan, Hasil).
    ``icd`` goes into ``penyakit_khusus`` — I10 marks hipertensi, E11 diabetes."""
    lipid = {}
    if kol is not None:
        lipid["Cholesterol Total"] = str(kol)
    if hdl is not None:
        lipid["HDL"] = str(hdl)
    if ldl is not None:
        lipid["LDL"] = str(ldl)
    if trig is not None:
        lipid["Trigliserida"] = str(trig)

    tabs = {"Resep": (
        {"tables": {"Resep": [{"Nama Obat": o} for o in obat]}}
        if obat else {"fields": {}}
    )}
    if lipid:
        tabs["PTM"] = {"fields": {"Profil Lipid": lipid}}
    if lab:
        tabs["Laboratorium"] = {
            "tables": {
                "Ubah Data Laboratorium": [
                    {"Pemeriksaan": p, "Hasil": str(h)} for p, h in lab
                ]
            }
        }
    icds = icd if isinstance(icd, (list, tuple)) else ([icd] if icd else [])
    return {
        "data_pasien": {
            "NIK": nik,
            "Nama Pasien": nama,
            "Jenis Kelamin": "P",
            "Tempat/Tgl Lahir": "BEKASI/ 27-04-1980",
            "No Telp / HP": "08123456789",
            "Alamat": "jl pakis RT RW Kel PEKAYON",
        },
        "penyakit_khusus": [{"ICDX": c, "Penyakit": ""} for c in icds],
        "tabs": tabs,
    }


def _asik(
    *, kol=None, ldl=None, hdl=None, trig=None, riwayat_lipid=None,
    riwayat_ht=None, riwayat_dm=None, obat=None, identitas=None,
    interpretasi=None,
):
    """Raw ASIK blob. The lipid values live on the POCT Lipid Panel form
    (reached via a "… Profil Lipid" layanan); the HT/DM riwayat answers live on
    their own layanan, exactly as the sibling registries read them."""
    form = {}
    if riwayat_lipid is not None:
        form[_RIWAYAT_LIPID_Q] = riwayat_lipid
    if kol is not None:
        form["Kolesterol Total"] = str(kol)
    if hdl is not None:
        form["Masukkan nilai HDL"] = str(hdl)
    if ldl is not None:
        form["LDL"] = str(ldl)
    if trig is not None:
        form["Trigliserida"] = str(trig)
    if interpretasi is not None:
        form["Interpretasi Dislipidemia"] = interpretasi

    nakes = [{"layanan": _LIPID_LAYANAN, "form_data": form}]
    if riwayat_ht is not None:
        nakes.append({
            "layanan": "Tekanan Darah Dewasa Lansia",
            "form_data": {_RIWAYAT_HT_Q: riwayat_ht},
        })
    if riwayat_dm is not None:
        nakes.append({
            "layanan": "Pemeriksaan Gula Darah Dewasa Lansia",
            "form_data": {_RIWAYAT_DM_Q: riwayat_dm},
        })

    out = {"pelayanan_nakes": nakes, "pemeriksaan_mandiri": []}
    if identitas is not None:
        out["detail_data"] = identitas
    if obat is not None:
        tf = {}
        for i, o in enumerate(obat):
            tf["Pilih obat" if i == 0 else f"Pilih obat ({i + 1})"] = o
        out["tatalaksana"] = {
            "rows": [{
                "kelompok_skrinning": "POCT Lipid Panel",
                "tatalaksana_name": "Dislipidemia",
                "status": "selesai",
                "form_data": tf,
            }]
        }
    return out


def _row(nik, fd, status, epus, *, asik=None, group=None, nama="PASIEN X"):
    return SimpleNamespace(
        id=uuid.uuid4(),
        nik=nik,
        nama=nama,
        filter_date=fd,
        match_status=status,
        match_group_id=group,
        scraped_epus_data=encrypt_json(epus),
        scraped_asik_data=encrypt_json(asik) if asik is not None else None,
    )


class _FakeResult:
    def __init__(self, rows):
        self._rows = rows

    def __iter__(self):
        return iter(self._rows)


class _FakeDB:
    """Returns the main blob rows on the first execute(), and the baseline ASIK
    rows on the second (the scan issues exactly those two queries)."""

    def __init__(self, rows):
        self._main = sorted(rows, key=lambda r: (r.nik, r.filter_date))
        self._asik = [
            SimpleNamespace(id=r.id, scraped_asik_data=r.scraped_asik_data)
            for r in rows
        ]
        self.calls = 0

    def execute(self, _stmt):
        self.calls += 1
        return _FakeResult(self._main if self.calls == 1 else self._asik)


def _scan(rows):
    payload = scan_lipid_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    return {p["nik"]: p for p in payload["patients"]}


M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY


def _g(**kw):
    base = {"kol_total": None, "ldl": None, "hdl": None, "trigliserida": None}
    base.update(kw)
    return base


# ── pure classification: the four thresholds, at their boundaries ──────────
@pytest.mark.parametrize(
    "reading,expected",
    [
        # Kolesterol Total: >= 200 (inclusive)
        (_g(kol_total=199), []),
        (_g(kol_total=200), ["kol_total"]),
        # LDL: >= 130 (inclusive)
        (_g(ldl=129), []),
        (_g(ldl=130), ["ldl"]),
        # HDL: < 40 (the inverted one — LOW is the problem)
        (_g(hdl=40), []),
        (_g(hdl=39.9), ["hdl"]),
        # Trigliserida: > 150 — the sheet's ONLY strict comparator
        (_g(trigliserida=150), []),
        (_g(trigliserida=150.1), ["trigliserida"]),
        # order is sheet order, and several can breach at once
        (_g(kol_total=250, ldl=200, hdl=30, trigliserida=400),
         ["kol_total", "ldl", "hdl", "trigliserida"]),
        (_g(kol_total=180, hdl=30), ["hdl"]),
    ],
)
def test_thresholds_at_their_exact_boundaries(reading, expected):
    assert _abnormal_analytes(reading) == expected


def test_trigliserida_150_is_normal_not_abnormal():
    """Regression guard for the one comparator that differs. The sheet says
    'Trigliserida >150', so exactly 150 is NORMAL. Tidying this into >= (to
    match the other three) silently admits a whole band of healthy patients."""
    assert _abnormal_analytes(_g(trigliserida=150)) == []
    assert _interpretasi_baseline(_g(trigliserida=150)) == _INTERP_NORMAL


def test_baseline_interpretasi_labels():
    assert _interpretasi_baseline(_g()) == ""  # nothing measured
    assert _interpretasi_baseline(_g(kol_total=180)) == _INTERP_NORMAL
    assert _interpretasi_baseline(_g(kol_total=180, ldl=200)) == _INTERP_DISLIPIDEMIA


def test_baseline_empty_is_blank_not_normal():
    """An unmeasured panel must never read 'Normal' — that would tell a
    puskesmas their patient's lipids are fine when nobody looked."""
    assert _interpretasi_baseline(_g()) == ""


def test_followup_any_failing_analyte_wins():
    """The N17 legend words BOTH tercapai and tidak-tercapai as an OR of four,
    which overlaps. Resolved as: one failing analyte beats three passing ones."""
    assert _interpretasi_followup(_g()) == ""
    assert _interpretasi_followup(_g(kol_total=180)) == _FU_TERKENDALI
    assert (
        _interpretasi_followup(_g(kol_total=180, ldl=100, hdl=55, trigliserida=200))
        == _FU_TIDAK_TERKENDALI
    )


# ── quarterly control months ───────────────────────────────────────────────
def test_control_months_are_quarterly_not_monthly():
    """The single biggest divergence from the DM registry."""
    today = date(2026, 7, 21)
    assert _control_months(1, 2024, today) == [4, 7, 10]
    assert _control_months(2, 2024, today) == [5, 8, 11]
    assert _control_months(3, 2024, today) == [6, 9, 12]
    assert _control_months(10, 2024, today) == []  # +3 would be next year


def test_control_months_capped_at_current_month_in_current_year():
    """A control that has not come due yet is not a missed one."""
    today = date(2026, 7, 21)
    assert _control_months(1, 2026, today) == [4, 7]
    assert _control_months(6, 2026, today) == []  # Sept is still in the future


# ── syarat: measurement-only ───────────────────────────────────────────────
def test_abnormal_panel_admits_patient():
    rows = [_row("A", date(_YEAR, 3, 1), M, _epus(nik="A"), asik=_asik(ldl=160))]
    p = _scan(rows)["A"]
    assert p["interpretasi"] == _INTERP_DISLIPIDEMIA
    assert p["ldl"] == 160


def test_normal_panel_is_excluded():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"),
             asik=_asik(kol=180, ldl=100, hdl=55, trig=120)),
    ]
    assert _scan(rows) == {}


def test_unmeasured_panel_is_excluded():
    rows = [_row("A", date(_YEAR, 3, 1), M, _epus(nik="A"), asik=_asik())]
    assert _scan(rows) == {}


def test_riwayat_ht_or_dm_alone_does_NOT_admit():
    """DIVERGENCE 1. In the hipertensi and DM registries a recorded diagnosis is
    itself a ticket in. Here it is not: the sheet admits on measurement
    ("ditegakkan diagnosis Dislipidemia melalui pengukuran"), and Riwayat HT/DM
    are columns H and I — reported, never filtered on. Gating on them would drop
    an LDL-190 patient who has neither diagnosis."""
    rows = [
        # both diagnoses recorded, panel clean → still excluded
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", icd=["I10", "E11"]),
             asik=_asik(kol=180, ldl=100, hdl=55, trig=120,
                        riwayat_ht="Ya", riwayat_dm="Ya")),
        # no diagnosis at all, panel abnormal → admitted
        _row("B", date(_YEAR, 3, 1), M, _epus(nik="B"), asik=_asik(ldl=190)),
    ]
    out = _scan(rows)
    assert set(out) == {"B"}


def test_riwayat_does_not_pin_the_baseline_label():
    """DIVERGENCE 2. A diagnosed patient whose panel is abnormal is labelled from
    the panel, and the riwayat columns ride along as information only."""
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", icd=["I10", "E11"]),
             asik=_asik(hdl=30)),
    ]
    p = _scan(rows)["A"]
    assert p["riwayat_ht"] == "Ya" and p["riwayat_dm"] == "Ya"
    assert p["interpretasi"] == _INTERP_DISLIPIDEMIA  # from HDL 30, not from riwayat


def test_no_matched_visit_means_no_row():
    rows = [_row("A", date(_YEAR, 3, 1), E, _epus(nik="A", ldl=200))]
    assert _scan(rows) == {}


# ── riwayat HT / DM sourcing ───────────────────────────────────────────────
def test_riwayat_from_asik_answers():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"),
             asik=_asik(ldl=160, riwayat_ht="Ya", riwayat_dm="Tidak")),
    ]
    p = _scan(rows)["A"]
    assert p["riwayat_ht"] == "Ya"
    assert p["riwayat_dm"] == "Tidak"
    assert p["sources"]["riwayat_ht"] == "ASIK"
    assert p["sources"]["riwayat_dm"] is None


def test_riwayat_falls_back_to_epus_diagnosis():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", icd="E11"),
             asik=_asik(ldl=160)),
    ]
    p = _scan(rows)["A"]
    assert p["riwayat_dm"] == "Ya"
    assert p["sources"]["riwayat_dm"] == "EPUS"
    assert p["riwayat_ht"] == "Tidak"


def test_riwayat_is_sticky_across_visits():
    """A diagnosis recorded on ANY in-year visit counts, not just the baseline."""
    rows = [
        _row("A", date(_YEAR, 2, 1), M, _epus(nik="A"), asik=_asik(ldl=160)),
        _row("A", date(_YEAR, 5, 1), E, _epus(nik="A", icd="I10")),
    ]
    assert _scan(rows)["A"]["riwayat_ht"] == "Ya"


# ── value sourcing: ASIK first, EPUS per-field fallback ────────────────────
def test_asik_wins_per_field_and_epus_fills_the_gaps():
    rows = [
        _row("A", date(_YEAR, 3, 1), M,
             _epus(nik="A", kol=999, trig=180),
             asik=_asik(kol=230, ldl=160)),
    ]
    p = _scan(rows)["A"]
    assert p["kol_total"] == 230       # ASIK wins
    assert p["ldl"] == 160             # ASIK only
    assert p["trigliserida"] == 180    # EPUS fills the gap
    assert p["hdl"] is None
    assert p["sources"]["lipid"] == "ASIK"


def test_epus_only_marks_the_source_as_epus():
    rows = [_row("A", date(_YEAR, 3, 1), M, _epus(nik="A", ldl=160), asik=_asik())]
    p = _scan(rows)["A"]
    assert p["ldl"] == 160
    assert p["sources"]["lipid"] == "EPUS"


def test_asik_riwayat_question_is_not_parsed_as_a_cholesterol_value():
    """TRAP. The CKG question 'Apakah Anda pernah didiagnosa atau mendapatkan
    hasil pemeriksaan KOLESTEROL (lemak darah) tinggi?' contains the same word as
    the value label. Without the question guard the answer string lands in
    Kolesterol Total (and a non-numeric answer would blank a real reading)."""
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"),
             asik=_asik(riwayat_lipid="Ya", ldl=160)),
    ]
    p = _scan(rows)["A"]
    assert p["kol_total"] is None
    assert p["ldl"] == 160


def test_asik_interpretasi_field_is_not_parsed_as_a_value():
    """ASIK ships its own 'Interpretasi Dislipidemia' verdict string. We compute
    our own from the four numbers so the label always matches the thresholds the
    export's live formulas apply."""
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"),
             asik=_asik(ldl=160, interpretasi="Dislipidemia")),
    ]
    p = _scan(rows)["A"]
    assert p["interpretasi"] == _INTERP_DISLIPIDEMIA
    assert p["kol_total"] is None  # the verdict string went nowhere


# ── EPUS lab table gap-fill ────────────────────────────────────────────────
def test_lab_table_fills_gaps_but_ptm_stays_source_of_truth():
    rows = [
        _row("A", date(_YEAR, 3, 1), M,
             _epus(nik="A", kol=210, lab=[("Kolesterol Total", 999),
                                          ("Trigliserida", 300)]),
             asik=_asik()),
    ]
    p = _scan(rows)["A"]
    assert p["kol_total"] == 210        # PTM wins
    assert p["trigliserida"] == 300     # lab fills the gap


def test_lab_table_last_matching_row_wins():
    """Collect-then-apply, per-array latest — a re-test on the same visit
    supersedes the earlier result. This is the rule commit 5592dc9 had to
    restore in the DM scan after an in-loop gap-fill froze the FIRST row."""
    rows = [
        _row("A", date(_YEAR, 3, 1), M,
             _epus(nik="A", lab=[("Kolesterol LDL", 140), ("Kolesterol LDL", 155)]),
             asik=_asik()),
    ]
    assert _scan(rows)["A"]["ldl"] == 155


def test_lab_row_named_kolesterol_hdl_is_hdl_not_total():
    """TRAP. Lab rows are routinely named 'Kolesterol HDL' / 'Kolesterol LDL'.
    Testing 'kolesterol' before 'hdl'/'ldl' files every HDL result as total
    cholesterol — which then reads as normal (35 < 200) and drops the patient."""
    rows = [
        _row("A", date(_YEAR, 3, 1), M,
             _epus(nik="A", lab=[("Kolesterol HDL", 35)]), asik=_asik()),
    ]
    p = _scan(rows)["A"]
    assert p["hdl"] == 35
    assert p["kol_total"] is None
    assert p["interpretasi"] == _INTERP_DISLIPIDEMIA


# ── follow-up ──────────────────────────────────────────────────────────────
def test_followup_missed_visit_only_on_control_months():
    """DIVERGENCE 3. Baseline in January → controls in April, July, October.
    February/March/May/... must stay ABSENT, not 'Pasien Missed Visit'."""
    rows = [_row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(ldl=160))]
    fu = _scan(rows)["A"]["followup"]
    assert sorted(int(k) for k in fu) == [4, 7, 10]
    for m in ("4", "7", "10"):
        assert fu[m][0]["interpretasi"] == _FU_MISSED_VISIT
        assert fu[m][0]["tanggal"] is None


def test_followup_records_visits_outside_control_months_too():
    """A lipid panel taken in a non-control month is still listed and still
    evaluated — the quarterly rule governs MISSED visits, not recorded ones."""
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(ldl=160)),
        _row("A", date(_YEAR, 2, 15), E, _epus(nik="A", ldl=100, kol=180)),
    ]
    fu = _scan(rows)["A"]["followup"]
    assert "2" in fu
    assert fu["2"][0]["tanggal"] == f"{_YEAR}-02-15"
    assert fu["2"][0]["interpretasi"] == _FU_TERKENDALI
    # …and February does NOT become a control month
    assert sorted(int(k) for k in fu) == [2, 4, 7, 10]


def test_followup_visit_in_a_control_month_replaces_the_missed_filler():
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(ldl=160)),
        _row("A", date(_YEAR, 4, 20), E, _epus(nik="A", ldl=200)),
    ]
    fu = _scan(rows)["A"]["followup"]
    assert len(fu["4"]) == 1
    assert fu["4"][0]["tanggal"] == f"{_YEAR}-04-20"
    assert fu["4"][0]["interpretasi"] == _FU_TIDAK_TERKENDALI


def test_followup_skips_the_baseline_month():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"), asik=_asik(ldl=160)),
        _row("A", date(_YEAR, 3, 25), E, _epus(nik="A", ldl=140)),
    ]
    assert "3" not in _scan(rows)["A"]["followup"]


def test_followup_lists_every_visit_in_a_month():
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(ldl=160)),
        _row("A", date(_YEAR, 4, 5), E, _epus(nik="A", ldl=200)),
        _row("A", date(_YEAR, 4, 25), E, _epus(nik="A", ldl=100, kol=180)),
    ]
    fu4 = _scan(rows)["A"]["followup"]["4"]
    assert [v["tanggal"] for v in fu4] == [f"{_YEAR}-04-05", f"{_YEAR}-04-25"]
    assert [v["interpretasi"] for v in fu4] == [_FU_TIDAK_TERKENDALI, _FU_TERKENDALI]


def test_followup_visit_without_lipid_values_is_not_listed():
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(ldl=160)),
        _row("A", date(_YEAR, 4, 5), E, _epus(nik="A")),  # visit, no lipid panel
    ]
    fu4 = _scan(rows)["A"]["followup"]["4"]
    assert len(fu4) == 1
    assert fu4[0]["interpretasi"] == _FU_MISSED_VISIT


def test_cross_date_scrape_twins_collapse():
    gid = uuid.uuid4()
    rows = [
        _row("C", date(_YEAR, 1, 5), M, _epus(nik="C"), asik=_asik(ldl=160)),
        _row("C", date(_YEAR, 4, 5), E, _epus(nik="C", ldl=150), group=gid),
        _row("C", date(_YEAR, 4, 9), E, _epus(nik="C", ldl=150), group=gid),
    ]
    fu4 = _scan(rows)["C"]["followup"]["4"]
    assert len(fu4) == 1
    assert fu4[0]["tanggal"] == f"{_YEAR}-04-05"


# ── drugs ──────────────────────────────────────────────────────────────────
def test_only_lipid_drugs_are_kept():
    rows = [
        _row("A", date(_YEAR, 3, 1), M,
             _epus(nik="A", obat=["SIMVASTATIN 20 MG", "AMLODIPIN 10 MG",
                                  "METFORMIN 500 MG", "GEMFIBROZIL 300 MG"]),
             asik=_asik(ldl=160)),
    ]
    p = _scan(rows)["A"]
    assert p["obat"] == ["SIMVASTATIN 20 MG", "GEMFIBROZIL 300 MG"]
    assert p["sources"]["obat"] == "EPUS"


def test_asik_prescription_wins_over_epus_resep():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", obat=["SIMVASTATIN 20 MG"]),
             asik=_asik(ldl=160, obat=["Atorvastatin 20 mg", "Vitamin C"])),
    ]
    p = _scan(rows)["A"]
    assert p["obat"] == ["Atorvastatin 20 mg"]  # Vitamin C filtered out
    assert p["sources"]["obat"] == "ASIK"


# ── identitas + ordering ───────────────────────────────────────────────────
def test_identitas_asik_wins_epus_fills_gaps():
    rows = [
        _row("I", date(_YEAR, 3, 1), M, _epus(nik="I", nama="EPUS NAMA"),
             asik=_asik(ldl=160, identitas={
                 "data_individu": {
                     "Nama": "ASIK NAMA",
                     "Jenis Kelamin": "Perempuan",
                     "Tanggal Lahir": "25 Januari 2000",
                 },
                 "data_domisili": {},
             })),
    ]
    p = _scan(rows)["I"]
    assert p["nama"] == "ASIK NAMA"
    assert p["tanggal_lahir"] == "2000-01-25"
    assert p["no_tlp"] == "08123456789"  # ASIK lacks it → EPUS fills
    assert p["alamat"].startswith("jl pakis")


def test_tanggal_berkunjung_is_the_earliest_matched_visit():
    rows = [
        _row("A", date(_YEAR, 5, 1), M, _epus(nik="A"), asik=_asik(ldl=160)),
        _row("A", date(_YEAR, 2, 1), M, _epus(nik="A"), asik=_asik(ldl=160)),
    ]
    assert _scan(rows)["A"]["tanggal_berkunjung"] == f"{_YEAR}-02-01"


def test_rows_sorted_by_nama_then_nik():
    rows = [
        _row("2", date(_YEAR, 3, 1), M, _epus(nik="2", nama="ZULFA"),
             asik=_asik(ldl=160)),
        _row("1", date(_YEAR, 3, 1), M, _epus(nik="1", nama="ANDI"),
             asik=_asik(ldl=160)),
    ]
    payload = scan_lipid_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    assert [p["nama"] for p in payload["patients"]] == ["ANDI", "ZULFA"]


# ── band counts ────────────────────────────────────────────────────────────
def test_analyte_bands_overlap_and_do_not_sum_to_total():
    """The counterpart of the DM registry's 'bands sum to total' invariant —
    here the opposite is guaranteed, and the UI must never sum them."""
    rows = [
        _row("X1", date(_YEAR, 3, 1), M, _epus(nik="X1"),
             asik=_asik(kol=250, hdl=30)),          # breaches TWO
        _row("X2", date(_YEAR, 3, 1), M, _epus(nik="X2"), asik=_asik(ldl=140)),
        _row("X3", date(_YEAR, 3, 1), M, _epus(nik="X3"), asik=_asik(trig=300)),
    ]
    payload = scan_lipid_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    patients = payload["patients"]
    bands = count_analyte_bands(patients)
    assert bands == {
        "kol_total_tinggi": 1, "ldl_tinggi": 1,
        "hdl_rendah": 1, "trigliserida_tinggi": 1,
    }
    assert sum(bands.values()) == 4
    assert len(patients) == 3  # …deliberately NOT equal to the sum


# ── routes ─────────────────────────────────────────────────────────────────
def _sample_payload():
    """A realistic scan payload produced by the real scan, so the route tests
    exercise the actual serialization contract."""
    rows = [
        _row("111", date(_YEAR, 1, 10), M,
             _epus(nik="111", nama="BUDI", icd="I10", obat=["SIMVASTATIN 20 MG"]),
             asik=_asik(kol=230, ldl=160, hdl=45, trig=120, riwayat_ht="Ya")),
        _row("111", date(_YEAR, 4, 20), E, _epus(nik="111", ldl=200)),
        _row("222", date(_YEAR, 2, 11), M, _epus(nik="222", nama="ANI"),
             asik=_asik(hdl=30)),
    ]
    return scan_lipid_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)


@pytest.fixture
def lipid_client(monkeypatch, stub_redis, warm_calls):
    pid = uuid.uuid4()
    fake_db = MagicMock()
    fake_db.scalar.return_value = "PUSKESMAS TEBET"

    app.dependency_overrides[get_db] = lambda: fake_db
    app.dependency_overrides[get_dashboard_principal] = lambda: Principal(
        "admin", uuid.uuid4(), None
    )

    import app.api.routes.lipid_report as lipid

    monkeypatch.setattr(lipid, "redis_client", lambda: stub_redis)
    monkeypatch.setattr(
        lipid,
        "request_warm",
        lambda rc, cache_key, report_type, kwargs: warm_calls.append(
            (cache_key, report_type)
        ),
    )
    with TestClient(app) as client:
        yield client, pid, stub_redis
    app.dependency_overrides.clear()


def _seed(rc, pid, payload):
    from app.api.routes.lipid_report import _registry_cache_key

    rc.set(_registry_cache_key(pid, _YEAR), json.dumps(payload))


def test_route_cold_cache_returns_computing_and_enqueues_warm(
    lipid_client, warm_calls
):
    client, pid, _rc = lipid_client
    r = client.get("/lipid-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["computing"] is True
    assert body["items"] == [] and body["total"] == 0
    assert body["cache_hit"] is False
    # The warm must be dispatched under the report_type cron.warm_one knows.
    assert [t for _k, t in warm_calls] == ["lipid_registry"]


def test_route_serves_cached_payload(lipid_client):
    client, pid, rc = lipid_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/lipid-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["computing"] is False and body["cache_hit"] is True
    assert body["total"] == 2
    assert body["kol_total_tinggi"] == 1
    assert body["ldl_tinggi"] == 1
    assert body["hdl_rendah"] == 1
    assert body["trigliserida_tinggi"] == 0
    assert [i["nama"] for i in body["items"]] == ["ANI", "BUDI"]  # sorted by nama
    budi = next(i for i in body["items"] if i["nama"] == "BUDI")
    assert budi["puskesmas_name"] == "PUSKESMAS TEBET"
    assert budi["tahun_pelaporan"] == _YEAR
    assert budi["riwayat_ht"] == "Ya" and budi["riwayat_dm"] == "Tidak"
    assert budi["kol_total"] == 230 and budi["ldl"] == 160
    assert budi["interpretasi"] == _INTERP_DISLIPIDEMIA
    # followup dict keys survive the str->int coercion, values keep their shape
    assert budi["followup"]["4"][0]["ldl"] == 200
    assert budi["followup"]["4"][0]["interpretasi"] == _FU_TIDAK_TERKENDALI


def test_route_search_filters_and_recounts_bands(lipid_client):
    client, pid, rc = lipid_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/lipid-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR, "q": "ani"})
    body = r.json()
    assert body["total"] == 1
    assert [i["nama"] for i in body["items"]] == ["ANI"]
    # Bands respect the filter, like `total`.
    assert body["hdl_rendah"] == 1 and body["kol_total_tinggi"] == 0


def test_route_search_matches_nik(lipid_client):
    client, pid, rc = lipid_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/lipid-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR, "q": "111"})
    assert [i["nik"] for i in r.json()["items"]] == ["111"]


def test_route_summary(lipid_client):
    client, pid, rc = lipid_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/lipid-reports/registry/summary",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 2
    assert body["kol_total_tinggi"] == 1 and body["hdl_rendah"] == 1
    assert body["riwayat_ht_ya"] == 1
    assert body["riwayat_dm_ya"] == 0
    assert body["dalam_pengobatan"] == 1  # only BUDI has a drug
    assert body["computing"] is False


def test_route_summary_cold_cache_enqueues_warm(lipid_client, warm_calls):
    client, pid, _rc = lipid_client
    r = client.get("/lipid-reports/registry/summary",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.json()["computing"] is True
    assert [t for _k, t in warm_calls] == ["lipid_registry"]


def test_route_clear_cache_deletes_and_requeues(lipid_client, warm_calls):
    from app.api.routes.lipid_report import _registry_cache_key

    client, pid, rc = lipid_client
    _seed(rc, pid, _sample_payload())
    key = _registry_cache_key(pid, _YEAR)
    assert rc.get(key) is not None

    r = client.request("DELETE", "/lipid-reports/registry/cache",
                       params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 204
    assert rc.get(key) is None
    assert [t for _k, t in warm_calls] == ["lipid_registry"]


def test_route_404_when_puskesmas_missing(lipid_client):
    client, pid, _rc = lipid_client
    app.dependency_overrides[get_db] = lambda: SimpleNamespace(
        scalar=lambda *_a, **_k: None
    )
    r = client.get("/lipid-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 404


# ── xlsx export ────────────────────────────────────────────────────────────
def _wb(patients=None):
    payload = _sample_payload()
    content = build_lipid_registry_workbook(
        "PUSKESMAS TEBET", _YEAR,
        patients if patients is not None else payload["patients"],
    )
    return openpyxl.load_workbook(io.BytesIO(content))


def test_export_grid_geometry_matches_the_dirjen_layout():
    """Identitas A-F, baseline G-N (incl. BOTH riwayat columns), obat O, then
    12 months x 7 columns starting at P (16) → last column CU (99)."""
    assert _MONTH_STRIDE == 7
    assert _month_base_col(1) == 16
    assert _month_base_col(12) == 16 + 11 * 7
    assert _LAST_COL == 99


def test_export_sheets_and_headers():
    wb = _wb()
    assert wb.sheetnames == ["Registri Dislipidemia", "Formula & Logika"]
    ws = wb["Registri Dislipidemia"]
    assert ws.cell(7, 1).value == "NIK"
    assert ws.cell(7, 7).value == "Tanggal Berkunjung"
    assert ws.cell(7, 8).value == "Riwayat diagnosis HT"
    assert ws.cell(7, 9).value == "Riwayat diagnosis DM"
    # Sheet analyte order is Kol Total, LDL, HDL, Trigliserida — LDL BEFORE HDL.
    assert [ws.cell(7, c).value.split("\n")[0] for c in range(10, 14)] == [
        "Kolesterol Total", "LDL", "HDL", "Trigliserida",
    ]
    assert ws.cell(7, 14).value == "Interpretasi hasil"
    assert ws.cell(5, 15).value == "Jenis Obat"
    assert ws.cell(6, 16).value == "Januari"
    bc = _month_base_col(4)
    assert ws.cell(7, bc).value.startswith("Tanggal Pemeriksaan")
    assert [ws.cell(7, bc + i).value.split("\n")[0] for i in range(1, 5)] == [
        "Kolesterol Total", "LDL", "HDL", "Trigliserida",
    ]
    assert ws.cell(7, bc + 5).value == "Interpretasi hasil"
    assert ws.cell(7, bc + 6).value == "Jenis Obat"


def test_export_writes_patient_values():
    wb = _wb()
    ws = wb["Registri Dislipidemia"]
    # Patients are sorted by nama → ANI first, then BUDI.
    assert ws.cell(8, 1).value == "222"
    assert ws.cell(8, 2).value == "ANI"
    budi_row = next(
        r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == "BUDI"
    )
    assert ws.cell(budi_row, 8).value == "Ya"      # Riwayat HT
    assert ws.cell(budi_row, 9).value == "Tidak"   # Riwayat DM
    assert ws.cell(budi_row, 10).value == 230      # Kolesterol Total
    assert ws.cell(budi_row, 11).value == 160      # LDL
    assert ws.cell(budi_row, 15).value == "SIMVASTATIN 20 MG"
    # April follow-up: LDL 200 in the month-4 group.
    bc = _month_base_col(4)
    assert ws.cell(budi_row, bc + 2).value == 200


def test_export_non_control_months_stay_empty():
    """The workbook must agree with the dashboard: a non-control month has no
    cells at all, so it can never render as 'Pasien Missed Visit'."""
    ws = _wb()["Registri Dislipidemia"]
    budi_row = next(
        r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == "BUDI"
    )
    # BUDI's baseline is January → controls are April, July, October.
    for m in (2, 3, 5, 6, 8, 9, 11, 12):
        bc = _month_base_col(m)
        assert ws.cell(budi_row, bc + 5).value is None, f"month {m} should be blank"
    for m in (7, 10):
        bc = _month_base_col(m)
        assert str(ws.cell(budi_row, bc + 5).value).startswith("=")


def test_export_interpretasi_columns_are_live_formulas():
    wb = _wb()
    ws = wb["Registri Dislipidemia"]
    baseline = ws.cell(8, 14).value
    assert isinstance(baseline, str) and baseline.startswith("=")
    assert _INTERP_DISLIPIDEMIA in baseline and _INTERP_NORMAL in baseline
    # Riwayat (H, I) must NOT appear — a diagnosis does not pin the label here.
    assert "H8" not in baseline and "I8" not in baseline


def test_export_baseline_formula_guards_empty_hdl():
    """Regression: N("") is 0, so an unguarded `<40` test would mark every
    unmeasured patient 'Dislipidemia'. The COUNT guard must be present."""
    ws = _wb()["Registri Dislipidemia"]
    baseline = ws.cell(8, 14).value
    assert "AND(COUNT(L8)>0,N(L8)<40)" in baseline


def test_export_formula_keeps_the_strict_trigliserida_comparator():
    ws = _wb()["Registri Dislipidemia"]
    baseline = ws.cell(8, 14).value
    assert "N(M8)>150" in baseline
    assert "N(M8)>=150" not in baseline


def test_export_followup_formula_has_all_three_branches():
    ws = _wb()["Registri Dislipidemia"]
    budi_row = next(
        r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == "BUDI"
    )
    fu = ws.cell(budi_row, _month_base_col(4) + 5).value
    assert isinstance(fu, str) and fu.startswith("=")
    assert _FU_MISSED_VISIT in fu
    assert _FU_TERKENDALI in fu and _FU_TIDAK_TERKENDALI in fu


def test_export_survives_a_backslash_in_patient_text():
    """Regression: the sheetData substitution used to pass the rendered XML as a
    STRING replacement to re.sub, which parses it as a regex template. A single
    backslash in any free-text field (an ePuskesmas address like "JL X RT.01\\RW.02",
    a drug name, the puskesmas name) then raised re.PatternError and 500'd the
    download for the whole puskesmas. Every text column is exercised."""
    p = dict(_sample_payload()["patients"][0])
    p["nama"] = "BU\\DI"          # \D  -> "bad escape"
    p["alamat"] = "JL 1\\2 RT"    # \2  -> "invalid group reference"
    p["no_tlp"] = "08\\g<x>"      # \g< -> "unknown group name"
    p["obat"] = ["SIM\\VA"]
    content = build_lipid_registry_workbook("PKM A\\B", _YEAR, [p])
    ws = openpyxl.load_workbook(io.BytesIO(content))["Registri Dislipidemia"]
    assert ws.cell(8, 2).value == "BU\\DI"
    assert ws.cell(8, 6).value == "JL 1\\2 RT"
    assert ws.cell(8, _OBAT_COL).value == "SIM\\VA"
    assert "PKM A\\B" in str(ws.cell(3, 1).value)


def test_export_empty_registry_still_valid_workbook():
    wb = _wb(patients=[])
    ws = wb["Registri Dislipidemia"]
    assert ws.cell(7, 1).value == "NIK"  # headers survive
    assert ws.cell(8, 1).value is None   # no data rows


def test_export_formula_sheet_documents_the_divergences():
    """The three rules a reader would otherwise carry over wrong from the DM
    kertas kerja must be written down for the puskesmas staff reading the file."""
    ws = _wb()["Formula & Logika"]
    text = "\n".join(
        str(ws.cell(r, c).value or "")
        for r in range(1, ws.max_row + 1)
        for c in (1, 2)
    )
    assert "3 bulan" in text                       # quarterly control
    assert "bukan syarat" in text.casefold()       # riwayat is informational
    assert "pengukuran" in text                    # admission is by measurement
    assert "151 mg/dL baru dihitung tinggi" in text  # the strict TG comparator
    assert "Simvastatin" in text                   # drug families


def test_export_route_returns_xlsx(lipid_client):
    client, pid, rc = lipid_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/lipid-reports/registry/export",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "Registri Dislipidemia - PUSKESMAS TEBET" in (
        r.headers["content-disposition"]
    )
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Registri Dislipidemia", "Formula & Logika"]
