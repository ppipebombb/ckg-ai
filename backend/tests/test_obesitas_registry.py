"""Tests for the Obesitas registry scan.

Mirrors ``test_lipid_registry.py`` (real Fernet encryption over synthetic
EPUS/ASIK blobs plus a fake Session, so the pipeline runs without Postgres), and
adds dedicated coverage for what this registry does differently from all three
siblings:

  1. **IMT is derived**, never read — and the rounding that decides the boundary
     is applied before the bands, so the label always matches the printed number;
  2. the control window is **3–6 months** (not monthly like DM, not quarterly
     like Dislipidemia);
  3. the follow-up target is a **>5% loss from the CKG baseline weight**, strict —
     exactly 5.0% is "tidak tercapai";
  4. there is **no Jenis Obat column** at all.

Plus the label-collision trap that would silently corrupt values: the CKG
anamnesis question "Apakah berat badan Anda berkurang >3 kg…?" contains the words
"berat badan".

A **past** report year (2024) is used so the control-window fill (capped at the
current month only for the current year) runs deterministically regardless of
when the test runs.
"""

import io
import json
import uuid
from datetime import date
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.api.deps import Principal, get_dashboard_principal, get_db
from app.core.security import encrypt_json
from app.main import app
from app.models.patient import MatchStatus
from app.services.obesitas_registry_export import (
    _INTERP_COL,
    _LAST_COL,
    _MONTH_STRIDE,
    _month_base_col,
    build_obesitas_registry_workbook,
)
from app.services.obesitas_registry_scan import (
    _FU_MISSED_VISIT,
    _FU_TERKENDALI,
    _FU_TIDAK_TERKENDALI,
    _INTERP_NORMAL,
    _INTERP_OBESITAS_I,
    _INTERP_OBESITAS_II,
    _control_months,
    _interpretasi_baseline,
    _interpretasi_followup,
    count_interpretasi_bands,
    imt,
    penurunan_pct,
    scan_obesitas_registry,
)

_YEAR = 2024  # past year → the control-window fill runs to December

_BERAT_TURUN_Q = (
    "Apakah berat badan Anda berkurang >3 kg dalam 3 bulan terakhir atau "
    "pakaian menjadi lebih longgar?"
)
_RIWAYAT_HT_Q = (
    "Apakah Anda pernah dinyatakan tekanan darah tinggi/hipertensi oleh Dokter?"
)
_RIWAYAT_DM_Q = (
    "Apakah Anda pernah dinyatakan diabetes atau kencing manis oleh Dokter?"
)
_GIZI_LAYANAN = "Skrining Gizi, Tekanan Darah, dan Gula Darah Perempuan"


# ── blob builders ──────────────────────────────────────────────────────────
def _epus(*, nik, nama="PASIEN X", bb=None, tb=None, icd=None, ptm=False):
    """Synthetic EPUS visit blob. ``ptm=True`` puts BB/TB in the PTM fallback
    group instead of Anamnesa → Periksa Fisik. ``icd`` goes into
    ``penyakit_khusus`` — I10 marks hipertensi, E11 diabetes."""
    antro = {}
    if bb is not None:
        antro["Berat Badan"] = str(bb)
    if tb is not None:
        antro["Tinggi Badan"] = str(tb)

    tabs: dict = {}
    if antro:
        if ptm:
            tabs["PTM"] = {"fields": {"Tekanan Darah & IMT": antro}}
        else:
            tabs["Anamnesa"] = {"fields": {"Periksa Fisik": antro}}
    icds = icd if isinstance(icd, (list, tuple)) else ([icd] if icd else [])
    return {
        "data_pasien": {
            "NIK": nik,
            "Nama Pasien": nama,
            "Jenis Kelamin": "P",
            "Tempat/Tgl Lahir": "BEKASI/ 27-04-1980",
            "No Telp / HP": "08123456789",
            "Alamat": "jl pakis RT RW Kel PEKAYON",
        },
        "penyakit_khusus": [{"ICDX": c, "Penyakit": ""} for c in icds],
        "tabs": tabs,
    }


def _asik(
    *, bb=None, tb=None, lingkar=None, berat_turun=None, riwayat_ht=None,
    riwayat_dm=None, identitas=None,
):
    """Raw ASIK blob. BB/TB live on the Gizi (BB - TB - Lingkar Perut) form
    (reached via a "Skrining Gizi…" layanan); the HT/DM riwayat answers live on
    their own layanan, exactly as the sibling registries read them."""
    form = {}
    if berat_turun is not None:
        form[_BERAT_TURUN_Q] = berat_turun
    if bb is not None:
        form["Berat Badan (Kg)"] = str(bb)
    if tb is not None:
        form["Pengukuran Tinggi Badan (cm)"] = str(tb)
    if lingkar is not None:
        form["Pengukuran Lingkar Perut"] = str(lingkar)

    nakes = [{"layanan": _GIZI_LAYANAN, "form_data": form}]
    if riwayat_ht is not None:
        nakes.append({
            "layanan": "Tekanan Darah Dewasa Lansia",
            "form_data": {_RIWAYAT_HT_Q: riwayat_ht},
        })
    if riwayat_dm is not None:
        nakes.append({
            "layanan": "Pemeriksaan Gula Darah Dewasa Lansia",
            "form_data": {_RIWAYAT_DM_Q: riwayat_dm},
        })

    out = {"pelayanan_nakes": nakes, "pemeriksaan_mandiri": []}
    if identitas is not None:
        out["detail_data"] = identitas
    return out


def _row(nik, fd, status, epus, *, asik=None, group=None, nama="PASIEN X"):
    return SimpleNamespace(
        id=uuid.uuid4(),
        nik=nik,
        nama=nama,
        filter_date=fd,
        match_status=status,
        match_group_id=group,
        scraped_epus_data=encrypt_json(epus),
        scraped_asik_data=encrypt_json(asik) if asik is not None else None,
    )


class _FakeResult:
    def __init__(self, rows):
        self._rows = rows

    def __iter__(self):
        return iter(self._rows)


class _FakeDB:
    """Returns the main blob rows on the first execute(), and the baseline ASIK
    rows on the second (the scan issues exactly those two queries)."""

    def __init__(self, rows):
        self._main = sorted(rows, key=lambda r: (r.nik, r.filter_date))
        self._asik = [
            SimpleNamespace(id=r.id, scraped_asik_data=r.scraped_asik_data)
            for r in rows
        ]
        self.calls = 0

    def execute(self, _stmt):
        self.calls += 1
        return _FakeResult(self._main if self.calls == 1 else self._asik)


def _scan(rows):
    payload = scan_obesitas_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    return {p["nik"]: p for p in payload["patients"]}


M, E = MatchStatus.MATCHED, MatchStatus.EPUS_ONLY


def _a(bb=None, tb=None):
    return {"bb": bb, "tb": tb}


# ── IMT derivation ─────────────────────────────────────────────────────────
def test_imt_formula_and_rounding():
    assert imt(78, 165) == 28.7      # 78 / 1.65^2 = 28.65…
    assert imt(90, 170) == 31.1
    assert imt(50, 170) == 17.3
    assert imt(None, 170) is None
    assert imt(70, None) is None
    assert imt(70, 0) is None


def test_imt_rounds_before_the_bands_are_applied():
    """The label must match the number printed next to it. 68 kg / 165 cm is
    24.977…, which the sheet's own 1-decimal precision makes 25.0 — so the row
    IS Obesitas I. Classifying the unrounded value would show '25.0 / Normal'
    and drop a patient the export's live formula keeps."""
    assert imt(68, 165) == 25.0
    assert _interpretasi_baseline(_a(68, 165)) == _INTERP_OBESITAS_I


# ── pure classification: the bands, at their exact boundaries ──────────────
@pytest.mark.parametrize(
    "bb,tb,expected",
    [
        # IMT 24.9 → below the syarat
        (67.8, 165, _INTERP_NORMAL),
        # IMT exactly 25.0 → Obesitas I (inclusive lower bound)
        (68.1, 165, _INTERP_OBESITAS_I),
        # IMT 29.9 → still Obesitas I
        (81.4, 165, _INTERP_OBESITAS_I),
        # IMT exactly 30.0 → Obesitas II (inclusive)
        (81.7, 165, _INTERP_OBESITAS_II),
        (120, 165, _INTERP_OBESITAS_II),
    ],
)
def test_band_boundaries(bb, tb, expected):
    assert _interpretasi_baseline(_a(bb, tb)) == expected


def test_the_29_9_to_30_gap_is_closed_upward():
    """B1 writes 'IMT 25-29.9' and 'IMT >= 30', which literally leaves
    29.9 < IMT < 30 unlabelled. Obesitas I must own it — a patient at 29.95 may
    not silently fall out of the registry."""
    assert imt(81.5, 165) == 29.9
    assert _interpretasi_baseline(_a(81.5, 165)) == _INTERP_OBESITAS_I


def test_baseline_empty_is_blank_not_normal():
    """An unmeasured patient must never read 'Normal' — that would tell a
    puskesmas their patient's weight is fine when nobody weighed them."""
    assert _interpretasi_baseline(_a()) == ""
    assert _interpretasi_baseline(_a(bb=90)) == ""  # weight but no height
    assert _interpretasi_baseline(_a(tb=165)) == ""  # height but no weight


# ── follow-up target ───────────────────────────────────────────────────────
def test_penurunan_pct_is_positive_for_a_loss():
    assert penurunan_pct(76, 80) == 5.0
    assert penurunan_pct(84, 80) == -5.0  # gained weight
    assert penurunan_pct(None, 80) is None
    assert penurunan_pct(76, None) is None
    assert penurunan_pct(76, 0) is None


def test_followup_target_is_strictly_more_than_five_percent():
    """The legend words the two sides as '>5%' and '<=5%', so exactly 5.0% is
    NOT tercapai. Tidying this into >= silently promotes a whole band."""
    assert _interpretasi_followup(_a(bb=76), 80) == _FU_TIDAK_TERKENDALI  # 5.0%
    assert _interpretasi_followup(_a(bb=75.9), 80) == _FU_TERKENDALI      # 5.1%
    assert _interpretasi_followup(_a(bb=85), 80) == _FU_TIDAK_TERKENDALI  # gained


def test_followup_without_a_weight_is_blank_not_failed():
    """A height-only follow-up visit is listed (the sheet has a TB column every
    month) but carries no verdict — there is nothing to compare."""
    assert _interpretasi_followup(_a(tb=165), 80) == ""
    assert _interpretasi_followup(_a(bb=76), None) == ""


# ── the 3-6 month control window ───────────────────────────────────────────
def test_control_months_are_a_three_to_six_month_window():
    """The divergence from BOTH siblings: DM stamps every month, Dislipidemia
    stamps every 3rd."""
    today = date(2026, 7, 21)
    assert _control_months(1, 2024, today) == [4, 5, 6, 7]
    assert _control_months(2, 2024, today) == [5, 6, 7, 8]
    assert _control_months(9, 2024, today) == [12]  # +4 would be next year
    assert _control_months(10, 2024, today) == []


def test_control_months_capped_at_current_month_in_current_year():
    """A control that has not come due yet is not a missed one."""
    today = date(2026, 7, 21)
    assert _control_months(1, 2026, today) == [4, 5, 6, 7]
    assert _control_months(3, 2026, today) == [6, 7]
    assert _control_months(6, 2026, today) == []  # Sept is still in the future


# ── syarat: IMT-only ───────────────────────────────────────────────────────
def test_obese_patient_is_admitted():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
    ]
    p = _scan(rows)["A"]
    assert p["bb"] == 90 and p["tb"] == 165
    assert p["imt"] == 33.1
    assert p["interpretasi"] == _INTERP_OBESITAS_II


def test_non_obese_patient_is_excluded():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"), asik=_asik(bb=55, tb=165)),
    ]
    assert _scan(rows) == {}


def test_patient_without_both_measurements_is_excluded():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"), asik=_asik(bb=90)),
        _row("B", date(_YEAR, 3, 1), M, _epus(nik="B"), asik=_asik(tb=165)),
        _row("C", date(_YEAR, 3, 1), M, _epus(nik="C"), asik=_asik()),
    ]
    assert _scan(rows) == {}


def test_riwayat_ht_or_dm_alone_does_NOT_admit():
    """Like the Dislipidemia registry and unlike HT/DM: a recorded diagnosis is
    NOT a ticket in. The sheet admits on the IMT measurement alone; Riwayat HT/DM
    are columns H and I — reported, never filtered on."""
    rows = [
        # both diagnoses recorded, IMT normal → still excluded
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", icd=["I10", "E11"]),
             asik=_asik(bb=55, tb=165, riwayat_ht="Ya", riwayat_dm="Ya")),
        # no diagnosis at all, IMT 33 → admitted
        _row("B", date(_YEAR, 3, 1), M, _epus(nik="B"), asik=_asik(bb=90, tb=165)),
    ]
    assert set(_scan(rows)) == {"B"}


def test_riwayat_does_not_pin_the_baseline_label():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", icd=["I10", "E11"]),
             asik=_asik(bb=70, tb=165)),  # IMT 25.7 → Obesitas I
    ]
    p = _scan(rows)["A"]
    assert p["riwayat_ht"] == "Ya" and p["riwayat_dm"] == "Ya"
    assert p["interpretasi"] == _INTERP_OBESITAS_I  # from IMT, not from riwayat


def test_no_matched_visit_means_no_row():
    rows = [_row("A", date(_YEAR, 3, 1), E, _epus(nik="A", bb=90, tb=165))]
    assert _scan(rows) == {}


# ── riwayat HT / DM sourcing ───────────────────────────────────────────────
def test_riwayat_from_asik_answers():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"),
             asik=_asik(bb=90, tb=165, riwayat_ht="Ya", riwayat_dm="Tidak")),
    ]
    p = _scan(rows)["A"]
    assert p["riwayat_ht"] == "Ya" and p["riwayat_dm"] == "Tidak"
    assert p["sources"]["riwayat_ht"] == "ASIK"
    assert p["sources"]["riwayat_dm"] is None


def test_riwayat_falls_back_to_epus_diagnosis():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", icd="E11"),
             asik=_asik(bb=90, tb=165)),
    ]
    p = _scan(rows)["A"]
    assert p["riwayat_dm"] == "Ya"
    assert p["sources"]["riwayat_dm"] == "EPUS"
    assert p["riwayat_ht"] == "Tidak"


# ── value sourcing: ASIK first, EPUS per-field fallback ────────────────────
def test_asik_wins_per_field_and_epus_fills_the_gaps():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", bb=61, tb=165),
             asik=_asik(bb=90)),
    ]
    p = _scan(rows)["A"]
    assert p["bb"] == 90    # ASIK wins (61 would have been IMT 22.4 → excluded)
    assert p["tb"] == 165   # EPUS fills the gap — without it there is no IMT
    assert p["sources"]["antropometri"] == "ASIK"


def test_epus_only_marks_the_source_as_epus():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", bb=90, tb=165),
             asik=_asik()),
    ]
    p = _scan(rows)["A"]
    assert p["imt"] == 33.1
    assert p["sources"]["antropometri"] == "EPUS"


def test_epus_ptm_group_is_the_fallback_for_periksa_fisik():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", bb=90, tb=165, ptm=True),
             asik=_asik()),
    ]
    assert _scan(rows)["A"]["imt"] == 33.1


def test_asik_berat_badan_QUESTION_is_not_parsed_as_a_weight():
    """TRAP. The CKG anamnesis asks 'Apakah BERAT BADAN Anda berkurang >3 kg…?'
    — the same words as the value label. Without the question guard the answer
    lands in BB, and `_num("Ya")` → None then masks the real ePuskesmas weight,
    silently dropping the patient from the registry."""
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A", bb=90, tb=165),
             asik=_asik(berat_turun="Ya")),
    ]
    p = _scan(rows)["A"]
    assert p["bb"] == 90
    assert p["interpretasi"] == _INTERP_OBESITAS_II


def test_lingkar_perut_is_not_parsed_as_a_measurement():
    """The Gizi form ships waist circumference alongside BB/TB. The sheet has no
    waist column and 'Pengukuran Lingkar Perut' contains neither trap word — but
    a loosened matcher would file it as TB (it is a 'Pengukuran … ' in cm)."""
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"),
             asik=_asik(bb=90, tb=165, lingkar=102)),
    ]
    p = _scan(rows)["A"]
    assert p["tb"] == 165
    assert p["imt"] == 33.1


def test_implausible_values_are_treated_as_unmeasured():
    """A BB typed into the TB field (or a height in metres) must not silently
    produce an IMT that admits or excludes a patient."""
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"), asik=_asik(bb=90, tb=1.65)),
        _row("B", date(_YEAR, 3, 1), M, _epus(nik="B"), asik=_asik(bb=900, tb=165)),
    ]
    assert _scan(rows) == {}


# ── follow-up ──────────────────────────────────────────────────────────────
def test_followup_missed_visit_only_inside_the_control_window():
    """Baseline in January → controls in April–July. February, March, and
    August onward must stay ABSENT, not 'Pasien Missed Visit'."""
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
    ]
    fu = _scan(rows)["A"]["followup"]
    assert sorted(int(k) for k in fu) == [4, 5, 6, 7]
    for m in ("4", "5", "6", "7"):
        assert fu[m][0]["interpretasi"] == _FU_MISSED_VISIT
        assert fu[m][0]["tanggal"] is None


def test_followup_records_visits_outside_the_control_window_too():
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 2, 15), E, _epus(nik="A", bb=84, tb=165)),
    ]
    fu = _scan(rows)["A"]["followup"]
    assert fu["2"][0]["tanggal"] == f"{_YEAR}-02-15"
    assert fu["2"][0]["penurunan_pct"] == 6.7
    assert fu["2"][0]["interpretasi"] == _FU_TERKENDALI
    # …and February does NOT become a control month
    assert sorted(int(k) for k in fu) == [2, 4, 5, 6, 7]


def test_followup_is_measured_against_the_CKG_baseline_not_the_prior_visit():
    """Two 3% steps down are still 'tidak tercapai' against the baseline on the
    first visit and 'tercapai' on the second — measured from 90, never from the
    previous month's weight."""
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 4, 5), E, _epus(nik="A", bb=87.3, tb=165)),
        _row("A", date(_YEAR, 5, 5), E, _epus(nik="A", bb=84.7, tb=165)),
    ]
    fu = _scan(rows)["A"]["followup"]
    assert fu["4"][0]["penurunan_pct"] == 3.0
    assert fu["4"][0]["interpretasi"] == _FU_TIDAK_TERKENDALI
    assert fu["5"][0]["penurunan_pct"] == 5.9
    assert fu["5"][0]["interpretasi"] == _FU_TERKENDALI


def test_followup_uses_the_FINAL_baseline_weight_after_asik_wins():
    """Regression guard for the ordering bug this scan is written to avoid: the
    follow-up verdict is computed AFTER the ASIK pass, so it is measured against
    the weight actually shown in the baseline column. Here ePuskesmas says 80 and
    ASIK says 90 — a 85 kg follow-up is a LOSS against 90 (5.6%, tercapai) but a
    GAIN against 80."""
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A", bb=80, tb=165),
             asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 4, 5), E, _epus(nik="A", bb=85, tb=165)),
    ]
    p = _scan(rows)["A"]
    assert p["bb"] == 90
    assert p["followup"]["4"][0]["penurunan_pct"] == 5.6
    assert p["followup"]["4"][0]["interpretasi"] == _FU_TERKENDALI


def test_followup_visit_in_a_control_month_replaces_the_missed_filler():
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 4, 20), E, _epus(nik="A", bb=89, tb=165)),
    ]
    fu = _scan(rows)["A"]["followup"]
    assert len(fu["4"]) == 1
    assert fu["4"][0]["tanggal"] == f"{_YEAR}-04-20"
    assert fu["4"][0]["interpretasi"] == _FU_TIDAK_TERKENDALI


def test_followup_height_only_visit_is_listed_but_not_judged():
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 4, 5), E, _epus(nik="A", tb=165)),
    ]
    v = _scan(rows)["A"]["followup"]["4"][0]
    assert v["tanggal"] == f"{_YEAR}-04-05"
    assert v["tb"] == 165 and v["bb"] is None
    assert v["imt"] is None
    assert v["interpretasi"] == ""


def test_followup_skips_the_baseline_month():
    rows = [
        _row("A", date(_YEAR, 3, 1), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 3, 25), E, _epus(nik="A", bb=80, tb=165)),
    ]
    assert "3" not in _scan(rows)["A"]["followup"]


def test_followup_lists_every_visit_in_a_month():
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 4, 5), E, _epus(nik="A", bb=89, tb=165)),
        _row("A", date(_YEAR, 4, 25), E, _epus(nik="A", bb=84, tb=165)),
    ]
    fu4 = _scan(rows)["A"]["followup"]["4"]
    assert [v["tanggal"] for v in fu4] == [f"{_YEAR}-04-05", f"{_YEAR}-04-25"]
    assert [v["interpretasi"] for v in fu4] == [
        _FU_TIDAK_TERKENDALI, _FU_TERKENDALI,
    ]


def test_followup_visit_without_any_measurement_is_not_listed():
    rows = [
        _row("A", date(_YEAR, 1, 10), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 4, 5), E, _epus(nik="A")),  # visit, nobody weighed
    ]
    fu4 = _scan(rows)["A"]["followup"]["4"]
    assert len(fu4) == 1
    assert fu4[0]["interpretasi"] == _FU_MISSED_VISIT


def test_cross_date_scrape_twins_collapse():
    gid = uuid.uuid4()
    rows = [
        _row("C", date(_YEAR, 1, 5), M, _epus(nik="C"), asik=_asik(bb=90, tb=165)),
        _row("C", date(_YEAR, 4, 5), E, _epus(nik="C", bb=88, tb=165), group=gid),
        _row("C", date(_YEAR, 4, 9), E, _epus(nik="C", bb=88, tb=165), group=gid),
    ]
    fu4 = _scan(rows)["C"]["followup"]["4"]
    assert len(fu4) == 1
    assert fu4[0]["tanggal"] == f"{_YEAR}-04-05"


# ── identitas + ordering ───────────────────────────────────────────────────
def test_identitas_asik_wins_epus_fills_gaps():
    rows = [
        _row("I", date(_YEAR, 3, 1), M, _epus(nik="I", nama="EPUS NAMA"),
             asik=_asik(bb=90, tb=165, identitas={
                 "data_individu": {
                     "Nama": "ASIK NAMA",
                     "Jenis Kelamin": "Perempuan",
                     "Tanggal Lahir": "25 Januari 2000",
                 },
                 "data_domisili": {},
             })),
    ]
    p = _scan(rows)["I"]
    assert p["nama"] == "ASIK NAMA"
    assert p["tanggal_lahir"] == "2000-01-25"
    assert p["no_tlp"] == "08123456789"  # ASIK lacks it → EPUS fills
    assert p["alamat"].startswith("jl pakis")


def test_tanggal_berkunjung_is_the_earliest_matched_visit():
    rows = [
        _row("A", date(_YEAR, 5, 1), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
        _row("A", date(_YEAR, 2, 1), M, _epus(nik="A"), asik=_asik(bb=90, tb=165)),
    ]
    assert _scan(rows)["A"]["tanggal_berkunjung"] == f"{_YEAR}-02-01"


def test_rows_sorted_by_nama_then_nik():
    rows = [
        _row("2", date(_YEAR, 3, 1), M, _epus(nik="2", nama="ZULFA"),
             asik=_asik(bb=90, tb=165)),
        _row("1", date(_YEAR, 3, 1), M, _epus(nik="1", nama="ANDI"),
             asik=_asik(bb=90, tb=165)),
    ]
    payload = scan_obesitas_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)
    assert [p["nama"] for p in payload["patients"]] == ["ANDI", "ZULFA"]


# ── band counts ────────────────────────────────────────────────────────────
def test_bands_are_exclusive_and_sum_to_total():
    """The counterpart of the Dislipidemia registry's overlapping analyte counts:
    here the two bands split one IMT scale, so they must always add up."""
    rows = [
        _row("X1", date(_YEAR, 3, 1), M, _epus(nik="X1"), asik=_asik(bb=70, tb=165)),
        _row("X2", date(_YEAR, 3, 1), M, _epus(nik="X2"), asik=_asik(bb=90, tb=165)),
        _row("X3", date(_YEAR, 3, 1), M, _epus(nik="X3"), asik=_asik(bb=100, tb=165)),
    ]
    patients = scan_obesitas_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)["patients"]
    bands = count_interpretasi_bands(patients)
    assert bands == {"obesitas_1": 1, "obesitas_2": 2}
    assert sum(bands.values()) == len(patients)


# ── routes ─────────────────────────────────────────────────────────────────
def _sample_payload():
    """A realistic scan payload produced by the real scan, so the route tests
    exercise the actual serialization contract."""
    rows = [
        _row("111", date(_YEAR, 1, 10), M, _epus(nik="111", nama="BUDI", icd="I10"),
             asik=_asik(bb=90, tb=165, riwayat_ht="Ya")),
        _row("111", date(_YEAR, 4, 20), E, _epus(nik="111", bb=84, tb=165)),
        _row("222", date(_YEAR, 2, 11), M, _epus(nik="222", nama="ANI"),
             asik=_asik(bb=70, tb=165)),
    ]
    return scan_obesitas_registry(_FakeDB(rows), uuid.uuid4(), _YEAR)


@pytest.fixture
def obesitas_client(monkeypatch, stub_redis, warm_calls):
    pid = uuid.uuid4()
    fake_db = MagicMock()
    fake_db.scalar.return_value = "PUSKESMAS TEBET"

    app.dependency_overrides[get_db] = lambda: fake_db
    app.dependency_overrides[get_dashboard_principal] = lambda: Principal(
        "admin", uuid.uuid4(), None
    )

    import app.api.routes.obesitas_report as obesitas

    monkeypatch.setattr(obesitas, "redis_client", lambda: stub_redis)
    monkeypatch.setattr(
        obesitas,
        "request_warm",
        lambda rc, cache_key, report_type, kwargs: warm_calls.append(
            (cache_key, report_type)
        ),
    )
    with TestClient(app) as client:
        yield client, pid, stub_redis
    app.dependency_overrides.clear()


def _seed(rc, pid, payload):
    from app.api.routes.obesitas_report import _registry_cache_key

    rc.set(_registry_cache_key(pid, _YEAR), json.dumps(payload))


def test_route_cold_cache_returns_computing_and_enqueues_warm(
    obesitas_client, warm_calls
):
    client, pid, _rc = obesitas_client
    r = client.get("/obesitas-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["computing"] is True
    assert body["items"] == [] and body["total"] == 0
    assert body["cache_hit"] is False
    # The warm must be dispatched under the report_type cron.warm_one knows.
    assert [t for _k, t in warm_calls] == ["obesitas_registry"]


def test_route_serves_cached_payload(obesitas_client):
    client, pid, rc = obesitas_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/obesitas-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["computing"] is False and body["cache_hit"] is True
    assert body["total"] == 2
    assert body["obesitas_1"] == 1
    assert body["obesitas_2"] == 1
    assert body["obesitas_1"] + body["obesitas_2"] == body["total"]
    assert [i["nama"] for i in body["items"]] == ["ANI", "BUDI"]  # sorted by nama
    budi = next(i for i in body["items"] if i["nama"] == "BUDI")
    assert budi["puskesmas_name"] == "PUSKESMAS TEBET"
    assert budi["tahun_pelaporan"] == _YEAR
    assert budi["riwayat_ht"] == "Ya" and budi["riwayat_dm"] == "Tidak"
    assert budi["bb"] == 90 and budi["tb"] == 165 and budi["imt"] == 33.1
    assert budi["interpretasi"] == _INTERP_OBESITAS_II
    # followup dict keys survive the str->int coercion, values keep their shape
    assert budi["followup"]["4"][0]["bb"] == 84
    assert budi["followup"]["4"][0]["penurunan_pct"] == 6.7
    assert budi["followup"]["4"][0]["interpretasi"] == _FU_TERKENDALI


def test_route_search_filters_and_recounts_bands(obesitas_client):
    client, pid, rc = obesitas_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/obesitas-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR, "q": "ani"})
    body = r.json()
    assert body["total"] == 1
    assert [i["nama"] for i in body["items"]] == ["ANI"]
    # Bands respect the filter, like `total`.
    assert body["obesitas_1"] == 1 and body["obesitas_2"] == 0


def test_route_search_matches_nik(obesitas_client):
    client, pid, rc = obesitas_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/obesitas-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR, "q": "111"})
    assert [i["nik"] for i in r.json()["items"]] == ["111"]


def test_route_summary(obesitas_client):
    client, pid, rc = obesitas_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/obesitas-reports/registry/summary",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 2
    assert body["obesitas_1"] == 1 and body["obesitas_2"] == 1
    assert body["riwayat_ht_ya"] == 1
    assert body["riwayat_dm_ya"] == 0
    assert body["computing"] is False


def test_route_summary_cold_cache_enqueues_warm(obesitas_client, warm_calls):
    client, pid, _rc = obesitas_client
    r = client.get("/obesitas-reports/registry/summary",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.json()["computing"] is True
    assert [t for _k, t in warm_calls] == ["obesitas_registry"]


def test_route_clear_cache_deletes_and_requeues(obesitas_client, warm_calls):
    from app.api.routes.obesitas_report import _registry_cache_key

    client, pid, rc = obesitas_client
    _seed(rc, pid, _sample_payload())
    key = _registry_cache_key(pid, _YEAR)
    assert rc.get(key) is not None

    r = client.request("DELETE", "/obesitas-reports/registry/cache",
                       params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 204
    assert rc.get(key) is None
    assert [t for _k, t in warm_calls] == ["obesitas_registry"]


def test_route_404_when_puskesmas_missing(obesitas_client):
    client, pid, _rc = obesitas_client
    app.dependency_overrides[get_db] = lambda: SimpleNamespace(
        scalar=lambda *_a, **_k: None
    )
    r = client.get("/obesitas-reports/registry",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 404


# ── xlsx export ────────────────────────────────────────────────────────────
def _wb(patients=None):
    payload = _sample_payload()
    content = build_obesitas_registry_workbook(
        "PUSKESMAS TEBET", _YEAR,
        patients if patients is not None else payload["patients"],
    )
    return openpyxl.load_workbook(io.BytesIO(content))


def test_export_grid_geometry_matches_the_dirjen_layout():
    """Identitas A-I (incl. Tanggal Berkunjung + BOTH riwayat columns, which this
    sheet files under Identitas), antropometri J-L, interpretasi M, then 12
    months x 5 columns starting at N (14) → last column BU (73). No Jenis Obat
    column anywhere."""
    assert _MONTH_STRIDE == 5
    assert _INTERP_COL == 13
    assert _month_base_col(1) == 14
    assert _month_base_col(12) == 14 + 11 * 5
    assert _LAST_COL == 73


def test_export_sheets_and_headers():
    wb = _wb()
    assert wb.sheetnames == ["Registri Obesitas", "Formula & Logika"]
    ws = wb["Registri Obesitas"]
    assert ws.cell(7, 1).value == "NIK"
    assert ws.cell(7, 7).value == "Tanggal Berkunjung"
    assert ws.cell(7, 8).value == "Riwayat diagnosis HT"
    assert ws.cell(7, 9).value == "Riwayat diagnosis DM"
    assert [ws.cell(7, c).value.split("\n")[0] for c in range(10, 13)] == [
        "BB", "TB", "IMT",
    ]
    assert ws.cell(5, 10).value == "Hasil Pemeriksaan Antopometri"
    assert ws.cell(6, 14).value == "Januari"
    bc = _month_base_col(4)
    assert ws.cell(7, bc).value.startswith("Tanggal Pemeriksaan")
    assert [ws.cell(7, bc + i).value.split("\n")[0] for i in range(1, 4)] == [
        "BB", "TB", "IMT",
    ]
    assert ws.cell(7, bc + 4).value == "Interpretasi hasil"


def test_export_writes_patient_values():
    wb = _wb()
    ws = wb["Registri Obesitas"]
    # Patients are sorted by nama → ANI first, then BUDI.
    assert ws.cell(8, 1).value == "222"
    assert ws.cell(8, 2).value == "ANI"
    budi_row = next(
        r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == "BUDI"
    )
    assert ws.cell(budi_row, 8).value == "Ya"      # Riwayat HT
    assert ws.cell(budi_row, 9).value == "Tidak"   # Riwayat DM
    assert ws.cell(budi_row, 10).value == 90       # BB
    assert ws.cell(budi_row, 11).value == 165      # TB
    # April follow-up: BB 84 in the month-4 group.
    bc = _month_base_col(4)
    assert ws.cell(budi_row, bc + 1).value == 84


def test_export_months_outside_the_control_window_stay_empty():
    """The workbook must agree with the dashboard: a month outside the 3-6 window
    has no cells at all, so it can never render as 'Pasien Missed Visit'."""
    ws = _wb()["Registri Obesitas"]
    budi_row = next(
        r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == "BUDI"
    )
    # BUDI's baseline is January → controls are April, May, June, July.
    for m in (2, 3, 8, 9, 10, 11, 12):
        bc = _month_base_col(m)
        assert ws.cell(budi_row, bc + 4).value is None, f"month {m} should be blank"
    for m in (5, 6, 7):
        bc = _month_base_col(m)
        assert str(ws.cell(budi_row, bc + 4).value).startswith("=")


def test_export_imt_columns_are_live_formulas():
    ws = _wb()["Registri Obesitas"]
    baseline_imt = ws.cell(8, 12).value
    assert isinstance(baseline_imt, str) and baseline_imt.startswith("=")
    assert "ROUND(N(J8)/(N(K8)/100)^2,1)" in baseline_imt


def test_export_interpretasi_reads_the_imt_cell_not_the_raw_values():
    """The label must never disagree with the IMT printed beside it — including
    at the boundary, where the 1-decimal rounding decides."""
    ws = _wb()["Registri Obesitas"]
    baseline = ws.cell(8, _INTERP_COL).value
    assert isinstance(baseline, str) and baseline.startswith("=")
    assert "N(L8)>=30" in baseline and "N(L8)>=25" in baseline
    assert _INTERP_OBESITAS_I in baseline and _INTERP_OBESITAS_II in baseline
    # Riwayat (H, I) must NOT appear — a diagnosis does not pin the label here.
    assert "H8" not in baseline and "I8" not in baseline


def test_export_followup_formula_compares_against_the_block_baseline_bb():
    ws = _wb()["Registri Obesitas"]
    budi_row = next(
        r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == "BUDI"
    )
    fu = ws.cell(budi_row, _month_base_col(4) + 4).value
    assert isinstance(fu, str) and fu.startswith("=")
    assert f"J{budi_row}" in fu  # the merged baseline BB cell
    assert _FU_MISSED_VISIT in fu
    assert _FU_TERKENDALI in fu and _FU_TIDAK_TERKENDALI in fu


def test_export_followup_formula_keeps_the_strict_five_percent_comparator():
    ws = _wb()["Registri Obesitas"]
    budi_row = next(
        r for r in range(8, ws.max_row + 1) if ws.cell(r, 2).value == "BUDI"
    )
    fu = ws.cell(budi_row, _month_base_col(4) + 4).value
    assert "*100,1)>5," in fu
    assert ">=5," not in fu


def test_export_survives_a_backslash_in_patient_text():
    """Regression shared with the Dislipidemia export: the sheetData substitution
    must pass a CALLABLE to re.sub, not the rendered XML string — otherwise one
    backslash in any free-text field ("JL X RT.01\\RW.02") raises re.PatternError
    and 500s the download for the whole puskesmas."""
    p = dict(_sample_payload()["patients"][0])
    p["nama"] = "BU\\DI"          # \D  -> "bad escape"
    p["alamat"] = "JL 1\\2 RT"    # \2  -> "invalid group reference"
    p["no_tlp"] = "08\\g<x>"      # \g< -> "unknown group name"
    content = build_obesitas_registry_workbook("PKM A\\B", _YEAR, [p])
    ws = openpyxl.load_workbook(io.BytesIO(content))["Registri Obesitas"]
    assert ws.cell(8, 2).value == "BU\\DI"
    assert ws.cell(8, 6).value == "JL 1\\2 RT"
    assert "PKM A\\B" in str(ws.cell(3, 1).value)


def test_export_empty_registry_still_valid_workbook():
    wb = _wb(patients=[])
    ws = wb["Registri Obesitas"]
    assert ws.cell(7, 1).value == "NIK"  # headers survive
    assert ws.cell(8, 1).value is None   # no data rows


def test_export_formula_sheet_documents_the_rules():
    """The rules a reader would otherwise carry over wrong from the sibling
    kertas kerja must be written down for the puskesmas staff reading the file."""
    ws = _wb()["Formula & Logika"]
    text = "\n".join(
        str(ws.cell(r, c).value or "")
        for r in range(1, ws.max_row + 1)
        for c in (1, 2)
    )
    assert "3-6 bulan" in text                    # the control window
    assert "bukan syarat" in text.casefold()      # riwayat is informational
    assert "IMT = Berat Badan (kg) dibagi kuadrat Tinggi Badan (meter)" in text
    assert "penurunan tepat 5,0% belum dihitung tercapai" in text
    assert "1 Januari 2026" in text               # the cut-off note from B1


def test_export_route_returns_xlsx(obesitas_client):
    client, pid, rc = obesitas_client
    _seed(rc, pid, _sample_payload())
    r = client.get("/obesitas-reports/registry/export",
                   params={"puskesmas_id": str(pid), "year": _YEAR})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "Registri Obesitas - PUSKESMAS TEBET" in (
        r.headers["content-disposition"]
    )
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Registri Obesitas", "Formula & Logika"]
