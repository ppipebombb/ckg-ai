"""Build the CKG-Sekolah form/question catalog from the Kemenkes metadata xlsx.

Source: `documents/Metadata Program Rutin(1).xlsx`, sheet
`Update USEKREM Padanan Paket Od` (the "[Update] USEKREM Padanan Paket Odoo
2025-2026" tab). That whole sheet is the single module "Usia Sekolah dan Remaja"
(the other col-1 labels — Gula Darah, Keswa, Malaria, … — are sub-groups, not
separate modules), so reading the sheet end-to-end captures EVERY school question.

THE XLSX IS A REFERENCE, NOT GROUND TRUTH. Live ASIK is the authority (the scraper
reads forms dynamically). Nothing here is hardcoded: every form, question, answer
option and klaster membership is read from the sheet, and live-verification +
spec↔live drift are COMPUTED against the live scrape map (asik_sekolah_form_mapping.json)
passed in — so the catalog re-derives cleanly when either input changes. The scraper
does NOT consume this file; it exists so coverage/drift is auditable with no data missed.

Sheet encoding (discovered 2026-06-29):
- Row 1: one klaster name per COLUMN PAIR from col 23 (`Kelas 1`, `Kelas 2 & 3`, …
  then age-based `Anak/Remaja N Tahun`); each klaster spans 2 cols (STG, PROD).
  Membership = `✅` in the klaster's PROD (EVEN) column.
- Per row: Nama Layanan (2), Kode Layanan (3), Nama Paket (5), **Kode Paket = FRM
  form code (6)** = the scraper's `pemeriksaan_code`, Nama Parameter (8),
  **Label Pertanyaan (11)** = the question, **Parameter Line (12)** = one answer
  option, Tipe (21) ∈ {`Layanan` = Pelayanan Nakes, `Skrining Mandiri` = self-report}.
  Merged Layanan/Paket cells carry forward; a question's options are the col-12
  values on its row + the following rows until the next Label Pertanyaan.

Usage:
  python build_catalog_from_xlsx.py <metadata.xlsx> <out_catalog.json> [<live_form_map.json>]
"""

import json
import sys
from pathlib import Path

import openpyxl

SHEET = "Update USEKREM Padanan Paket Od"
MODULE = "Usia Sekolah dan Remaja"

CAVEATS = [
    "REFERENCE ONLY — the 2026 spec (Padanan Odoo), not live ASIK. Live ASIK is "
    "ground truth; the scraper reads forms dynamically and this file is never "
    "consumed by it.",
    "The klaster ✅-matrix is INCOMPLETE: some live Nakes forms carry no klaster "
    "mark here (see reconciliation.live_not_in_spec).",
    "FRM codes/grouping drift vs live (e.g. NTDs split live but consolidated in "
    "spec; Hepatitis renumbered) — see reconciliation.",
    "Question lists include server-COMPUTED interpretations and CONDITIONAL "
    "follow-ups. ASIK derives/gates these; the scraper captures the nurse-entered "
    "INPUT questions only, so a live form legitimately shows fewer fields.",
    "Scope: the scraper captures Tipe='Layanan' (Pelayanan Nakes) only. "
    "Tipe='Skrining Mandiri' (self-report) is catalogued here for completeness but "
    "intentionally not scraped (pelayanan_nakes_only).",
    "Other CKG programs (BBL / Balita-Prasekolah / Dewasa-Lansia) live in separate "
    "sheets and are out of CKG-Sekolah scope.",
]


def _clean_code(v) -> str | None:
    return "".join(str(v).split()) if v not in (None, "") else None


def _norm_klaster(name: str) -> str:
    # spacing differs (xlsx "Kelas 11&12 Laki-laki" vs live "Kelas 11 & 12 Laki-laki")
    return " ".join(str(name).replace("&", " & ").split()).lower()


def extract(xlsx_path: Path) -> tuple[dict, list[str], list[str]]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET]

    klaster_cols: list[tuple[int, str]] = []
    for c in range(23, ws.max_column + 1, 2):
        nm = ws.cell(1, c).value
        if nm:
            klaster_cols.append((c + 1, str(nm).strip()))
    all_klaster = [n for _, n in klaster_cols]
    school_klaster = [n for n in all_klaster if n.lower().startswith("kelas")]

    forms: dict[str, dict] = {}
    cur_lay = cur_pak = cur_frm = cur_q = None
    for r in range(6, ws.max_row + 1):
        if ws.cell(r, 2).value:
            cur_lay = str(ws.cell(r, 2).value).strip()
        if ws.cell(r, 5).value:
            cur_pak = str(ws.cell(r, 5).value).strip()
            cur_frm = _clean_code(ws.cell(r, 6).value)
            # Drop broken-reference / non-FRM cells (e.g. "#REF!") — rows under
            # them carry no real form code, so skip until the next valid paket.
            if cur_frm and not cur_frm.startswith("FRM"):
                cur_frm = None
        label = ws.cell(r, 11).value
        line = ws.cell(r, 12).value  # one answer option
        if cur_frm is None:
            continue
        f = forms.setdefault(cur_frm, {
            "layanan": cur_lay, "paket": cur_pak,
            "tipe": None, "klaster": set(), "questions": [],
        })
        tipe = ws.cell(r, 21).value
        if not f["tipe"] and tipe:
            f["tipe"] = str(tipe).strip()
        f["klaster"] |= {n for pc, n in klaster_cols if ws.cell(r, pc).value not in (None, "")}
        if label not in (None, ""):
            # new question on this form
            cur_q = {
                "parameter": (str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value else None),
                "label": str(label).strip(),
                "options": [],
            }
            f["questions"].append(cur_q)
        if line not in (None, "") and f["questions"]:
            opt = str(line).strip()
            if opt not in f["questions"][-1]["options"]:
                f["questions"][-1]["options"].append(opt)
        cur_q = cur_q  # (options attach to the most-recent question of THIS form)
    return forms, all_klaster, school_klaster


def reconcile(forms: dict, live_map_path: Path | None) -> dict:
    """Compute spec↔live drift for whatever klaster the live map covers."""
    if not live_map_path or not live_map_path.exists():
        return {"note": "no live map supplied — drift not computed"}
    live = json.loads(live_map_path.read_text())
    live_forms = live.get("forms") or {}
    live_codes = set(live_forms)
    # live map is single-klaster today; read it from any form's klaster list
    live_klaster = None
    for d in live_forms.values():
        kl = d.get("klaster") or []
        if kl:
            live_klaster = kl[0]
            break
    out = {"live_klaster": live_klaster, "live_form_count": len(live_codes)}
    if live_klaster:
        target = _norm_klaster(live_klaster)
        spec_nakes = {
            frm for frm, d in forms.items()
            if d["tipe"] == "Layanan" and any(_norm_klaster(k) == target for k in d["klaster"])
        }
        out["spec_nakes_for_klaster"] = sorted(spec_nakes)
        out["live_not_in_spec"] = sorted(live_codes - spec_nakes)   # live forms the spec matrix omits
        out["spec_not_in_live"] = sorted(spec_nakes - live_codes)   # spec forms absent/renumbered live
    return out


def build(xlsx_path: Path, live_map_path: Path | None) -> dict:
    forms, all_klaster, school_klaster = extract(xlsx_path)
    live_codes = set()
    if live_map_path and live_map_path.exists():
        live_codes = set((json.loads(live_map_path.read_text()).get("forms") or {}))

    order = {n: i for i, n in enumerate(all_klaster)}
    ordered = {}
    total_q = total_opt = 0
    for frm in sorted(forms):
        d = forms[frm]
        total_q += len(d["questions"])
        total_opt += sum(len(q["options"]) for q in d["questions"])
        ordered[frm] = {
            "layanan": d["layanan"],
            "paket": d["paket"],
            "tipe": d["tipe"],
            "klaster": sorted(d["klaster"], key=lambda n: order.get(n, 999)),
            "live_verified": frm in live_codes,
            "questions": d["questions"],
        }
    nakes = sum(1 for d in ordered.values() if d["tipe"] == "Layanan")
    mandiri = sum(1 for d in ordered.values() if d["tipe"] == "Skrining Mandiri")
    return {
        "version": 2,
        "generated_from": f"documents/Metadata Program Rutin(1).xlsx :: '{SHEET}'",
        "module": MODULE,
        "caveats": CAVEATS,
        "counts": {
            "forms_total": len(ordered),
            "nakes_layanan": nakes,
            "skrining_mandiri": mandiri,
            "untyped": len(ordered) - nakes - mandiri,
            "questions_total": total_q,
            "answer_options_total": total_opt,
            "klaster_total": len(all_klaster),
            "school_klaster": len(school_klaster),
        },
        "all_klaster": all_klaster,
        "school_klaster": school_klaster,
        "reconciliation": reconcile(forms, live_map_path),
        "forms": ordered,
    }


def main() -> None:
    if len(sys.argv) not in (3, 4):
        print(__doc__)
        sys.exit(1)
    xlsx, out = Path(sys.argv[1]), Path(sys.argv[2])
    live = Path(sys.argv[3]) if len(sys.argv) == 4 else None
    doc = build(xlsx, live)
    out.write_text(json.dumps(doc, ensure_ascii=False, indent=2))
    c = doc["counts"]
    print(f"wrote {c['forms_total']} forms ({c['nakes_layanan']} Nakes / "
          f"{c['skrining_mandiri']} Mandiri / {c['untyped']} untyped) · "
          f"{c['questions_total']} questions · {c['answer_options_total']} options · "
          f"{c['klaster_total']} klaster → {out}")
    r = doc["reconciliation"]
    if "live_not_in_spec" in r:
        print(f"reconcile[{r['live_klaster']}]: live={r['live_form_count']} "
              f"live_not_in_spec={r['live_not_in_spec']} spec_not_in_live={r['spec_not_in_live']}")


if __name__ == "__main__":
    main()
